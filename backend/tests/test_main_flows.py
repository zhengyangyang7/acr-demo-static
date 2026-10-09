from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook, load_workbook

from database import get_conn, hash_password


def auth_header(token):
    return {'Authorization': f'Bearer {token}'}


def login_admin(client):
    r = client.post('/api/auth/login', json={'username': 'Cookie.gu', 'password': 'TestPass1'})
    data = r.get_json()
    assert r.status_code == 200 and data.get('success'), data
    return data['token']


def _insert_employee(name='张三'):
    conn = get_conn()
    cur = conn.execute(
        """INSERT INTO employees (name, role, dept, email, password, lark_only, status, created_at)
           VALUES (?, '普通用户', 'IT', ?, ?, 0, '在职', datetime('now','localtime'))""",
        (name, f'{name}@test.com', hash_password('UserPass1')),
    )
    conn.commit()
    return cur.lastrowid


def _insert_asset(**kwargs):
    defaults = {
        'asset_no': 'A-001',
        'category': '笔记本',
        'status': '在库',
        'location': '深圳',
        'assignee': '',
        'purchase_price': None,
        'purchase_date': None,
        'asset_type_code': None,
        'license_type': None,
        'expire_date': None,
        'renew_remind_days': 30,
        'warranty_expire': None,
    }
    defaults.update(kwargs)
    conn = get_conn()
    cur = conn.execute(
        """INSERT INTO assets (asset_no, category, status, location, assignee,
             purchase_price, purchase_date, asset_type_code, license_type, expire_date,
             renew_remind_days, warranty_expire)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
        (defaults['asset_no'], defaults['category'], defaults['status'], defaults['location'],
         defaults['assignee'], defaults['purchase_price'], defaults['purchase_date'],
         defaults['asset_type_code'], defaults['license_type'], defaults['expire_date'],
         defaults['renew_remind_days'], defaults['warranty_expire']),
    )
    conn.commit()
    return cur.lastrowid


def _insert_consumable(item_no='C-001', quantity=10, min_stock=5):
    conn = get_conn()
    cur = conn.execute(
        """INSERT INTO consumables (item_no, name, category, quantity, min_stock, status, location)
           VALUES (?, '硒鼓', '硒鼓', ?, ?, '正常', '仓库')""",
        (item_no, quantity, min_stock),
    )
    conn.commit()
    return cur.lastrowid


def test_admin_login_ok_and_bad_password(client):
    ok = client.post('/api/auth/login', json={'username': 'Cookie.gu', 'password': 'TestPass1'})
    assert ok.status_code == 200
    assert ok.get_json()['success'] is True
    bad = client.post('/api/auth/login', json={'username': 'Cookie.gu', 'password': 'wrong'})
    assert bad.status_code == 401
    assert bad.get_json()['success'] is False


def test_stock_out_stays_in_stock_until_approved(client):
    token = login_admin(client)
    aid = _insert_asset(asset_no='SO-001')
    r = client.post('/api/stock/out', json={
        'asset_id': aid, 'target_user': '李四', 'to_location': '深圳',
    }, headers=auth_header(token))
    assert r.status_code == 200
    assert r.get_json().get('success') is True
    before = client.get(f'/api/assets/{aid}', headers=auth_header(token)).get_json()
    assert before['status'] == '在库'
    listed = client.get('/api/approvals?apply_type=资产出库', headers=auth_header(token)).get_json()
    app_id = listed['data'][0]['id']
    appr = client.post(f'/api/approvals/{app_id}/approve', json={
        'apply_type': '资产出库', 'status': '已通过',
    }, headers=auth_header(token))
    assert appr.status_code == 200
    after = client.get(f'/api/assets/{aid}', headers=auth_header(token)).get_json()
    assert after['status'] == '在用'
    assert after['assignee'] == '李四'


def test_borrow_approve_and_reject(client, app_mod):
    admin = login_admin(client)
    eid = _insert_employee()
    user_token = app_mod.generate_token(eid)
    aid_ok = _insert_asset(asset_no='BR-OK')
    aid_no = _insert_asset(asset_no='BR-NO')

    created = client.post('/api/my/applications', json={
        'app_type': 'borrow', 'asset_id': aid_ok, 'asset_no': 'BR-OK', 'reason': '出差',
    }, headers=auth_header(user_token))
    assert created.status_code == 200
    assert created.get_json()['success']
    still = client.get(f'/api/assets/{aid_ok}', headers=auth_header(admin)).get_json()
    assert still['status'] == '在库'
    assert not still.get('assignee')

    listed = client.get('/api/approvals?apply_type=资产申请', headers=auth_header(admin)).get_json()
    app_id = next(x['id'] for x in listed['data'] if x['info'].get('asset_no') == 'BR-OK')
    appr = client.post(f'/api/approvals/{app_id}/approve', json={
        'apply_type': '资产申请', 'status': '已通过',
    }, headers=auth_header(admin))
    assert appr.status_code == 200, appr.get_json()
    used = client.get(f'/api/assets/{aid_ok}', headers=auth_header(admin)).get_json()
    assert used['status'] == '在用'
    assert used['assignee'] == '张三'

    created2 = client.post('/api/my/applications', json={
        'app_type': 'borrow', 'asset_id': aid_no, 'asset_no': 'BR-NO',
    }, headers=auth_header(user_token))
    listed2 = client.get('/api/approvals?apply_type=资产申请', headers=auth_header(admin)).get_json()
    rej_id = next(x['id'] for x in listed2['data'] if x['info'].get('asset_no') == 'BR-NO')
    client.post(f'/api/approvals/{rej_id}/approve', json={
        'apply_type': '资产申请', 'status': '已拒绝',
    }, headers=auth_header(admin))
    rejected = client.get(f'/api/assets/{aid_no}', headers=auth_header(admin)).get_json()
    assert rejected['status'] == '在库'
    assert not rejected.get('assignee')


def test_consumable_approve_and_insufficient(client, app_mod):
    admin = login_admin(client)
    eid = _insert_employee()
    user_token = app_mod.generate_token(eid)
    cid = _insert_consumable(quantity=10, min_stock=2)

    r = client.post('/api/profile/apply/consumable', json={
        'item_id': cid, 'quantity': 3, 'purpose': '打印',
    }, headers=auth_header(user_token))
    assert r.status_code == 200
    assert get_conn().execute("SELECT quantity FROM consumables WHERE id=?", (cid,)).fetchone()['quantity'] == 10

    listed = client.get('/api/approvals?apply_type=资产申请', headers=auth_header(admin)).get_json()
    app_id = next(x['id'] for x in listed['data'] if x['info'].get('app_type') == 'consumable')
    ok = client.post(f'/api/approvals/{app_id}/approve', json={
        'apply_type': '资产申请', 'status': '已通过',
    }, headers=auth_header(admin))
    assert ok.status_code == 200, ok.get_json()
    assert get_conn().execute("SELECT quantity FROM consumables WHERE id=?", (cid,)).fetchone()['quantity'] == 7

    fail_apply = client.post('/api/profile/apply/consumable', json={
        'item_id': cid, 'quantity': 100, 'purpose': '超量',
    }, headers=auth_header(user_token))
    assert fail_apply.status_code == 400

    conn = get_conn()
    conn.execute("UPDATE consumables SET quantity=1 WHERE id=?", (cid,))
    conn.commit()
    from app import insert_asset_application
    user = conn.execute("SELECT * FROM employees WHERE id=?", (eid,)).fetchone()
    app_id2, _ = insert_asset_application(
        conn, user, 'consumable',
        detail={'item_id': cid, 'quantity': 5, 'qty': 5})
    conn.commit()
    fail = client.post(f'/api/approvals/{app_id2}/approve', json={
        'apply_type': '资产申请', 'status': '已通过',
    }, headers=auth_header(admin))
    assert fail.status_code == 400
    conn = get_conn()
    assert conn.execute("SELECT quantity FROM consumables WHERE id=?", (cid,)).fetchone()['quantity'] == 1
    assert conn.execute("SELECT status FROM asset_applications WHERE id=?", (app_id2,)).fetchone()['status'] == 'pending'


def test_book_value_computed_and_missing(client):
    token = login_admin(client)
    as_of = date.today()
    start = date(as_of.year - 1, as_of.month, min(as_of.day, 28))
    aid = _insert_asset(
        asset_no='BV-OK', purchase_price=10000,
        purchase_date=start.isoformat(), asset_type_code='LT')
    aid_empty = _insert_asset(asset_no='BV-NA')
    ok = client.get(f'/api/assets/{aid}', headers=auth_header(token)).get_json()
    assert ok['depreciation_months'] == 36
    from app import compute_book_value
    expected = compute_book_value(10000, start.isoformat(), 36, 0.05, as_of=as_of)
    assert ok['book_value'] == expected
    empty = client.get(f'/api/assets/{aid_empty}', headers=auth_header(token)).get_json()
    assert empty['book_value'] is None


def test_alerts_export_xlsx(client):
    token = login_admin(client)
    yesterday = (date.today() - timedelta(days=1)).isoformat()
    _insert_asset(
        asset_no='EXP-1', license_type='subscription',
        expire_date=yesterday, renew_remind_days=30, category='软件')
    _insert_asset(
        asset_no='WAR-1', warranty_expire=yesterday, category='笔记本')
    _insert_consumable(item_no='LOW-1', quantity=1, min_stock=10)
    r = client.get('/api/alerts/export', headers=auth_header(token))
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.data))
    assert '订阅到期' in wb.sheetnames and '低库存耗材' in wb.sheetnames
    assert '保修到期' in wb.sheetnames
    expire_rows = list(wb['订阅到期'].iter_rows(min_row=2, values_only=True))
    low_rows = list(wb['低库存耗材'].iter_rows(min_row=2, values_only=True))
    war_rows = list(wb['保修到期'].iter_rows(min_row=2, values_only=True))
    assert any(row and row[0] == 'EXP-1' for row in expire_rows)
    assert any(row and row[0] == 'LOW-1' for row in low_rows)
    assert any(row and row[0] == 'WAR-1' for row in war_rows)


def test_employee_password_login(client):
    eid = _insert_employee(name='员工甲')
    r = client.post('/api/auth/login', json={'username': '员工甲', 'password': 'UserPass1'})
    assert r.status_code == 200
    data = r.get_json()
    assert data['success'] and data['user']['role'] == '普通用户'
    assert data['user']['id'] == eid


def test_return_and_repair_approve(client, app_mod):
    admin = login_admin(client)
    eid = _insert_employee(name='员工乙')
    user_token = app_mod.generate_token(eid)
    aid = _insert_asset(asset_no='RT-1', status='在用', assignee='员工乙')
    aid2 = _insert_asset(asset_no='RP-1', status='在用', assignee='员工乙')

    ret = client.post('/api/profile/apply/return', json={
        'asset_id': aid, 'reason': '离职归还',
    }, headers=auth_header(user_token))
    assert ret.status_code == 200, ret.get_json()
    mid = client.get(f'/api/assets/{aid}', headers=auth_header(admin)).get_json()
    assert mid['status'] == '在用'

    listed = client.get('/api/approvals?apply_type=资产申请', headers=auth_header(admin)).get_json()
    ret_id = next(x['id'] for x in listed['data'] if x['info'].get('asset_no') == 'RT-1')
    ok = client.post(f'/api/approvals/{ret_id}/approve', json={
        'apply_type': '资产申请', 'status': '已通过',
    }, headers=auth_header(admin))
    assert ok.status_code == 200, ok.get_json()
    after = client.get(f'/api/assets/{aid}', headers=auth_header(admin)).get_json()
    assert after['status'] == '在库'
    assert not after.get('assignee')

    rep = client.post('/api/profile/apply/repair', json={
        'asset_id': aid2, 'fault_desc': '屏幕坏了',
    }, headers=auth_header(user_token))
    assert rep.status_code == 200
    listed2 = client.get('/api/approvals?apply_type=资产申请', headers=auth_header(admin)).get_json()
    rep_id = next(x['id'] for x in listed2['data'] if x['info'].get('asset_no') == 'RP-1')
    ok2 = client.post(f'/api/approvals/{rep_id}/approve', json={
        'apply_type': '资产申请', 'status': '已通过',
    }, headers=auth_header(admin))
    assert ok2.status_code == 200, ok2.get_json()
    repaired = client.get(f'/api/assets/{aid2}', headers=auth_header(admin)).get_json()
    assert repaired['status'] == '维修中'


def test_double_approve_rejected(client, app_mod):
    admin = login_admin(client)
    eid = _insert_employee(name='员工丙')
    user_token = app_mod.generate_token(eid)
    aid = _insert_asset(asset_no='DUP-1')
    client.post('/api/my/applications', json={
        'app_type': 'borrow', 'asset_id': aid, 'asset_no': 'DUP-1',
    }, headers=auth_header(user_token))
    listed = client.get('/api/approvals?apply_type=资产申请', headers=auth_header(admin)).get_json()
    app_id = next(x['id'] for x in listed['data'] if x['info'].get('asset_no') == 'DUP-1')
    first = client.post(f'/api/approvals/{app_id}/approve', json={
        'apply_type': '资产申请', 'status': '已通过',
    }, headers=auth_header(admin))
    assert first.status_code == 200
    second = client.post(f'/api/approvals/{app_id}/approve', json={
        'apply_type': '资产申请', 'status': '已通过',
    }, headers=auth_header(admin))
    assert second.status_code == 400
    asset = client.get(f'/api/assets/{aid}', headers=auth_header(admin)).get_json()
    assert asset['assignee'] == '员工丙'
    logs = get_conn().execute(
        "SELECT COUNT(*) as c FROM stock_logs WHERE asset_id=? AND log_type='借用出库'",
        (aid,)).fetchone()['c']
    assert logs == 1


def test_category_infers_type_code_for_book_value(client):
    token = login_admin(client)
    as_of = date.today()
    start = date(as_of.year - 1, as_of.month, min(as_of.day, 28))
    # 只填类目「笔记本」，不传 asset_type_code，创建时应回填 LT
    r = client.post('/api/assets', json={
        'category': '笔记本',
        'brand': 'Dell',
        'purchase_price': 10000,
        'purchase_date': start.isoformat(),
        'status': '在库',
        'location': '深圳',
    }, headers=auth_header(token))
    assert r.status_code == 201, r.get_json()
    aid = r.get_json()['id']
    detail = client.get(f'/api/assets/{aid}', headers=auth_header(token)).get_json()
    assert detail.get('asset_type_code') == 'LT' or detail.get('book_value') is not None
    assert detail['book_value'] is not None
    assert detail['depreciation_months'] == 36


def test_fulfill_new_asset_application(client, app_mod):
    admin = login_admin(client)
    eid = _insert_employee(name='员工丁')
    user_token = app_mod.generate_token(eid)
    stock_id = _insert_asset(asset_no='FF-STOCK', status='在库')
    created = client.post('/api/profile/apply/asset', json={
        'type': 'existing', 'category': '笔记本', 'reason': '新人入职',
    }, headers=auth_header(user_token))
    assert created.status_code == 200, created.get_json()
    app_id = created.get_json()['id']
    appr = client.post(f'/api/approvals/{app_id}/approve', json={
        'apply_type': '资产申请', 'status': '已通过',
    }, headers=auth_header(admin))
    assert appr.status_code == 200, appr.get_json()
    pending = client.get('/api/approvals/fulfillable', headers=auth_header(admin)).get_json()
    assert any(x['id'] == app_id for x in pending['data'])
    done = client.post(f'/api/approvals/{app_id}/fulfill', json={
        'asset_id': stock_id,
    }, headers=auth_header(admin))
    assert done.status_code == 200, done.get_json()
    asset = client.get(f'/api/assets/{stock_id}', headers=auth_header(admin)).get_json()
    assert asset['status'] == '在用'
    assert asset['assignee'] == '员工丁'
    again = client.get('/api/approvals/fulfillable', headers=auth_header(admin)).get_json()
    assert not any(x['id'] == app_id for x in again['data'])


def test_profile_update(client, app_mod):
    eid = _insert_employee(name='员工戊')
    token = app_mod.generate_token(eid)
    r = client.put('/api/profile/update', json={
        'dept': '行政部', 'email': 'wu@test.com', 'location': '北京',
    }, headers=auth_header(token))
    assert r.status_code == 200, r.get_json()
    user = r.get_json()['user']
    assert user['dept'] == '行政部'
    assert user['email'] == 'wu@test.com'
    assert user['location'] == '北京'


def test_stats_overview_money_warranty_fulfill(client, app_mod):
    admin = login_admin(client)
    as_of = date.today()
    start = date(as_of.year - 1, as_of.month, min(as_of.day, 28))
    yesterday = (as_of - timedelta(days=1)).isoformat()
    _insert_asset(
        asset_no='OV-1', category='笔记本', status='在库',
        purchase_price=10000, purchase_date=start.isoformat(), asset_type_code='LT')
    _insert_asset(
        asset_no='OV-2', category='笔记本', status='在库',
        purchase_price=8000, purchase_date=start.isoformat(), asset_type_code='LT')
    _insert_asset(
        asset_no='OV-SCRAP', category='笔记本', status='已报废',
        purchase_price=5000, purchase_date=start.isoformat(), asset_type_code='LT')
    _insert_asset(
        asset_no='OV-WAR', category='笔记本', status='在用',
        warranty_expire=yesterday)

    eid = _insert_employee(name='概览员工')
    user_token = app_mod.generate_token(eid)
    stock_id = _insert_asset(asset_no='OV-FF', status='在库')
    created = client.post('/api/profile/apply/asset', json={
        'type': 'existing', 'category': '笔记本', 'reason': '概览待发放',
    }, headers=auth_header(user_token))
    assert created.status_code == 200, created.get_json()
    app_id = created.get_json()['id']
    appr = client.post(f'/api/approvals/{app_id}/approve', json={
        'apply_type': '资产申请', 'status': '已通过',
    }, headers=auth_header(admin))
    assert appr.status_code == 200, appr.get_json()

    ov = client.get('/api/stats/overview', headers=auth_header(admin)).get_json()
    assert ov['purchase_value'] >= 18000
    assert ov['purchase_value'] < 23000  # 不含已报废 5000
    assert ov['book_value_total'] > 0
    assert ov['warranty_alert_count'] >= 1
    assert any(x.get('asset_no') == 'OV-WAR' for x in (ov.get('warranty_alerts') or []))
    assert ov['fulfillable_count'] >= 1

    done = client.post(f'/api/approvals/{app_id}/fulfill', json={
        'asset_id': stock_id,
    }, headers=auth_header(admin))
    assert done.status_code == 200, done.get_json()
    ov2 = client.get('/api/stats/overview', headers=auth_header(admin)).get_json()
    assert ov2['fulfillable_count'] == ov['fulfillable_count'] - 1


def test_inventory_check_and_complete(client):
    admin = login_admin(client)
    aid = _insert_asset(asset_no='INV-1', status='在库', location='深圳')
    created = client.post('/api/inventory/tasks', json={
        'task_name': '第四期盘点',
        'scope_type': 'in_stock',
        'created_by': 'Cookie.gu',
    }, headers=auth_header(admin))
    assert created.status_code == 200, created.get_json()
    task = created.get_json()['task']
    task_id = task['id']
    items = client.get(f'/api/inventory/tasks/{task_id}/items', headers=auth_header(admin)).get_json()
    assert items['success']
    item = next(x for x in items['data'] if x['asset_id'] == aid)
    assert (item.get('check_status') or 'pending') == 'pending'

    checked = client.post(f'/api/inventory/items/{item["id"]}/check', json={
        'check_method': 'manual', 'checked_by': 'Cookie.gu',
    }, headers=auth_header(admin))
    assert checked.status_code == 200, checked.get_json()
    after = client.get(f'/api/inventory/tasks/{task_id}/items', headers=auth_header(admin)).get_json()
    item2 = next(x for x in after['data'] if x['id'] == item['id'])
    assert item2['check_status'] == 'normal'

    done = client.post(f'/api/inventory/tasks/{task_id}/complete', headers=auth_header(admin))
    assert done.status_code == 200, done.get_json()
    detail = client.get(f'/api/inventory/tasks/{task_id}', headers=auth_header(admin)).get_json()
    assert detail['task']['status'] == 'completed'

    ov = client.get('/api/stats/overview', headers=auth_header(admin)).get_json()
    assert ov['inventory_ongoing'] == 0


def test_idle_and_value_by_category(client):
    admin = login_admin(client)
    as_of = date.today()
    old = (as_of - timedelta(days=120)).isoformat()
    recent = as_of.isoformat()
    start = date(as_of.year - 1, as_of.month, min(as_of.day, 28))
    _insert_asset(
        asset_no='IDLE-OLD', category='笔记本', status='在库',
        purchase_date=old, purchase_price=9000, asset_type_code='LT', location='深圳')
    _insert_asset(
        asset_no='IDLE-NEW', category='笔记本', status='在库',
        purchase_date=recent, location='深圳')

    ov = client.get('/api/stats/overview', headers=auth_header(admin)).get_json()
    assert ov['idle_count'] >= 1
    idle_nos = [x['asset_no'] for x in (ov.get('idle_assets') or [])]
    assert 'IDLE-OLD' in idle_nos
    assert 'IDLE-NEW' not in idle_nos

    cats = {x['category']: x['book_value'] for x in (ov.get('value_by_category') or [])}
    assert cats.get('笔记本', 0) > 0
    assert 'monthly_in' in ov
    assert 'repair_stats' in ov


def test_approval_status_filter_accepts_chinese(client, app_mod):
    admin = login_admin(client)
    eid = _insert_employee(name='筛选员工')
    user_token = app_mod.generate_token(eid)
    client.post('/api/profile/apply/asset', json={
        'type': 'existing', 'category': '笔记本', 'reason': '筛选测试',
    }, headers=auth_header(user_token))
    zh = client.get('/api/approvals?apply_type=资产申请&status=' + '待审批',
                    headers=auth_header(admin)).get_json()
    en = client.get('/api/approvals?apply_type=资产申请&status=pending',
                    headers=auth_header(admin)).get_json()
    assert zh['success'] and en['success']
    assert zh['total'] >= 1
    assert zh['total'] == en['total']


def test_inventory_mark_abnormal_and_self_confirm(client, app_mod):
    admin = login_admin(client)
    eid = _insert_employee(name='盘点员工')
    user_token = app_mod.generate_token(eid)
    stock_id = _insert_asset(asset_no='AB-STOCK', status='在库', location='深圳')
    used_id = _insert_asset(
        asset_no='AB-USED', status='在用', assignee='盘点员工', location='深圳')

    created = client.post('/api/inventory/tasks', json={
        'task_name': '异常与自助确认',
        'scope_type': 'all',
        'created_by': 'Cookie.gu',
    }, headers=auth_header(admin))
    assert created.status_code == 200, created.get_json()
    task_id = created.get_json()['task']['id']
    items = client.get(f'/api/inventory/tasks/{task_id}/items',
                       headers=auth_header(admin)).get_json()['data']
    stock_item = next(x for x in items if x['asset_id'] == stock_id)
    used_item = next(x for x in items if x['asset_id'] == used_id)
    assert used_item.get('confirm_required') == 1

    marked = client.post(f'/api/inventory/items/{stock_item["id"]}/abnormal', json={
        'abnormal_type': 'missing', 'abnormal_desc': '仓库找不到',
    }, headers=auth_header(admin))
    assert marked.status_code == 200, marked.get_json()
    assert marked.get_json()['check_status'] == 'missing'

    pending = client.get('/api/profile/inventory-pending',
                         headers=auth_header(user_token)).get_json()
    assert pending['success']
    assert any(x['id'] == used_item['id'] for x in pending['data'])

    confirmed = client.post(f'/api/inventory/items/{used_item["id"]}/confirm', json={
        'confirmed_by': '盘点员工',
    }, headers=auth_header(user_token))
    assert confirmed.status_code == 200, confirmed.get_json()
    after = client.get(f'/api/inventory/tasks/{task_id}/items',
                       headers=auth_header(admin)).get_json()['data']
    used_after = next(x for x in after if x['id'] == used_item['id'])
    assert used_after['check_status'] == 'normal'


def test_inventory_complete_marks_pending_abnormal(client):
    admin = login_admin(client)
    a1 = _insert_asset(asset_no='CMP-1', status='在库', location='深圳')
    a2 = _insert_asset(asset_no='CMP-2', status='在库', location='深圳')
    created = client.post('/api/inventory/tasks', json={
        'task_name': '完成时标未盘',
        'scope_type': 'in_stock',
        'created_by': 'Cookie.gu',
    }, headers=auth_header(admin))
    assert created.status_code == 200, created.get_json()
    task_id = created.get_json()['task']['id']
    items = client.get(f'/api/inventory/tasks/{task_id}/items',
                       headers=auth_header(admin)).get_json()['data']
    item1 = next(x for x in items if x['asset_id'] == a1)
    item2 = next(x for x in items if x['asset_id'] == a2)
    client.post(f'/api/inventory/items/{item1["id"]}/check', json={
        'check_method': 'manual', 'checked_by': 'Cookie.gu',
    }, headers=auth_header(admin))

    done = client.post(f'/api/inventory/tasks/{task_id}/complete',
                       headers=auth_header(admin))
    assert done.status_code == 200, done.get_json()
    after = client.get(f'/api/inventory/tasks/{task_id}/items',
                       headers=auth_header(admin)).get_json()['data']
    i1 = next(x for x in after if x['id'] == item1['id'])
    i2 = next(x for x in after if x['id'] == item2['id'])
    assert i1['check_status'] == 'normal'
    assert i2['check_status'] == 'abnormal'
    assert '未盘点' in (i2.get('abnormal_desc') or '')
    detail = client.get(f'/api/inventory/tasks/{task_id}',
                        headers=auth_header(admin)).get_json()
    assert detail['task']['status'] == 'completed'


def test_stats_overview_pending_po(client):
    admin = login_admin(client)
    open1 = client.post('/api/orders', json={
        'title': '未到货无票', 'total_amount': 3000, 'status': '待到货',
        'items': [{'item_name': '笔记本', 'quantity': 1, 'unit_price': 3000}],
    }, headers=auth_header(admin))
    assert open1.status_code == 200, open1.get_json()
    open2 = client.post('/api/orders', json={
        'title': '部分到货有票', 'total_amount': 1500, 'status': '部分到货',
        'invoice_no': 'INV-1',
        'items': [{'item_name': '显示器', 'quantity': 1, 'unit_price': 1500}],
    }, headers=auth_header(admin))
    assert open2.status_code == 200, open2.get_json()
    done = client.post('/api/orders', json={
        'title': '已到货无票', 'total_amount': 800, 'status': '已到货',
        'items': [{'item_name': '键盘', 'quantity': 1, 'unit_price': 800}],
    }, headers=auth_header(admin))
    assert done.status_code == 200, done.get_json()
    cancelled = client.post('/api/orders', json={
        'title': '已取消', 'total_amount': 9999, 'status': '已取消',
        'items': [{'item_name': '取消品', 'quantity': 1, 'unit_price': 9999}],
    }, headers=auth_header(admin))
    assert cancelled.status_code == 200, cancelled.get_json()

    ov = client.get('/api/stats/overview', headers=auth_header(admin)).get_json()
    assert ov['pending_po_count'] >= 2
    assert ov['pending_po_amount'] >= 4500
    # 未到票：待到货无票 + 已到货无票（取消不计）
    assert ov['uninvoiced_po_count'] >= 2


def test_create_order_from_fulfillable_application(client, app_mod):
    admin = login_admin(client)
    eid = _insert_employee(name='采购申请人')
    user_token = app_mod.generate_token(eid)
    created = client.post('/api/profile/apply/asset', json={
        'type': 'existing', 'category': '笔记本', 'reason': '一键建单测试',
    }, headers=auth_header(user_token))
    assert created.status_code == 200, created.get_json()
    app_id = created.get_json()['id']
    appr = client.post(f'/api/approvals/{app_id}/approve', json={
        'apply_type': '资产申请', 'status': '已通过',
    }, headers=auth_header(admin))
    assert appr.status_code == 200, appr.get_json()

    before = client.get('/api/stats/overview', headers=auth_header(admin)).get_json()
    made = client.post(f'/api/approvals/{app_id}/create-order', json={},
                       headers=auth_header(admin))
    assert made.status_code == 200, made.get_json()
    body = made.get_json()
    assert body['success'] and body.get('order_no')

    pending = client.get('/api/approvals/fulfillable', headers=auth_header(admin)).get_json()
    row = next(x for x in pending['data'] if x['id'] == app_id)
    assert row.get('po_order_no') == body['order_no']

    after = client.get('/api/stats/overview', headers=auth_header(admin)).get_json()
    assert after['pending_po_count'] == before['pending_po_count'] + 1
    assert after['uninvoiced_po_count'] == before['uninvoiced_po_count'] + 1

    again = client.post(f'/api/approvals/{app_id}/create-order', json={},
                        headers=auth_header(admin))
    assert again.status_code == 400
    assert again.get_json().get('order_no') == body['order_no']
    after2 = client.get('/api/stats/overview', headers=auth_header(admin)).get_json()
    assert after2['pending_po_count'] == after['pending_po_count']


def test_dispose_list_requires_auth_and_shows_pending(client):
    assert client.get('/api/dispose').status_code == 401
    admin = login_admin(client)
    empty = client.get('/api/dispose', headers=auth_header(admin))
    assert empty.status_code == 200
    assert empty.get_json()['data'] == []

    aid = _insert_asset(asset_no='DSP-001', status='在用')
    before = client.get('/api/stats/overview', headers=auth_header(admin)).get_json()
    posted = client.post(
        '/api/dispose',
        json={'asset_id': aid, 'dispose_type': '报废', 'reason': '到期报废',
              'dispose_date': '2026-10-06'},
        headers=auth_header(admin),
    )
    assert posted.status_code == 200, posted.get_json()
    assert posted.get_json().get('success') is True
    after = client.get('/api/stats/overview', headers=auth_header(admin)).get_json()
    assert after['pending_approvals'] == (before.get('pending_approvals') or 0) + 1

    listed = client.get('/api/dispose', headers=auth_header(admin))
    assert listed.status_code == 200
    body = listed.get_json()
    assert body['total'] >= 1
    row = body['data'][0]
    assert row.get('pending') is True
    assert row.get('record_status') == '待审批'
    assert row.get('asset_no') == 'DSP-001'
    assert row.get('dispose_type') == '报废'


def _xlsx_file(headers, rows, filename='import.xlsx'):
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = filename
    return buf, filename


def test_asset_excel_template_and_strict_import(client, app_mod):
    admin = login_admin(client)
    tpl = client.get('/api/assets/template', headers=auth_header(admin))
    assert tpl.status_code == 200
    empty_tpl = client.post(
        '/api/assets/import',
        data={'file': (BytesIO(tpl.data), '资产导入模板.xlsx')},
        headers=auth_header(admin),
    )
    assert empty_tpl.status_code == 400, empty_tpl.get_json()
    err0 = empty_tpl.get_json().get('error') or ''
    assert '打不开' not in err0
    assert '读不出来' not in err0
    assert any(k in err0 for k in ('空', '数据', '表头', '模板'))

    filled_wb = load_workbook(BytesIO(tpl.data))
    filled_wb.active.append([
        'IMP-TPL-01', '笔记本', 'Dell', 'XPS', '', '', '', '', '在库',
        '深圳', '', '', '', '2026-01-01', '', 1000, '',
    ])
    filled_buf = BytesIO()
    filled_wb.save(filled_buf)
    filled_buf.seek(0)
    filled = client.post(
        '/api/assets/import',
        data={'file': (filled_buf, '资产导入模板.xlsx')},
        headers=auth_header(admin),
    )
    assert filled.status_code == 200, filled.get_json()
    assert filled.get_json().get('imported') == 1

    wb = load_workbook(BytesIO(tpl.data))
    assert [c.value for c in wb.active[1]] == list(app_mod.ASSET_IMPORT_HEADERS)
    assert wb.active.max_row == 1

    headers = app_mod.ASSET_IMPORT_HEADERS
    empty_buf, empty_name = _xlsx_file(headers, [])
    empty = client.post(
        '/api/assets/import',
        data={'file': (empty_buf, empty_name)},
        headers=auth_header(admin),
    )
    assert empty.status_code == 400

    bad_buf, bad_name = _xlsx_file(['错列'], [['x']])
    bad = client.post(
        '/api/assets/import',
        data={'file': (bad_buf, bad_name)},
        headers=auth_header(admin),
    )
    assert bad.status_code == 400
    assert '模板' in (bad.get_json().get('error') or '')

    row = ['IMP-001', '笔记本', 'Dell', 'XPS', '', '', '', '', '在库', '深圳', '', '', '',
           '2026-01-01', '', 12000, '']
    ok_buf, ok_name = _xlsx_file(headers, [row])
    ok = client.post(
        '/api/assets/import',
        data={'file': (ok_buf, ok_name)},
        headers=auth_header(admin),
    )
    body = ok.get_json()
    assert ok.status_code == 200, body
    assert body.get('success') is True
    assert body.get('imported') == 1
    listed = client.get('/api/assets?q=IMP-001', headers=auth_header(admin)).get_json()
    assert any(x.get('asset_no') == 'IMP-001' for x in listed.get('data') or [])

    dup_buf, dup_name = _xlsx_file(headers, [row])
    dup = client.post(
        '/api/assets/import',
        data={'file': (dup_buf, dup_name)},
        headers=auth_header(admin),
    )
    assert dup.status_code == 400


def test_consumable_excel_template_and_import(client, app_mod):
    admin = login_admin(client)
    tpl = client.get('/api/consumables/template', headers=auth_header(admin))
    assert tpl.status_code == 200
    wb = load_workbook(BytesIO(tpl.data))
    assert [c.value for c in wb.active[1]] == list(app_mod.CONSUMABLE_IMPORT_HEADERS)
    assert wb.active.max_row == 1

    headers = app_mod.CONSUMABLE_IMPORT_HEADERS
    row = ['CONS-IMP-1', '硒鼓', '硒鼓', 'HP', '26A', '个', 8, 2, '仓库', '正常', '', '']
    ok_buf, ok_name = _xlsx_file(headers, [row])
    ok = client.post(
        '/api/consumables/import',
        data={'file': (ok_buf, ok_name)},
        headers=auth_header(admin),
    )
    body = ok.get_json()
    assert ok.status_code == 200, body
    assert body.get('imported') == 1

