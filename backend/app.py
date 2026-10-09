"""
资产管理系统 - Flask 后端
"""
from flask import Flask, request, jsonify, send_from_directory, send_file
from flask_cors import CORS
import sqlite3, os, re, json, hashlib, hmac, time, base64, io, csv, secrets
from datetime import datetime, timedelta, date
import urllib.request, urllib.error
from werkzeug.utils import secure_filename
from functools import wraps

from database import get_conn, init_db, close_conn_on_teardown, hash_password, verify_password

BASE_DIR  = os.path.dirname(__file__)
FRONT_DIR = os.path.join(BASE_DIR, 'frontend') if os.path.exists(os.path.join(BASE_DIR, 'frontend')) else os.path.join(BASE_DIR, '..', 'frontend')
DATA_DIR  = os.environ.get('DATA_DIR') or os.path.join(BASE_DIR, '..', 'data')
UPLOAD_DIR = os.path.join(DATA_DIR, 'uploads')
try:
    os.makedirs(UPLOAD_DIR, exist_ok=True)
except OSError as e:
    # 目录不可写（如容器以非 root 运行且宿主挂载目录无写权限）不应导致整个后端启动失败，
    # 仅告警，上传功能在真正需要时再报错。
    print(f"[WARN] 无法创建上传目录 {UPLOAD_DIR}: {e}", flush=True)

app = Flask(__name__, static_folder=FRONT_DIR, static_url_path='')
CORS(app)
# 安全加固（P2-5）：上传请求体上限 16MB，超限直接 413 不落盘，防超大文件占满磁盘（DoS）。
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024
# 请求结束时真正关闭数据库连接（get_conn 返回的共享连接已被改为 .close() 不真正关闭）
app.teardown_appcontext(close_conn_on_teardown)

# 强制 HTML 文件不缓存（避免浏览器加载旧版 index.html）
@app.after_request
def disable_html_cache(resp):
    ct = resp.headers.get('Content-Type', '')
    if 'text/html' in ct:
        resp.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        resp.headers['Pragma'] = 'no-cache'
        resp.headers['Expires'] = '0'
    # 静态 JS/CSS 也强制 no-store，防止旧版本被 ETag 验证通过
    elif any(t in ct for t in ('javascript', 'css', 'image')):
        resp.headers['Cache-Control'] = 'no-store, must-revalidate'
    # 移除 ETag 和 Last-Modified，防止 304 缓存命中
    resp.headers.pop('ETag', None)
    resp.headers.pop('Last-Modified', None)
    # 安全响应头（纵深防御 P2-7）：
    resp.headers['X-Content-Type-Options'] = 'nosniff'          # 防 MIME 嗅探
    resp.headers['X-Frame-Options'] = 'SAMEORIGIN'              # 防点击劫持（允许同源 iframe 子页）
    resp.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'  # 收敛 Referer 泄露
    # CSP 观察模式（Report-Only，不拦截仅上报）：既能观察 XSS 尝试，又避免误伤
    # Lark 集成与大量内联脚本/eval（Vue compiler）；稳定运行后可升级为强制 CSP。
    resp.headers['Content-Security-Policy-Report-Only'] = (
        "default-src 'self'; script-src 'self' 'unsafe-inline' 'unsafe-eval'; "
        "style-src 'self' 'unsafe-inline'; img-src 'self' data: blob:; "
        "connect-src 'self' https://open.larksuite.com https://accounts.larksuite.com; "
        "frame-ancestors 'self'")
    return resp

# favicon：返回 204，避免浏览器自动请求产生 404 噪音
@app.route('/favicon.ico')
def favicon():
    return '', 204

# 启动时自动建表（gunicorn 不会走 __main__ 分支，必须在这里调用）
try:
    init_db()
    print(f"[OK] 数据库已就绪: {DATA_DIR}/assets.db", flush=True)
except Exception as e:
    print(f"[ERROR] 数据库初始化失败: {e}", flush=True)

# 自定义静态文件路由，排除 /api/ 路径
# 原因：Flask 默认的 /<path:filename> 会匹配带查询字符串的 API 请求
@app.route('/<path:filename>', methods=['GET', 'HEAD'])
def serve_static(filename):
    """静态文件服务（排除 /api/ 路径）"""
    from werkzeug.exceptions import NotFound
    # 移除查询字符串
    pure_path = filename.split('?')[0] if '?' in filename else filename
    # 检查是否是 API 请求
    if pure_path.startswith('api/') or pure_path == 'api':
        raise NotFound()
    resp = send_from_directory(FRONT_DIR, pure_path)
    # HTML 文件强制不缓存（iframe 子页与入口页均生效）
    if pure_path.endswith('.html'):
        resp.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        resp.headers['Pragma'] = 'no-cache'
        resp.headers['Expires'] = '0'
    return resp

# 同时保留 /uploads/ 路由用于上传文件
@app.route('/api/uploads/<path:filename>', methods=['GET', 'HEAD'])
def serve_upload_file(filename):
    """上传文件服务"""
    return send_from_directory(UPLOAD_DIR, filename)

# 禁用静态文件缓存，方便开发调试
@app.after_request
def add_no_cache(response):
    if 'static' in request.path or request.path == '/':
        response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
    return response

# ─────────────────────────────────────────────
# 工具函数
# ─────────────────────────────────────────────

def rows_to_list(rows):
    return [dict(r) for r in rows]


# 资产敏感字段：关联账号的明文密码绝不返回给前端（P0-4）。
# account_name（账号名）保留用于业务展示；account_password / account_pwd 含明文密码，必须剔除。
ASSET_SENSITIVE_FIELDS = ('account_password', 'account_pwd')

def strip_sensitive_asset(d):
    """剔除资产记录中的敏感字段（明文密码），原地修改并返回，安全用于返回前的脱敏。"""
    for f in ASSET_SENSITIVE_FIELDS:
        d.pop(f, None)
    return d

def assets_to_safe_list(rows):
    """资产行列表 → 剔除敏感字段后的 dict 列表（返回前端统一入口）。"""
    return [strip_sensitive_asset(dict(r)) for r in rows]


def _months_between(start_date, end_date):
    """日历月差（未满一个月不计）。start/end 为 date。"""
    months = (end_date.year - start_date.year) * 12 + (end_date.month - start_date.month)
    if end_date.day < start_date.day:
        months -= 1
    return max(0, months)


def compute_book_value(purchase_price, purchase_date, depreciation_months, residual_rate, as_of=None):
    """直线法账面价值。缺采购价、入库日或折旧月数时返回 None。
    账面 = max(采购价 × 残值率, 采购价 − 采购价 × (1 − 残值率) × 已用月数 / 折旧月数)
    """
    try:
        price = float(purchase_price)
    except (TypeError, ValueError):
        return None
    if price < 0 or not purchase_date:
        return None
    try:
        dep_months = int(depreciation_months)
    except (TypeError, ValueError):
        return None
    if dep_months <= 0:
        return None
    try:
        rate = float(residual_rate) if residual_rate is not None else 0.0
    except (TypeError, ValueError):
        rate = 0.0
    rate = min(max(rate, 0.0), 1.0)
    try:
        start = datetime.strptime(str(purchase_date)[:10], '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return None
    as_of = as_of or datetime.now().date()
    used = _months_between(start, as_of)
    floor = price * rate
    book = price - price * (1.0 - rate) * used / dep_months
    return round(max(floor, book), 2)


def load_type_params_map(conn):
    rows = conn.execute(
        "SELECT type_code, type_name, depreciation_months, residual_rate FROM asset_type_params"
    ).fetchall()
    return {r['type_code']: r for r in rows}


def resolve_asset_type_code(conn, category, asset_type_code=''):
    """若未填类型码，按类目名称对照 asset_type_params.type_name 回填。"""
    code = (asset_type_code or '').strip()
    if code:
        return code
    cat = (category or '').strip()
    if not cat:
        return ''
    row = conn.execute(
        "SELECT type_code FROM asset_type_params WHERE type_name=? OR type_code=? LIMIT 1",
        (cat, cat)).fetchone()
    return row['type_code'] if row else ''


def category_filter_clause(conn, category_param):
    """类目筛选：同时兼容 field_code（如 LT）与中文类目名（如「笔记本」）。

    台账里 category 多为中文名，asset_type_code 为编码；快速操作按类型浏览传的是编码。
    返回 (sql片段, params)，无筛选时返回 ('', [])。
    """
    cat = (category_param or '').strip()
    if not cat:
        return '', []
    row = conn.execute(
        'SELECT field_code, field_name FROM field_options '
        'WHERE field_type="category" AND (field_code=? OR field_name=?) LIMIT 1',
        (cat, cat)
    ).fetchone()
    if row:
        code, name = row['field_code'], row['field_name']
        return (
            '(asset_type_code=? OR category=? OR category=?)',
            [code, name, code]
        )
    # 未在 field_options 中：仍兼容直接按 category / asset_type_code 精确匹配
    return '(category=? OR asset_type_code=?)', [cat, cat]


def attach_book_value(d, params_map):
    code = d.get('asset_type_code')
    p = params_map.get(code) if code else None
    if not p:
        d['book_value'] = None
        d['depreciation_months'] = None
        return d
    d['depreciation_months'] = p['depreciation_months']
    d['book_value'] = compute_book_value(
        d.get('purchase_price'), d.get('purchase_date'),
        p['depreciation_months'], p['residual_rate'])
    return d


def enrich_assets_book_value(dicts, conn):
    pmap = load_type_params_map(conn)
    for d in dicts:
        attach_book_value(d, pmap)
    return dicts


class AssetAppExecError(Exception):
    """资产申请审批执行失败（库存不足、状态不符等），单据应保持 pending。"""
    pass


_APP_TYPE_LABELS = {
    'new_asset': '新资产申请',
    'borrow': '借用申请',
    'return': '归还申请',
    'repair': '报修申请',
    'consumable': '耗材领用',
    'replace': '更换设备申请',
    'special': '特殊需求申请',
}


def insert_asset_application(conn, user, app_type, *, asset_id=None, asset_no='', title='', reason='', detail=None):
    """写入 asset_applications（pending）。调用方负责 commit。返回 (id, app_no)。"""
    from datetime import date
    today_str = date.today().strftime('%Y%m%d')
    count = conn.execute(
        "SELECT COUNT(*) as cnt FROM asset_applications WHERE app_no LIKE ?",
        (f'APP-{today_str}-%',)
    ).fetchone()['cnt']
    app_no = f"APP-{today_str}-{str(count + 1).zfill(4)}"
    if asset_id == '' or asset_id is None:
        asset_id = None
    else:
        try:
            asset_id = int(asset_id)
        except (TypeError, ValueError):
            asset_id = None
    if not title:
        title = f"{_APP_TYPE_LABELS.get(app_type, '申请')} - {user['name']}"
    if detail is None:
        detail = {}
    cursor = conn.execute('''
        INSERT INTO asset_applications
        (app_no, app_type, applicant_id, applicant_name, asset_id, asset_no,
         title, reason, detail, status, apply_time)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'pending', datetime('now','localtime'))
    ''', (
        app_no, app_type, user['id'], user['name'], asset_id, asset_no or '',
        title, reason or '', json.dumps(detail, ensure_ascii=False)
    ))
    return cursor.lastrowid, app_no


def execute_asset_application(conn, appr):
    """审批通过后按 app_type 执行。不 commit。失败抛 AssetAppExecError。"""
    app_type = appr['app_type']
    applicant_name = appr['applicant_name']
    try:
        detail = json.loads(appr['detail'] or '{}') if appr['detail'] else {}
    except (TypeError, json.JSONDecodeError):
        detail = {}
    if not isinstance(detail, dict):
        detail = {}

    if app_type in ('new_asset', 'replace', 'special'):
        return

    aid = appr['asset_id'] if appr['asset_id'] else detail.get('asset_id')
    if aid in ('', None):
        aid = None
    else:
        try:
            aid = int(aid)
        except (TypeError, ValueError):
            aid = None

    if app_type == 'borrow':
        if not aid:
            raise AssetAppExecError('借用申请缺少资产')
        asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
        if not asset:
            raise AssetAppExecError('资产不存在')
        if asset['status'] != '在库':
            raise AssetAppExecError(f'资产当前状态为【{asset["status"]}】，无法借用')
        conn.execute(
            "UPDATE assets SET status='在用', assignee=?, updated_at=datetime('now','localtime') WHERE id=?",
            (applicant_name, aid))
        conn.execute('''INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,target_user,from_location,to_location,remark)
                        VALUES(?,?,?,?,?,?,?,?)''',
                     (aid, asset['asset_no'], '借用出库', applicant_name, applicant_name,
                      asset['location'], asset['location'], appr.get('reason') or ''))
        return

    if app_type == 'return':
        if not aid:
            raise AssetAppExecError('归还申请缺少资产')
        asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
        if not asset:
            raise AssetAppExecError('资产不存在')
        if (asset['assignee'] or '') != applicant_name:
            raise AssetAppExecError('只能归还申请人名下的资产')
        if asset['status'] != '在用':
            raise AssetAppExecError(f'资产当前状态为【{asset["status"]}】，无法归还')
        conn.execute(
            "UPDATE assets SET status='在库', assignee='', updated_at=datetime('now','localtime') WHERE id=?",
            (aid,))
        conn.execute('''INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,target_user,from_location,to_location,remark)
                        VALUES(?,?,?,?,?,?,?,?)''',
                     (aid, asset['asset_no'], '归还', applicant_name, applicant_name,
                      asset['location'], asset['location'], appr.get('reason') or ''))
        return

    if app_type == 'repair':
        if not aid:
            raise AssetAppExecError('报修申请缺少资产')
        asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
        if not asset:
            raise AssetAppExecError('资产不存在')
        if (asset['assignee'] or '') != applicant_name:
            raise AssetAppExecError('只能报修申请人名下的资产')
        fault = detail.get('fault_desc') or appr.get('reason') or ''
        conn.execute(
            "UPDATE assets SET status='维修中', updated_at=datetime('now','localtime') WHERE id=?",
            (aid,))
        conn.execute('''INSERT INTO repair_logs(asset_id, asset_no, fault_desc, submit_by, repair_status)
                        VALUES(?,?,?,?, '待处理')''',
                     (aid, asset['asset_no'], fault, applicant_name))
        return

    if app_type == 'consumable':
        item_id = detail.get('item_id')
        qty = detail.get('quantity', detail.get('qty', 0))
        try:
            item_id = int(item_id) if item_id not in (None, '') else None
            qty = int(qty)
        except (TypeError, ValueError):
            raise AssetAppExecError('耗材申请数量或耗材无效')
        if not item_id or qty < 1:
            raise AssetAppExecError('耗材申请缺少耗材或数量')
        item = conn.execute("SELECT * FROM consumables WHERE id=?", (item_id,)).fetchone()
        if not item:
            raise AssetAppExecError('耗材不存在')
        if item['quantity'] < qty:
            raise AssetAppExecError(f'库存不足，当前库存: {item["quantity"]}')
        new_qty = item['quantity'] - qty
        new_status = '正常' if new_qty > item['min_stock'] else ('低库存' if new_qty > 0 else '缺货')
        conn.execute(
            "UPDATE consumables SET quantity=?, status=?, updated_at=datetime('now','localtime') WHERE id=?",
            (new_qty, new_status, item_id))
        conn.execute('''
            INSERT INTO consumable_logs(item_id, item_no, log_type, quantity, operator, target_user, location, remark)
            VALUES(?, ?, '出库', ?, ?, ?, ?, ?)
        ''', (item_id, item['item_no'], qty, applicant_name, applicant_name, item['location'],
              appr.get('reason') or detail.get('purpose') or ''))
        return

    raise AssetAppExecError('无效的申请类型')


def _rand_suffix():
    """生成不可枚举的上传文件随机段（安全加固 P2-4）：
    12 字节安全随机（约 2^96 空间，无法穷举），替代原可推算的时间戳，
    防止凭证/发票/资产照片等附件 URL 被枚举猜测后未授权下载。"""
    return secrets.token_urlsafe(12).replace('-', '').replace('_', '')



def fuzzy_match_clause(keyword, fields):
    """生成模糊匹配 SQL 片段，支持以下场景：
      - 整体子串匹配（连续子串）
      - 拆字 AND 匹配：输入"王存玮" → name 同时含 王/存/玮 才能命中
      - 大小写不敏感（输入含 ASCII 字母时自动启用 LOWER）
    fields: 字段名列表
    返回: (sql_clause, params_list)
    """
    if not keyword or not fields:
        return '', []
    kw = keyword.strip()
    if not kw:
        return '', []
    # 检测是否需要大小写不敏感
    has_ascii_alpha = any(c.isascii() and c.isalpha() for c in kw)

    field_groups = []
    field_params = []
    for field in fields:
        clauses = []
        if has_ascii_alpha:
            clauses.append(f"LOWER({field}) LIKE LOWER(?) ESCAPE '\\'")
        else:
            clauses.append(f"{field} LIKE ? ESCAPE '\\'")
        field_params.append(f'%{kw}%')
        # 拆字匹配
        for c in kw:
            if c.isspace():
                continue
            if has_ascii_alpha and c.isascii():
                clauses.append(f"LOWER({field}) LIKE LOWER(?) ESCAPE '\\'")
                field_params.append(f'%{c.lower()}%')
            else:
                clauses.append(f"{field} LIKE ? ESCAPE '\\'")
                field_params.append(f'%{c}%')
        field_groups.append('(' + ' AND '.join(clauses) + ')')
    return '(' + ' OR '.join(field_groups) + ')', field_params


def gen_asset_no(subject_code: str = '', region_code: str = '', asset_type_code: str = '') -> str:
    """生成资产编号: YYYYMMDD + 三位数序号 + 主体 + 地区 + 类型
    格式: 20260410001-NHWHK-SZ-LT
    序号按(日期+主体+地区+类型)组合唯一
    """
    conn = get_conn()
    today = datetime.now().strftime('%Y%m%d')
    
    # 序号: 当天该组合下的第N个
    like = f"{today}%{subject_code}%{region_code}%{asset_type_code}%"
    row = conn.execute(
        "SELECT COUNT(*) as cnt FROM assets WHERE asset_no LIKE ?", (like,)
    ).fetchone()
    seq = (row['cnt'] or 0) + 1
    conn.close()
    
    return f"{today}{seq:03d}-{subject_code}-{region_code}-{asset_type_code}"


def gen_consumable_no(category: str = '') -> str:
    """生成耗材编号: HC-YYYYMMDD-序号"""
    conn = get_conn()
    today = datetime.now().strftime('%Y%m%d')
    like = f'HC-{today}-%'
    row = conn.execute(
        "SELECT COUNT(*) as cnt FROM consumables WHERE item_no LIKE ?", (like,)
    ).fetchone()
    seq = (row['cnt'] or 0) + 1
    conn.close()
    return f"HC-{today}-{seq:03d}"


def _is_valid_lark_webhook(url: str) -> bool:
    """校验 webhook 地址是否为飞书官方域名（SSRF 防护 P2-1）：
    仅允许 open.feishu.cn / open.larksuite.com，拒绝内网/回环/任意 URL。"""
    try:
        from urllib.parse import urlparse
        u = urlparse((url or '').strip())
        if u.scheme != 'https':
            return False
        return u.netloc.lower() in ('open.feishu.cn', 'open.larksuite.com')
    except Exception:
        return False


def lark_send(webhook_url: str, secret: str | None, msg: str) -> bool:
    """发送飞书 Webhook 消息（支持签名）"""
    payload = {
        "msg_type": "text",
        "content": {"text": msg}
    }
    if secret:
        ts = str(int(time.time()))
        string_to_sign = f'{ts}\n{secret}'
        hmac_code = hmac.new(string_to_sign.encode('utf-8'), digestmod=hashlib.sha256).digest()
        sign = base64.b64encode(hmac_code).decode('utf-8')
        payload['timestamp'] = ts
        payload['sign'] = sign

    data = json.dumps(payload).encode('utf-8')
    try:
        req = urllib.request.Request(
            webhook_url,
            data=data,
            headers={'Content-Type': 'application/json; charset=utf-8'}
        )
        resp = urllib.request.urlopen(req, timeout=10)
        result = json.loads(resp.read().decode())
        return result.get('code', -1) == 0 or result.get('StatusCode', -1) == 0
    except Exception as e:
        print(f"[Lark] 发送失败: {e}")
        return False


def notify_lark(event_type: str, msg: str):
    """根据配置推送飞书通知"""
    conn = get_conn()
    # event_type: stock, repair, consumable_low_stock, expire
    if event_type in ('consumable_low_stock',):
        configs = conn.execute(
            "SELECT * FROM lark_config WHERE enabled=1 AND notify_stock=1"
        ).fetchall()
    elif event_type == 'expire':
        configs = conn.execute(
            "SELECT * FROM lark_config WHERE enabled=1 AND notify_expire=1"
        ).fetchall()
    else:
        field = 'notify_stock' if event_type == 'stock' else 'notify_repair'
        configs = conn.execute(
            f"SELECT * FROM lark_config WHERE enabled=1 AND {field}=1"
        ).fetchall()
    conn.close()
    for cfg in configs:
        try:
            # 安全加固（复审2-①）：sqlite3.Row 无 .get()，改 cfg['secret']；
            # 通知发送异常隔离——通知是辅助功能，失败不应导致审批/出入库等业务回滚失败。
            lark_send(cfg['webhook_url'], cfg['secret'], msg)
        except Exception as _e:
            print(f"[Lark] 通知发送异常(已隔离不影响业务): {_e}")


def notify_lark_all(msg: str):
    """向所有已启用的飞书配置推送消息（用于管理员通知，不限事件类型）"""
    conn = get_conn()
    configs = conn.execute(
        "SELECT * FROM lark_config WHERE enabled=1"
    ).fetchall()
    conn.close()
    for cfg in configs:
        try:
            # 安全加固（复审2-①）：sqlite3.Row 无 .get()，改 cfg['secret']；
            # 通知发送异常隔离——通知是辅助功能，失败不应导致审批/出入库等业务回滚失败。
            lark_send(cfg['webhook_url'], cfg['secret'], msg)
        except Exception as _e:
            print(f"[Lark] 通知发送异常(已隔离不影响业务): {_e}")


def push_expire_reminder():
    """每日定时：推送即将到期/已过期的订阅制资产提醒"""
    from datetime import date
    today = date.today().isoformat()
    conn = get_conn()
    rows = conn.execute("""
        SELECT asset_no, category, brand, model, assignee,
               expire_date, seat_count, renew_remind_days
        FROM assets
        WHERE license_type = 'subscription'
          AND expire_date IS NOT NULL AND expire_date != ''
          AND status != '已报废'
    """).fetchall()
    conn.close()

    expired_list = []
    warning_list = []
    for r in rows:
        r = dict(r)
        remind_days = r.get('renew_remind_days') or 30
        try:
            delta = (date.fromisoformat(r['expire_date']) - date.fromisoformat(today)).days
        except Exception:
            continue
        if delta < 0:
            expired_list.append((r, delta))
        elif delta <= remind_days:
            warning_list.append((r, delta))

    if not expired_list and not warning_list:
        return  # 没有需要提醒的，不推送

    lines = [f"🔔 【订阅续期提醒】{today}"]
    if expired_list:
        lines.append(f"\n❌ 已过期（{len(expired_list)} 项）：")
        for r, delta in expired_list:
            label = f"{r['brand']} {r['model']}".strip() or r['category']
            lines.append(f"  · {r['asset_no']} {label}  到期日 {r['expire_date']}（已超 {abs(delta)} 天）")
    if warning_list:
        lines.append(f"\n⚠️ 即将到期（{len(warning_list)} 项）：")
        for r, delta in warning_list:
            label = f"{r['brand']} {r['model']}".strip() or r['category']
            lines.append(f"  · {r['asset_no']} {label}  到期日 {r['expire_date']}（还剩 {delta} 天）")

    notify_lark('expire', '\n'.join(lines))


@app.route('/api/expire-reminder/push', methods=['POST'])
def manual_push_expire():
    """手动触发续期提醒推送（测试用）"""
    try:
        push_expire_reminder()
        return jsonify({'ok': True, 'msg': '推送完成'})
    except Exception as e:
        return jsonify({'ok': False, 'msg': str(e)}), 500



# ─────────────────────────────────────────────
# 静态页面
# ─────────────────────────────────────────────

@app.route('/')
def index():
    return send_from_directory(FRONT_DIR, 'main.html')

@app.route('/index.html')
def index_html_alias():
    # 旧版 index.html 重定向到新版 main.html，强制浏览器加载新文件
    return send_from_directory(FRONT_DIR, 'main.html')


# 其他旧版整页副本统一重定向到最新 main.html（避免暴露过期逻辑；
# 注意：stock.html/repair.html/dispose.html/fields.html 等是 iframe 子页，需正常提供）
_LEGACY_ENTRY_HTML = {'app.html', 'index-old.html'}
for _legacy in _LEGACY_ENTRY_HTML:
    def _legacy_redirect():
        return send_from_directory(FRONT_DIR, 'main.html')
    _legacy_redirect.__name__ = f'legacy_redirect_{_legacy.replace(".", "_")}'
    app.add_url_rule(f'/{_legacy}', view_func=_legacy_redirect)


# ═══════════════════════════════════════════════════════════════
# 认证 & Lark 登录集成
# ═══════════════════════════════════════════════════════════════

# 简单的 token 存储（生产环境应使用 Redis 或数据库存储）
# Token 已持久化到数据库（active_tokens 表），gunicorn worker 回收不会丢失

def generate_token(user_id):
    """生成登录 token（持久化到数据库，gunicorn worker 回收不会丢失）"""
    import secrets
    token = secrets.token_urlsafe(32)
    conn = get_conn()
    conn.execute(
        "INSERT OR REPLACE INTO active_tokens(token, user_id, created_at) VALUES(?,?,?)",
        (token, user_id, datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
    )
    conn.commit()
    conn.close()
    return token

def verify_token(token):
    """验证 token 是否有效"""
    if not token:
        return None
    conn = get_conn()
    row = conn.execute(
        "SELECT user_id, created_at FROM active_tokens WHERE token=?",
        (token,)
    ).fetchone()
    if not row:
        conn.close()
        return None
    created = datetime.strptime(row['created_at'], '%Y-%m-%d %H:%M:%S')
    # token 24小时过期
    if (datetime.now() - created).total_seconds() > 86400:
        conn.execute("DELETE FROM active_tokens WHERE token=?", (token,))
        conn.commit()
        conn.close()
        return None
    conn.close()
    return row['user_id']

def get_current_user_id():
    """从请求中获取当前用户ID（仅从 Authorization header 获取，不使用 cookie 避免循环跳转）"""
    token = request.headers.get('Authorization', '').replace('Bearer ', '')
    return verify_token(token)


def get_current_user():
    """获取当前登录用户的完整信息"""
    user_id = get_current_user_id()
    if not user_id:
        return None
    conn = get_conn()
    user = conn.execute("SELECT * FROM employees WHERE id = ?", (user_id,)).fetchone()
    conn.close()
    return dict(user) if user else None


def log_action(action_type, module, target_type=None, target_id=None, target_no=None, target_name=None,
               detail=None, before_data=None, after_data=None):
    """记录操作日志
    action_type: create/update/delete/approve/reject/arrive/cancel/upload/apply
    module: asset/consumable/order/repair/dispose/inventory/approval
    """
    # 从 login_required 设置的上下文取用户，避免调用 get_current_user()(视图函数) 导致的失效
    uid = getattr(request, 'current_user_id', None) or get_current_user_id()
    if not uid:
        return

    conn = get_conn()
    try:
        row = conn.execute("SELECT name, role FROM employees WHERE id=?", (uid,)).fetchone()
        if not row:
            return
        name = row['name']
        role = row['role'] or '未知'
        conn.execute('''
            INSERT INTO action_logs (action_type, module, target_type, target_id, target_no, target_name,
                detail, before_data, after_data, operator, operator_role, ip_address, user_agent)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            action_type, module, target_type, target_id, target_no, target_name,
            json.dumps(detail, ensure_ascii=False) if detail else None,
            json.dumps(before_data, ensure_ascii=False) if before_data else None,
            json.dumps(after_data, ensure_ascii=False) if after_data else None,
            name, role,
            request.remote_addr,
            request.headers.get('User-Agent', '')[:200]
        ))
        conn.commit()
    except Exception as e:
        print(f"[日志记录失败] {e}")
    # 不主动关闭共享连接，交由调用方/线程局部统一回收，避免破坏调用方事务

def get_current_name():
    """返回当前登录用户姓名（用于申请人/操作人记录），无则返回'系统'"""
    uid = getattr(request, 'current_user_id', None) or get_current_user_id()
    if not uid:
        return '系统'
    conn = get_conn()
    try:
        row = conn.execute("SELECT name FROM employees WHERE id=?", (uid,)).fetchone()
        return row['name'] if row else '系统'
    except Exception:
        return '系统'


# 操作类型 → 中文（用于审批中心展示）
OP_TYPE_ZH = {
    'stock_out': '资产出库', 'stock_return': '资产回收', 'transfer': '资产调拨',
    'asset_transfer': '资产转移', 'repair': '维修申请', 'dispose': '资产处置',
    'borrow_out': '借用出库', 'borrow_return': '借用归还'
}


def create_operation_application(op_type, asset_ids, payload, applicant=None, applicant_id=None):
    """创建资产操作审批单（出库/回收/调拨/转移/维修/处置），待超管审批后才执行实际动作"""
    conn = get_conn()
    try:
        asset_nos = []
        for aid in (asset_ids or []):
            a = conn.execute("SELECT asset_no FROM assets WHERE id=?", (aid,)).fetchone()
            if a:
                asset_nos.append(a['asset_no'])
        if not applicant:
            applicant = get_current_name()
        if applicant_id is None:
            applicant_id = getattr(request, 'current_user_id', None) or get_current_user_id()
        cur = conn.execute('''
            INSERT INTO operation_applications (op_type, asset_ids, asset_nos, payload, applicant, applicant_id, status)
            VALUES (?,?,?,?,?,?, 'pending')
        ''', (op_type, json.dumps(asset_ids or []), json.dumps(asset_nos),
              json.dumps(payload, ensure_ascii=False), applicant, applicant_id))
        app_id = cur.lastrowid
        conn.commit()
        try:
            log_action('apply', 'operation', 'operation_application', app_id,
                       '/'.join(asset_nos) if asset_nos else '', f"{OP_TYPE_ZH.get(op_type, op_type)}申请",
                       {'asset_ids': asset_ids, 'payload': payload})
        except Exception:
            pass
        return app_id
    except Exception:
        conn.rollback()
        raise


def save_operation_proof(app_id, files, uploader='系统'):
    """将报修/处置时上传的照片暂存为审批单证明（ref_type='operation_proof'），
    审批通过后由 execute_operation 迁移到 repair_log / dispose_log 对应记录"""
    if not files:
        return []
    conn = get_conn()
    ids = []
    for f in files:
        if not f or not f.filename:
            continue
        filename = secure_filename(f.filename)
        timestamp = _rand_suffix()  # 安全加固 P2-4：随机段替代可推算时间戳，防附件URL枚举
        new_filename = f"opproof_{app_id}_{timestamp}_{filename}"
        file_path = os.path.join(UPLOAD_DIR, new_filename)
        f.save(file_path)
        cur = conn.execute('''
            INSERT INTO photos(ref_type, ref_id, file_name, file_path, file_size, mime_type, uploaded_by)
            VALUES(?,?,?,?,?,?,?)
        ''', ('operation_proof', app_id, new_filename, file_path, os.path.getsize(file_path),
              f.content_type, uploader))
        ids.append(cur.lastrowid)
    conn.commit()
    return ids


def execute_operation(op):
    """审批通过后，按操作类型执行实际动作（与改造前立即执行的逻辑一致）"""
    conn = get_conn()
    op_type = op['op_type']
    try:
        payload = json.loads(op['payload']) if isinstance(op['payload'], str) else op['payload']
        asset_ids = json.loads(op['asset_ids']) if isinstance(op['asset_ids'], str) else op['asset_ids']
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        for aid in (asset_ids or []):
            asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
            if not asset:
                continue
            if op_type == 'stock_out':
                tgt = payload.get('target_user', '')
                loc = payload.get('to_location', asset['location'])
                conn.execute("UPDATE assets SET status='在用', assignee=?, location=?, updated_at=datetime('now','localtime') WHERE id=?", (tgt, loc, aid))
                conn.execute('''INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,target_user,from_location,to_location,remark)
                                VALUES(?,?,?,?,?,?,?,?)''',
                             (aid, asset['asset_no'], '出库', payload.get('operator', '系统'), tgt, asset['location'], loc, payload.get('remark', '')))
                notify_lark('stock', f"🚀 【资产出库】\n编号：{asset['asset_no']}\n类目：{asset['category']}\n品牌型号：{asset['brand']} {asset['model']}\n领用人：{tgt}\n存放地：{loc}\n经办人：{payload.get('operator', '系统')}\n时间：{now}")
            elif op_type == 'stock_return':
                loc = payload.get('to_location', asset['location'])
                conn.execute("UPDATE assets SET status='在库', assignee='', location=?, updated_at=datetime('now','localtime') WHERE id=?", (loc, aid))
                conn.execute('''INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,target_user,from_location,to_location,remark)
                                VALUES(?,?,?,?,?,?,?,?)''',
                             (aid, asset['asset_no'], '归还', payload.get('operator', '系统'), asset['assignee'], asset['location'], loc, payload.get('remark', '')))
                notify_lark('stock', f"↩️ 【资产归还】\n编号：{asset['asset_no']}\n类目：{asset['category']}\n归还人：{asset['assignee']}\n存入：{loc}\n经办人：{payload.get('operator', '系统')}\n时间：{now}")
            elif op_type == 'transfer':
                tgt = payload.get('target_user', asset['assignee'])
                loc = payload.get('to_location', asset['location'])
                conn.execute("UPDATE assets SET assignee=?, location=?, updated_at=datetime('now','localtime') WHERE id=?", (tgt, loc, aid))
                conn.execute('''INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,target_user,from_location,to_location,remark)
                                VALUES(?,?,?,?,?,?,?,?)''',
                             (aid, asset['asset_no'], '调拨', payload.get('operator', '系统'), tgt, asset['location'], loc, payload.get('remark', '')))
                notify_lark('stock', f"🔄 【资产调拨】\n编号：{asset['asset_no']}\n由 {asset['assignee'] or '库存'} → {tgt}\n地点：{asset['location']} → {loc}\n经办人：{payload.get('operator', '系统')}\n时间：{now}")
            elif op_type == 'asset_transfer':
                to_subject = payload.get('to_subject', '')
                to_region = payload.get('to_region', '')
                to_user = payload.get('target_user', asset['assignee'] or '')
                to_loc = payload.get('to_location', asset['location'])
                parts = asset['asset_no'].split('-')
                type_part = parts[-1] if len(parts) >= 5 else ''
                new_asset_no = '-'.join([parts[0], parts[1], to_subject, to_region, type_part])
                conn.execute("UPDATE assets SET asset_no=?, location=?, assignee=?, updated_at=datetime('now','localtime') WHERE id=?", (new_asset_no, to_loc, to_user, aid))
                conn.execute('''INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,target_user,from_location,to_location,remark)
                                VALUES(?,?,?,?,?,?,?,?)''',
                             (aid, new_asset_no, '资产转移', payload.get('operator', '系统'), to_user, asset['location'], to_loc, payload.get('remark', '')))
                notify_lark('stock', f"📦 【资产转移】\n旧编号：{asset['asset_no']}\n新编号：{new_asset_no}\n主体：{parts[2] if len(parts) > 2 else ''} → {to_subject}\n地点：{asset['location']} → {to_loc}\n经办人：{payload.get('operator', '系统')}\n时间：{now}")
            elif op_type == 'repair':
                cur = conn.execute('''INSERT INTO repair_logs(asset_id,asset_no,fault_desc,repair_vendor,submit_by,repair_status,repair_notes)
                                      VALUES(?,?,?,?,?,?,?)''',
                                   (aid, asset['asset_no'], payload.get('fault_desc', ''), payload.get('repair_vendor', ''),
                                    payload.get('submit_by', ''), '待处理', payload.get('repair_notes', '')))
                repair_id = cur.lastrowid
                conn.execute("UPDATE assets SET status='维修中', updated_at=datetime('now','localtime') WHERE id=?", (aid,))
                # 提交的证明/照片随审批通过挂到维修记录
                conn.execute("UPDATE photos SET ref_type='repair' WHERE ref_type='operation_proof' AND ref_id=?", (op['id'],))
                notify_lark('repair', f"🔧 【维修申请】\n编号：{asset['asset_no']}\n类目：{asset['category']} {asset['brand']} {asset['model']}\n故障描述：{payload.get('fault_desc', '')}\n提交人：{payload.get('submit_by', '')}\n时间：{now}")
            elif op_type == 'dispose':
                dispose_type = payload.get('dispose_type', '报废')
                cur = conn.execute('''INSERT INTO dispose_logs(asset_id, asset_no, dispose_type, reason, operator, approved_by, dispose_date, notes)
                                      VALUES(?,?,?,?,?,?,?,?)''',
                                   (aid, asset['asset_no'], dispose_type, payload.get('reason', ''), payload.get('operator', ''),
                                    payload.get('approved_by', ''), payload.get('dispose_date', ''), payload.get('notes', '')))
                dispose_id = cur.lastrowid
                new_status = '已报废' if dispose_type == '报废' else '已出库'
                conn.execute("UPDATE assets SET status=?, assignee='', updated_at=datetime('now','localtime') WHERE id=?", (new_status, aid))
                # 提交的处置证明随审批通过挂到处置记录
                conn.execute("UPDATE photos SET ref_type='dispose' WHERE ref_type='operation_proof' AND ref_id=?", (op['id'],))
                notify_lark('stock', f"🗑️ 【资产{dispose_type}】\n编号：{asset['asset_no']}\n类目：{asset['category']} {asset['brand']} {asset['model']}\n处理人：{payload.get('operator', '系统')}\n时间：{now}")
            elif op_type == 'borrow_out':
                tgt = payload.get('target_user', '')
                loc = payload.get('to_location', asset['location'])
                conn.execute("UPDATE assets SET status='在用', assignee=?, location=?, updated_at=datetime('now','localtime') WHERE id=?", (tgt, loc, aid))
                conn.execute('''INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,target_user,from_location,to_location,remark)
                                VALUES(?,?,?,?,?,?,?,?)''',
                             (aid, asset['asset_no'], '借用出库', payload.get('operator', '系统'), tgt, asset['location'], loc, payload.get('remark', '')))
                notify_lark('stock', f"📤 【借用出库】\n编号：{asset['asset_no']}\n类目：{asset['category']} {asset['brand']} {asset['model']}\n借用人：{tgt}\n经办人：{payload.get('operator', '系统')}\n时间：{now}")
            elif op_type == 'borrow_return':
                loc = payload.get('to_location', asset['location'])
                borrower = asset['assignee'] or ''
                conn.execute("UPDATE assets SET status='在库', assignee='', location=?, updated_at=datetime('now','localtime') WHERE id=?", (loc, aid))
                conn.execute('''INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,target_user,from_location,to_location,remark)
                                VALUES(?,?,?,?,?,?,?,?)''',
                             (aid, asset['asset_no'], '借用归还', payload.get('operator', '系统'), borrower, asset['location'], loc, payload.get('remark', '')))
                notify_lark('stock', f"📥 【借用归还】\n编号：{asset['asset_no']}\n类目：{asset['category']}\n归还人：{borrower}\n存入：{loc}\n经办人：{payload.get('operator', '系统')}\n时间：{now}")
        # 安全加固（复审P2-1）：不在此 commit——本函数仅被审批(approve_aggregated)调用，
        # 与其共享同一事务连接，由外层在置最终审批状态后统一 commit，保证原子性，
        # 避免提前提交 processing 中间态、进程崩溃后申请永久卡死无法再审批。
        return True
    except Exception as e:
        conn.rollback()
        print(f"[执行操作审批失败] {e}")
        return False


def login_required(f):
    """登录验证装饰器"""
    @wraps(f)
    def decorated(*args, **kwargs):
        user_id = get_current_user_id()
        if not user_id:
            return jsonify({'success': False, 'error': '未登录或登录已过期', 'code': 401}), 401
        # 将用户信息存入请求上下文
        request.current_user_id = user_id
        return f(*args, **kwargs)
    return decorated

def superadmin_required(f):
    """超管权限验证装饰器"""
    @wraps(f)
    def decorated(*args, **kwargs):
        user_id = get_current_user_id()
        if not user_id:
            return jsonify({'success': False, 'error': '未登录', 'code': 401}), 401
        
        conn = get_conn()
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (user_id,)).fetchone()
        conn.close()
        
        if not user or user['role'] != '超管':
            return jsonify({'success': False, 'error': '需要超管权限', 'code': 403}), 403
        
        request.current_user = dict(user)
        return f(*args, **kwargs)
    return decorated


# 可执行「通过/驳回」的审批角色（与菜单「审批中心」面向的角色对齐）
APPROVER_ROLES = ('超管', '管理员', '普通管理员')


def approver_required(f):
    """审批权限：超管 / 管理员 / 普通管理员可审批通过与驳回。"""
    @wraps(f)
    def decorated(*args, **kwargs):
        user_id = get_current_user_id()
        if not user_id:
            return jsonify({'success': False, 'error': '未登录', 'code': 401}), 401

        conn = get_conn()
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (user_id,)).fetchone()
        conn.close()

        if not user or user['role'] not in APPROVER_ROLES:
            return jsonify({
                'success': False,
                'error': '需要审批权限（超管/管理员/普通管理员）',
                'code': 403
            }), 403

        request.current_user = dict(user)
        return f(*args, **kwargs)
    return decorated


# ─────────────────────────────────────────────
# 全局接口鉴权（before_request）
# ─────────────────────────────────────────────
# 管理角色白名单 —— 与前端 frontend/ams-role.js 的 ADMIN_ROLES 完全一致。
# 判定原则：role 在白名单内 → 管理角色；否则一律视为普通用户（含历史脏值 'user'、空值等）。
ADMIN_ROLES = ('超管', '管理员', '普通管理员', '操作员', '财务')

# 免登录接口（登录前/OAuth 回调/状态探测等匿名必经路径）
PUBLIC_API_PATHS = {
    '/api/auth/login',          # 在职员工密码登录（需已设置密码）
    '/api/auth/logout',         # 退出登录（匿名调用仅删 token，无副作用）
    '/api/auth/check',          # 登录状态探测：匿名需返回 is_logged_in:false
    '/api/auth/lark-signature', # Lark 签名校验（匿名 OAuth）
    '/api/auth/lark-code-login',  # Lark code 换 token 登录
    '/api/meta',                # 元数据字典（类目/位置/编号配置），无业务数据
}
# 注意：/api/auth/lark-login 已从白名单移除并停用——它直接信任前端回传的 lark_user_id/name，
# 无签名校验，存在身份伪造风险（P0-3）。统一走 lark-code-login。

# 普通用户可用接口（需登录，管理角色同样可用）—— 前缀匹配
USER_API_PREFIXES = (
    '/api/auth/user',            # 当前用户信息
    '/api/auth/change-password', # 修改自己的密码
    '/api/my/',                  # 门户：我的资产/申请/耗材（GET/POST）
    '/api/profile/',             # 门户：我的资产/耗材/待办 + 自助申请（领用/报修）
)

# 普通用户可用接口（需登录）—— 精确 (METHOD, PATH) 匹配
USER_API_EXACT = {
    ('GET', '/api/assets/borrowable'),  # my.html 门户「可借用资产」目录
    ('GET', '/api/consumables'),        # profile.html 自助领用耗材目录（现状匿名，收为登录可见）
}


@app.before_request
def enforce_global_api_auth():
    """全局接口鉴权：所有 /api/* 接口强制鉴权，杜绝『只靠前端页面守卫、拿到 URL 直接调用』的越权。

    三层规则：
      1. /api/uploads/* 与 PUBLIC_API_PATHS → 免登录（图片 img src 无法携带 header）。
      2. 其余接口必须先携带有效登录 token，否则 401。
      3. 非普通用户可用接口（USER_API_* 之外）= 管理接口：仅 ADMIN_ROLES 角色可访问，否则 403。
    说明：各接口函数内更严格的角色校验（如超管专用）保持不变，本钩子只做统一兜底。
    """
    path = request.path
    if not path.startswith('/api/'):
        return
    if request.method == 'OPTIONS':
        return  # CORS 预检
    if path.startswith('/api/uploads/'):
        return
    if path in PUBLIC_API_PATHS:
        return

    user = get_current_user()
    if user is None:
        return jsonify({'success': False, 'error': '未登录或登录已过期', 'code': 401}), 401
    request.current_user_id = user['id']
    request.current_user = user

    # 普通用户可用接口：登录即可访问
    if path.startswith(USER_API_PREFIXES) or (request.method, path) in USER_API_EXACT:
        return
    # 责任人自助确认盘点：仅 POST .../inventory/items/<id>/confirm
    if (request.method == 'POST' and path.startswith('/api/inventory/items/')
            and path.endswith('/confirm')):
        return

    # 其余均为管理接口：要求管理角色
    if user.get('role') in ADMIN_ROLES:
        return
    return jsonify({'success': False, 'error': '需要管理权限', 'code': 403}), 403


@app.route('/api/auth/login', methods=['POST'])
def admin_login():
    """
    在职且已设置密码的员工可用姓名或邮箱登录。
    超管进管理端，其他角色进门户（由前端按角色跳转）。
    """
    data = request.json or {}
    username = data.get('username', '').strip()
    password = data.get('password', '')
    
    if not username or not password:
        return jsonify({'success': False, 'error': '请输入账号和密码'}), 400
    
    conn = get_conn()
    try:
        # 在职员工：支持 name 或 email；不再限制仅超管
        user = conn.execute(
            """SELECT * FROM employees 
               WHERE (name = ? OR email = ?) AND status = '在职'""",
            (username, username)
        ).fetchone()
        
        if not user:
            return jsonify({'success': False, 'error': '账号不存在或无权限'}), 401
        
        # 校验密码（PBKDF2 哈希；兼容存量明文）。已移除历史后门默认密码 admin123（P0-2）。
        stored_pwd = user['password'] if user['password'] else ''
        if not stored_pwd:
            # 空密码账号无法登录，需管理员重置后方可使用
            return jsonify({'success': False, 'error': '密码错误'}), 401
        if not verify_password(password, stored_pwd):
            return jsonify({'success': False, 'error': '密码错误'}), 401
        # 存量明文密码：校验通过后立即升级为哈希（平滑迁移，根治明文存储 P1-3）
        if not stored_pwd.startswith('pbkdf2$'):
            conn.execute("UPDATE employees SET password = ? WHERE id = ?",
                         (hash_password(password), user['id']))
        
        # 更新最后登录时间
        conn.execute(
            "UPDATE employees SET last_login = datetime('now','localtime') WHERE id = ?",
            (user['id'],)
        )
        conn.commit()
        
        # 生成 token
        token = generate_token(user['id'])
        
        user_dict = dict(user)
        user_dict.pop('password', None)  # 不返回密码
        
        response = jsonify({
            'success': True,
            'user': user_dict,
            'token': token,
            'message': '登录成功'
        })
        
        # 设置 cookie（可选）
        response.set_cookie('ams_token', token, httponly=True, max_age=86400)
        
        return response
        
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/auth/logout', methods=['POST'])
def logout():
    """退出登录"""
    token = request.headers.get('Authorization', '').replace('Bearer ', '')
    if not token:
        token = request.cookies.get('ams_token')
    
    if token:
        conn = get_conn()
        conn.execute("DELETE FROM active_tokens WHERE token=?", (token,))
        conn.commit()
        conn.close()
    
    response = jsonify({'success': True, 'message': '已退出登录'})
    response.delete_cookie('ams_token')
    return response


@app.route('/api/auth/check', methods=['GET'])
def check_auth():
    """检查当前登录状态"""
    user_id = get_current_user_id()
    if not user_id:
        return jsonify({'success': True, 'is_logged_in': False})
    
    conn = get_conn()
    user = conn.execute("SELECT * FROM employees WHERE id = ?", (user_id,)).fetchone()
    conn.close()
    
    if not user:
        return jsonify({'success': True, 'is_logged_in': False})
    
    user_dict = dict(user)
    user_dict.pop('password', None)
    
    return jsonify({
        'success': True,
        'is_logged_in': True,
        'user': user_dict
    })


@app.route('/api/auth/lark-login', methods=['POST'])
def lark_login():
    """
    Lark 用户登录/绑定接口（已停用）
    前端通过 Lark JSAPI 获取用户信息后，调用此接口进行身份验证
    """
    # 安全加固（P0-3）：此接口直接信任前端回传的 lark_user_id/lark_open_id/name，
    # 无任何签名/授权码校验，可被直接调用伪造任意身份登录。已停用。
    # 统一走 /api/auth/lark-code-login（code 换 token，服务端向飞书换取真实身份）。
    # 前端 auth.js 当前仅使用 lark-code-login，此接口无实际调用方（仅废弃的 index-old.html 曾调用）。
    return jsonify({
        'success': False,
        'error': '该登录方式已停用（存在身份伪造风险），请通过飞书授权码登录',
        'use_code_login': True
    }), 410

    data = request.json or {}
    lark_user_id = data.get('lark_user_id')
    lark_open_id = data.get('lark_open_id')
    name = data.get('name', '')
    avatar = data.get('avatar', '')
    department = data.get('department', '')
    
    if not lark_user_id:
        return jsonify({'success': False, 'error': '缺少 lark_user_id'}), 400
    
    conn = get_conn()
    try:
        # 1. 查找是否已有绑定关系
        user = conn.execute(
            """SELECT e.* FROM employees e 
               WHERE e.lark_user_id = ? OR e.lark_open_id = ?""",
            (lark_user_id, lark_open_id)
        ).fetchone()
        
        if user:
            # 已绑定，更新信息（注意 employees 表部门字段为 dept，非 department）
            conn.execute(
                """UPDATE employees SET 
                   lark_user_id = ?, lark_open_id = ?, name = ?, 
                   avatar = ?, dept = ?, last_login = datetime('now','localtime')
                   WHERE id = ?""",
                (lark_user_id, lark_open_id, name, avatar, department, user['id'])
            )
            conn.commit()
            
            # 生成 token
            token = generate_token(user['id'])
            user_dict = dict(user)
            user_dict.pop('password', None)
            
            return jsonify({
                'success': True,
                'user': user_dict,
                'token': token,
                'is_new': False,
                'message': '登录成功'
            })
        else:
            # 未精确绑定：按姓名匹配未绑定的同名员工
            candidates = conn.execute(
                """SELECT id, name, dept FROM employees 
                   WHERE name = ? AND (lark_user_id IS NULL OR lark_user_id = '')
                   AND (lark_open_id IS NULL OR lark_open_id = '')
                   AND (lark_union_id IS NULL OR lark_union_id = '')
                   ORDER BY id""",
                (name,)
            ).fetchall()

            if candidates:
                # 存在未绑定的同名员工 → 直接自动绑定。
                # lark 的 open_id/union_id 唯一，绑定后该飞书账号只会落在这条员工
                # 记录上，不会因重名而错乱，故即便有多位同名也直接绑定首位。
                cand = candidates[0]
                conn.execute(
                    """UPDATE employees SET 
                       lark_user_id = ?, lark_open_id = ?, name = ?, 
                       avatar = ?, dept = ?, last_login = datetime('now','localtime')
                       WHERE id = ?""",
                    (lark_user_id, lark_open_id, name, avatar, department, cand['id'])
                )
                conn.commit()
                # 重新拉取完整员工记录，保证返回给前端的 user 含 role 等全部字段
                user = conn.execute("SELECT * FROM employees WHERE id = ?", (cand['id'],)).fetchone()
                token = generate_token(cand['id'])
                user_dict = dict(user)
                user_dict.pop('password', None)
                multi_hint = (f'（检测到 {len(candidates)} 位同名，已绑定首位）'
                              if len(candidates) > 1 else '')
                return jsonify({
                    'success': True,
                    'user': user_dict,
                    'token': token,
                    'is_new': False,
                    'bind_type': 'auto',
                    'message': f'已自动绑定同名员工：{name}{multi_hint}'
                })

            # 多名同名（歧义）或系统中无同名：提交待绑定申请，通知超管人工确认
            candidate_id = candidates[0]['id'] if candidates else None
            # 去重：同一飞书用户已有 pending 记录则直接返回，避免重复提交
            existing = conn.execute(
                """SELECT * FROM lark_pending_bindings 
                   WHERE status='pending' AND (lark_user_id=? OR lark_open_id=?)""",
                (lark_user_id, lark_open_id)
            ).fetchone()
            if existing:
                return jsonify({
                    'success': False,
                    'pending': True,
                    'pending_id': existing['id'],
                    'message': '绑定申请已提交，等待管理员审核'
                })

            cursor = conn.execute(
                """INSERT INTO lark_pending_bindings
                   (lark_user_id, lark_open_id, name, avatar, department,
                    candidate_employee_id, created_at)
                   VALUES (?, ?, ?, ?, ?, ?, datetime('now','localtime'))""",
                (lark_user_id, lark_open_id, name, avatar, department, candidate_id)
            )
            conn.commit()
            pending_id = cursor.lastrowid

            cand_hint = (f"（系统匹配到 {len(candidates)} 位同名员工，请确认应绑定哪一位）"
                         if len(candidates) > 1
                         else "（系统中无同名员工，可能需新建账号）")
            try:
                notify_lark_all(
                    f"🔔 【飞书绑定待审核】\n"
                    f"用户「{name}」通过飞书首次登录，请求绑定系统账号。{cand_hint}\n"
                    f"请前往「用户管理 - 飞书待绑定」手动审核。"
                )
            except Exception:
                pass

            return jsonify({
                'success': False,
                'pending': True,
                'pending_id': pending_id,
                'message': ('存在多名同名员工，已提交管理员审核' if len(candidates) > 1
                            else '绑定申请已提交，等待管理员审核')
            })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/auth/lark-signature', methods=['GET'])
def lark_signature():
    """生成 Lark H5 SDK config 所需的签名"""
    import hashlib, time, random, string
    app_id = os.environ.get('LARK_APP_ID', '')
    app_secret = os.environ.get('LARK_APP_SECRET', '')
    if not app_id or not app_secret:
        return jsonify({'success': False, 'error': '飞书登录尚未配置完成，请联系管理员'}), 500
    
    url = request.args.get('url', '')
    timestamp = str(int(time.time()))
    nonce = ''.join(random.choices(string.ascii_letters + string.digits, k=16))
    sign_str = ''.join(sorted([timestamp, nonce, app_secret]))
    signature = hashlib.sha1(sign_str.encode()).hexdigest()
    
    return jsonify({
        'success': True,
        'url': url,
        'appId': app_id,
        'timestamp': timestamp,
        'nonceStr': nonce,
        'signature': signature
    })


@app.route('/api/auth/lark-code-login', methods=['POST'])
def lark_code_login():
    """用飞书授权 code 登录（推荐方式，更稳定）"""
    import requests as req_lib
    
    data = request.json or {}
    code = data.get('code')
    if not code:
        return jsonify({'success': False, 'error': '缺少授权码，请从飞书工作台重新打开'}), 400
    
    app_id = os.environ.get('LARK_APP_ID', '')
    app_secret = os.environ.get('LARK_APP_SECRET', '')
    
    if not app_secret:
        return jsonify({'success': False, 'error': '飞书登录尚未配置完成，请联系管理员'}), 500
    
    # redirect_uri 必须与前端发起授权时使用的完全一致，否则飞书拒绝换 token。
    # 优先使用前端回传的实际回调地址，其次取环境变量，最后回退历史默认值。
    redirect_uri = (data.get('redirect_uri')
                    or os.environ.get('LARK_REDIRECT_URI')
                    or 'https://aap.hknhw.com')
    try:
        # 1. code 换 access_token（国际版用 open.larksuite.com）
        token_resp = req_lib.post('https://open.larksuite.com/open-apis/authen/v2/oauth/token',
            json={'grant_type': 'authorization_code', 'code': code,
                  'client_id': app_id, 'client_secret': app_secret,
                  'redirect_uri': redirect_uri},
            timeout=10)
        raw_text = token_resp.text
        try:
            import json as json_lib
            token_data = json_lib.loads(raw_text)
        except:
            return jsonify({'success': False, 'error': '飞书暂时没有正确回应，请稍后从工作台重新打开本系统'}), 500
        if token_data.get('code') != 0:
            return jsonify({'success': False, 'error': '飞书授权失败，请关闭后重新打开本系统'}), 400
        
        access_token = token_data['access_token']
        union_id = token_data.get('union_id', '')
        
        # 2. 拿用户信息（GET v1 接口）
        user_resp = req_lib.get('https://open.larksuite.com/open-apis/authen/v1/user_info',
            headers={'Authorization': f'Bearer {access_token}'}, timeout=10)
        user_data = user_resp.json()
        if user_data.get('code') != 0:
            return jsonify({'success': False, 'error': '未能读取飞书用户信息，请联系管理员检查应用权限'}), 400
        
        u = user_data['data']
        name = u.get('name', '')
        avatar = u.get('avatar_url', '')
        open_id = u.get('open_id', '')
        email = u.get('email', '')
        # 真正的 Lark User ID（租户内唯一、跨应用一致、管理员在 Lark 管理后台可见）。
        # 仅当应用在 Lark 开放平台申请了「获取用户 user ID」权限后，user_info 才返回该字段；
        # 未申请权限时为空，此时降级用 open_id/union_id。官方不建议把邮箱当登录凭证。
        user_id = u.get('user_id', '')
    except Exception as e:
        return jsonify({'success': False, 'error': '连接飞书失败，请稍后重试'}), 500
    
    # 3. 查/创建/绑定用户
    conn = get_conn()
    try:
        user = conn.execute(
            "SELECT * FROM employees WHERE lark_user_id = ? OR lark_open_id = ? OR lark_union_id = ?",
            (user_id, open_id, union_id)
        ).fetchone()

        # 安全加固（复审⑦）：离职/停用员工禁止登录（密码登录已校验在职，此处补齐 Lark 登录）
        if user and user['status'] != '在职':
            return jsonify({'success': False, 'error': '账号已离职或停用，请联系管理员'}), 403

        is_new = False
        if not user:
            # 未精确绑定：按姓名匹配「未绑定任何 Lark 账号」的同名在职员工
            # 安全加固（三轮审计#6）：同名一律转人工审核（不再自动绑定），
            # 防止同名飞书账号接管他人身份（无论普通用户还是管理角色）——
            # 由超管在「用户管理-飞书待绑定」人工确认后再绑定。
            candidates_all = conn.execute(
                "SELECT id, name, dept, role FROM employees WHERE name=? "
                "AND status='在职' "
                "AND (lark_user_id IS NULL OR lark_user_id='') "
                "AND (lark_open_id IS NULL OR lark_open_id='') "
                "AND (lark_union_id IS NULL OR lark_union_id='') ORDER BY id",
                (name,)
            ).fetchall()

            # 同名（无论角色）或无同名：均提交待绑定申请，由超管人工审核
            candidate_id = candidates_all[0]['id'] if candidates_all else None
            # 修复(四轮N1)：user_id 未申请权限时为空串，OR 条件会 ''='' 误命中他人 pending 记录，
            # 导致第二个新用户被误判"已提交"、不发飞书通知且绑定申请未真正提交。
            # 只对非空字段拼接去重条件；三者皆空则不去重（视为新申请）。
            dedup_conds, dedup_vals = [], []
            if user_id:
                dedup_conds.append("lark_user_id=?")
                dedup_vals.append(user_id)
            if open_id:
                dedup_conds.append("lark_open_id=?")
                dedup_vals.append(open_id)
            if union_id:
                dedup_conds.append("lark_union_id=?")
                dedup_vals.append(union_id)
            existing = None
            if dedup_conds:
                existing = conn.execute(
                    "SELECT * FROM lark_pending_bindings WHERE status='pending' AND (" + " OR ".join(dedup_conds) + ")",
                    tuple(dedup_vals)
                ).fetchone()
            if existing:
                return jsonify({
                    'success': False,
                    'pending': True,
                    'pending_id': existing['id'],
                    'message': '绑定申请已提交，等待管理员审核'
                })

            cursor = conn.execute(
                """INSERT INTO lark_pending_bindings
                   (lark_user_id, lark_open_id, lark_union_id, name, avatar, department, email,
                    candidate_employee_id, created_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, datetime('now','localtime'))""",
                (user_id, open_id, union_id, name, avatar, '', email, candidate_id)
            )
            conn.commit()
            pending_id = cursor.lastrowid

            # 同名提示基于全部同名（含被排除自动绑定的管理角色），确保数量准确
            if len(candidates_all) > 1:
                cand_hint = f"（系统匹配到 {len(candidates_all)} 位同名员工，请确认应绑定哪一位）"
            elif len(candidates_all) == 1:
                cand_hint = "（系统匹配到 1 位同名员工，请确认是否绑定）"
            else:
                cand_hint = "（系统中无同名员工，可能需新建账号）"
            try:
                notify_lark_all(
                    f"🔔 【飞书绑定待审核】\n"
                    f"用户「{name}」通过飞书首次登录，请求绑定系统账号。{cand_hint}\n"
                    f"请前往「用户管理 - 飞书待绑定」手动审核。"
                )
            except Exception:
                pass

            return jsonify({
                'success': False,
                'pending': True,
                'pending_id': pending_id,
                # 安全加固（复审2-⑤）：多名同名判断用全部同名 candidates_all（原误用过滤后 candidates 恒为空）
                'message': ('存在多名同名员工，已提交管理员审核' if len(candidates_all) > 1
                            else '绑定申请已提交，等待管理员审核')
            })
        
        # 已绑定用户：回填最新的 user_id/open_id/union_id（申请「获取用户 user ID」权限后，
        # 首次登录即可自动补齐真正的 user_id；历史 open_id/union_id 绑定也一并刷新）。
        # 只更新 ID 字段，不覆盖管理员手动维护的 name/avatar/dept。
        if user_id or open_id or union_id:
            conn.execute(
                """UPDATE employees SET
                   lark_user_id = CASE WHEN ? != '' THEN ? ELSE lark_user_id END,
                   lark_open_id = CASE WHEN ? != '' THEN ? ELSE lark_open_id END,
                   lark_union_id = CASE WHEN ? != '' THEN ? ELSE lark_union_id END,
                   last_login = datetime('now','localtime')
                   WHERE id = ?""",
                (user_id, user_id, open_id, open_id, union_id, union_id, user['id'])
            )
            conn.commit()
            user = conn.execute("SELECT * FROM employees WHERE id = ?", (user['id'],)).fetchone()

        token = generate_token(user['id'])
        user_dict = dict(user)
        user_dict.pop('password', None)
        
        return jsonify({
            'success': True,
            'user': user_dict,
            'token': token,
            'is_new': is_new,
            'message': 'Lark登录成功' + ('（新用户）' if is_new else '')
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/auth/user', methods=['GET'])
@login_required
def get_current_user_api():
    """获取当前登录用户信息"""
    conn = get_conn()
    user = conn.execute(
        "SELECT * FROM employees WHERE id = ?", (request.current_user_id,)
    ).fetchone()
    conn.close()
    
    if not user:
        return jsonify({'success': False, 'error': '用户不存在', 'code': 404}), 404
    
    user_dict = dict(user)
    user_dict.pop('password', None)
    
    return jsonify({
        'success': True,
        'user': user_dict
    })


@app.route('/api/auth/change-password', methods=['POST'])
@login_required
def change_password():
    """修改密码"""
    data = request.json or {}
    old_password = data.get('old_password', '')
    new_password = data.get('new_password', '')
    
    # 只强制要求新密码；旧密码允许为空（员工未设置过本地密码时可直接设置初始密码），
    # 是否必须匹配由下方「已设置过旧密码才校验」逻辑决定
    if not new_password:
        return jsonify({'success': False, 'error': '请提供新密码'}), 400

    # 密码强度策略（修复历史仅要求 4 位的弱密码问题 P1-3）
    if len(new_password) < 8:
        return jsonify({'success': False, 'error': '新密码至少8位'}), 400
    if not (re.search(r'[A-Za-z]', new_password) and re.search(r'\d', new_password)):
        return jsonify({'success': False, 'error': '新密码需同时包含字母和数字'}), 400

    conn = get_conn()
    try:
        user = conn.execute(
            "SELECT * FROM employees WHERE id = ?", (request.current_user_id,)
        ).fetchone()

        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404

        # 验证旧密码（PBKDF2 哈希；兼容存量明文）
        stored_pwd = user['password'] if user['password'] else ''
        if stored_pwd and not verify_password(old_password, stored_pwd):
            return jsonify({'success': False, 'error': '旧密码错误'}), 401

        # 更新密码（加盐哈希存储，根治明文入库）
        conn.execute(
            "UPDATE employees SET password = ? WHERE id = ?",
            (hash_password(new_password), request.current_user_id)
        )
        conn.commit()
        
        return jsonify({'success': True, 'message': '密码修改成功'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


# ═══════════════════════════════════════════════════════════════
# 用户管理 API（超管专用）
# ═══════════════════════════════════════════════════════════════

@app.route('/api/admin/users', methods=['GET'])
@superadmin_required
def get_all_users():
    """获取所有用户列表（超管）"""
    conn = get_conn()
    try:
        users = conn.execute('''
            SELECT id, name, role, dept, location, email, active, lark_only, status, last_login, created_at,
                   lark_user_id, lark_open_id, lark_union_id, avatar, custom_fields
            FROM employees
            ORDER BY created_at DESC
        ''').fetchall()
        return jsonify({'success': True, 'data': rows_to_list(users)})
    finally:
        conn.close()


@app.route('/api/admin/lark-pending', methods=['GET'])
@superadmin_required
def list_lark_pending():
    """列出飞书待绑定申请（超管）"""
    conn = get_conn()
    try:
        rows = conn.execute('''
            SELECT p.*, e.name AS candidate_name, e.dept AS candidate_dept
            FROM lark_pending_bindings p
            LEFT JOIN employees e ON p.candidate_employee_id = e.id
            WHERE p.status = 'pending'
            ORDER BY p.created_at DESC
        ''').fetchall()
        return jsonify({'success': True, 'data': rows_to_list(rows)})
    finally:
        conn.close()


@app.route('/api/admin/lark-bind', methods=['POST'])
@superadmin_required
def handle_lark_bind():
    """手动处理飞书待绑定（超管）：绑定到指定员工 / 新建账号 / 拒绝"""
    data = request.json or {}
    pending_id = data.get('pending_id')
    action = data.get('action')  # approve / reject
    employee_id = data.get('employee_id')  # approve 时绑定到该员工
    create_new = data.get('create_new', False)  # approve 且无 employee_id 时新建账号

    if not pending_id or action not in ('approve', 'reject'):
        return jsonify({'success': False, 'error': '参数错误'}), 400

    conn = get_conn()
    try:
        pending = conn.execute(
            "SELECT * FROM lark_pending_bindings WHERE id=?", (pending_id,)
        ).fetchone()
        if not pending:
            return jsonify({'success': False, 'error': '待绑定记录不存在'}), 404
        if pending['status'] != 'pending':
            return jsonify({'success': False, 'error': '该申请已处理'}), 400

        handler_id = request.current_user['id']

        if action == 'reject':
            conn.execute(
                """UPDATE lark_pending_bindings 
                   SET status='rejected', handled_by=?, handled_at=datetime('now','localtime')
                   WHERE id=?""",
                (handler_id, pending_id)
            )
            conn.commit()
            try:
                notify_lark_all(
                    f"⚠️ 【飞书绑定被拒绝】\n用户「{pending['name']}」的飞书绑定申请已被超管拒绝。"
                )
            except Exception:
                pass
            return jsonify({'success': True, 'message': '已拒绝该绑定申请'})

        # approve：决定绑定到哪个员工
        target_employee_id = employee_id
        if not target_employee_id and create_new:
            cursor = conn.execute(
                """INSERT INTO employees 
                   (name, role, dept, email, avatar, lark_user_id, lark_open_id, lark_union_id,
                    status, created_at, last_login)
                   VALUES (?, '普通用户', '', ?, ?, ?, ?, ?, '在职',
                           datetime('now','localtime'), datetime('now','localtime'))""",
                (pending['name'], pending['email'] or '', pending['avatar'] or '',
                 pending['lark_user_id'] or '', pending['lark_open_id'] or '',
                 pending['lark_union_id'] or '')
            )
            target_employee_id = cursor.lastrowid
        elif not target_employee_id:
            # 默认绑定到系统预匹配的同名员工
            if not pending['candidate_employee_id']:
                return jsonify({
                    'success': False,
                    'error': '未指定绑定员工，且系统中无同名员工可绑定，请传入 employee_id 或 create_new=true'
                }), 400
            target_employee_id = pending['candidate_employee_id']

        conn.execute(
            """UPDATE employees 
               SET lark_user_id=?, lark_open_id=?, lark_union_id=?, avatar=?, last_login=datetime('now','localtime')
               WHERE id=?""",
            (pending['lark_user_id'] or '', pending['lark_open_id'] or '',
             pending['lark_union_id'] or '', pending['avatar'] or '', target_employee_id)
        )
        conn.execute(
            """UPDATE lark_pending_bindings 
               SET status='approved', handled_by=?, handled_at=datetime('now','localtime')
               WHERE id=?""",
            (handler_id, pending_id)
        )
        conn.commit()

        emp = conn.execute("SELECT * FROM employees WHERE id=?", (target_employee_id,)).fetchone()
        try:
            notify_lark_all(
                f"✅ 【飞书绑定已通过】\n用户「{pending['name']}」已绑定系统账号「{emp['name']}」，现可通过飞书登录。"
            )
        except Exception:
            pass

        return jsonify({
            'success': True,
            'message': f"已绑定至员工 {emp['name']}",
            'employee_id': target_employee_id
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/admin/users', methods=['POST'])
@superadmin_required
def create_user():
    """创建新用户（超管）"""
    data = request.json or {}
    name = data.get('name', '').strip()
    role = data.get('role', '普通用户')
    location = data.get('location', '').strip()
    dept = data.get('dept', '').strip()
    email = data.get('email', '').strip()
    # 安全加固（复审⑦）：未指定密码时生成随机初始密码（不硬编码 Aaa77377），
    # 通过响应返回给超管转告用户，首次登录后应修改。
    password = data.get('password', '').strip()
    generated_pwd = None
    if not password:
        password = 'Init@' + secrets.token_urlsafe(6)
        generated_pwd = password
    lark_only = data.get('lark_only', 1)

    if not name:
        return jsonify({'success': False, 'error': '姓名不能为空'}), 400

    conn = get_conn()
    try:
        if email:
            existing = conn.execute("SELECT id FROM employees WHERE email = ?", (email,)).fetchone()
            if existing:
                return jsonify({'success': False, 'error': '邮箱已被使用'}), 400

        cur = conn.execute('''
            INSERT INTO employees (name, role, dept, location, email, password, lark_only, status, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, '在职', datetime('now','localtime'))
        ''', (name, role, dept, location, email, hash_password(password), lark_only))
        conn.commit()
        operator = request.current_user.get('name') if getattr(request, 'current_user', None) else '系统'
        log_action('create', 'employee', 'employee', cur.lastrowid, name, name, operator)
        resp = {'success': True, 'id': cur.lastrowid, 'message': '用户创建成功'}
        if generated_pwd:
            resp['initial_password'] = generated_pwd
            resp['message'] = '用户创建成功，已生成随机初始密码（请转告用户，首次登录后修改）'
        return jsonify(resp)
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/admin/users/<int:user_id>', methods=['PUT'])
@superadmin_required
def update_user(user_id):
    """更新用户信息（超管）"""
    data = request.json or {}
    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (user_id,)).fetchone()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404

        # 超管的 lark_only 始终为 0，不可修改
        if user['role'] == '超管':
            data.pop('lark_only', None)

        fields, values = [], []
        for field in ['name', 'role', 'dept', 'location', 'email', 'status', 'active', 'lark_only', 'password', 'lark_user_id', 'lark_open_id', 'lark_union_id', 'custom_fields']:
            if field in data:
                # 安全加固（复审P1-1）：超管改密码必须哈希入库，不再明文落库
                if field == 'password':
                    fields.append("password = ?")
                    values.append(hash_password(data[field]))
                else:
                    fields.append(f"{field} = ?")
                    values.append(data[field])

        # active 与 status 保持同步：拨动在职/离职开关后，两字段一致（active=0 视为离职）
        if 'active' in data:
            fields.append("status = ?")
            values.append('在职' if data['active'] else '离职')

        if not fields:
            return jsonify({'success': False, 'error': '没有要更新的字段'}), 400

        values.append(user_id)
        conn.execute(f"UPDATE employees SET {', '.join(fields)} WHERE id = ?", values)
        conn.commit()
        log_action('update', 'employee', 'employee', user_id, user['name'], user['name'],
                   get_current_user()['name'] if get_current_user() else '系统')
        return jsonify({'success': True, 'message': '用户更新成功'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/admin/users/<int:user_id>/lark-only', methods=['PUT'])
@superadmin_required
def set_user_lark_only(user_id):
    """设置用户是否强制 Lark 访问（超管专属）"""
    data = request.json or {}
    lark_only = data.get('lark_only', 1)

    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (user_id,)).fetchone()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404
        if user['role'] == '超管':
            return jsonify({'success': False, 'error': '不能修改超管的 Lark 限制设置'}), 403

        conn.execute("UPDATE employees SET lark_only = ? WHERE id = ?", (lark_only, user_id))
        conn.commit()
        action = '开启' if lark_only else '关闭'
        return jsonify({'success': True, 'message': f'已{action}该用户的 Lark 强制访问限制'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/admin/users/<int:user_id>', methods=['DELETE'])
@superadmin_required
def delete_user(user_id):
    """删除用户（超管）"""
    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (user_id,)).fetchone()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404
        if user['role'] == '超管':
            return jsonify({'success': False, 'error': '不能删除超管账号'}), 403
        # 不能删除当前登录的账号
        if get_current_user_id() == user_id:
            return jsonify({'success': False, 'error': '不能删除当前登录的账号'}), 403

        # 名下有关联资产/耗材领用/出入库记录则禁止删除，避免数据孤儿
        asset_cnt = conn.execute(
            "SELECT COUNT(*) AS c FROM assets WHERE (assignee = ? AND assignee IS NOT NULL AND assignee <> '') OR (assignee_id = ?)",
            (user['name'], user_id)
        ).fetchone()['c']
        # 耗材领用记录（target_user 按姓名）
        consumable_cnt = conn.execute(
            "SELECT COUNT(*) AS c FROM consumable_logs WHERE target_user = ? AND target_user IS NOT NULL AND target_user <> ''",
            (user['name'],)
        ).fetchone()['c']
        # 资产出入库记录（target_user 按姓名）
        stock_cnt = conn.execute(
            "SELECT COUNT(*) AS c FROM stock_logs WHERE target_user = ? AND target_user IS NOT NULL AND target_user <> ''",
            (user['name'],)
        ).fetchone()['c']

        if asset_cnt > 0 or consumable_cnt > 0 or stock_cnt > 0:
            return jsonify({
                'success': False,
                'error': f'该员工名下仍有在用/关联资产({asset_cnt})或耗材领用记录({consumable_cnt})或出入库记录({stock_cnt})，无法删除。请先转移/归还相关物资后再删除。'
            }), 409

        # 清理该员工的登录态，避免遗留有效 token
        conn.execute("DELETE FROM active_tokens WHERE user_id = ?", (user_id,))
        conn.execute("DELETE FROM employees WHERE id = ?", (user_id,))
        conn.commit()
        log_action('delete', 'employee', 'employee', user_id, user['name'], user['name'],
                   get_current_user()['name'] if get_current_user() else '系统')
        return jsonify({'success': True, 'message': '用户已删除'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


# ═══════════════════════════════════════════════════════════════
# 采购订单 API
# ═══════════════════════════════════════════════════════════════

def gen_order_no():
    """生成订单编号，格式 PO-YYYYMMDD-XXX"""
    conn = get_conn()
    try:
        today = datetime.now().strftime('%Y%m%d')
        prefix = f'PO-{today}-'
        last = conn.execute(
            "SELECT order_no FROM purchase_orders WHERE order_no LIKE ? ORDER BY id DESC LIMIT 1",
            (prefix + '%',)
        ).fetchone()
        seq = 1
        if last:
            try:
                seq = int(last['order_no'].split('-')[-1]) + 1
            except:
                pass
        return f'{prefix}{seq:03d}'
    finally:
        conn.close()

def _refresh_order_status(conn, order_id):
    """根据明细到货情况自动更新订单状态"""
    if not order_id:
        return
    row = conn.execute(
        "SELECT SUM(quantity) as tq, SUM(arrived_qty) as aq FROM order_items WHERE order_id = ?",
        (order_id,)
    ).fetchone()
    if not row or not row['tq']:
        return
    tq, aq = row['tq'] or 0, row['aq'] or 0
    if aq == 0:
        status = '待到货'
    elif aq < tq:
        status = '部分到货'
    else:
        status = '已到货'
    conn.execute(
        "UPDATE purchase_orders SET status=?, updated_at=datetime('now','localtime') WHERE id=?",
        (status, order_id)
    )


    try:
        today = datetime.now().strftime('%Y%m%d')
        prefix = f'PO-{today}-'
        last = conn.execute(
            "SELECT order_no FROM purchase_orders WHERE order_no LIKE ? ORDER BY id DESC LIMIT 1",
            (prefix + '%',)
        ).fetchone()
        seq = 1
        if last:
            try:
                seq = int(last['order_no'].split('-')[-1]) + 1
            except:
                pass
        return f'{prefix}{seq:03d}'
    finally:
        conn.close()


@app.route('/api/orders', methods=['GET'])
@login_required
def get_orders():
    """获取采购订单列表"""
    status = request.args.get('status', '')
    keyword = request.args.get('keyword', '')
    page = int(request.args.get('page', 1))
    page_size = int(request.args.get('page_size', 20))
    offset = (page - 1) * page_size

    conn = get_conn()
    try:
        conds, params = [], []
        if status:
            conds.append("status = ?")
            params.append(status)
        if keyword:
            conds.append("(order_no LIKE ? OR title LIKE ? OR supplier LIKE ?)")
            params += [f'%{keyword}%'] * 3
        where = ('WHERE ' + ' AND '.join(conds)) if conds else ''

        total = conn.execute(f"SELECT COUNT(*) FROM purchase_orders {where}", params).fetchone()[0]
        rows = conn.execute(
            f"SELECT * FROM purchase_orders {where} ORDER BY id DESC LIMIT ? OFFSET ?",
            params + [page_size, offset]
        ).fetchall()

        # 每个订单附带明细数量和已到货情况
        result = []
        for r in rows:
            d = dict(r)
            items = conn.execute(
                "SELECT COUNT(*) as cnt, SUM(quantity) as total_qty, SUM(arrived_qty) as arrived_qty FROM order_items WHERE order_id = ?",
                (d['id'],)
            ).fetchone()
            d['items_count'] = items['cnt'] or 0
            d['total_qty'] = items['total_qty'] or 0
            d['arrived_qty'] = items['arrived_qty'] or 0
            result.append(d)

        return jsonify({'success': True, 'data': result, 'total': total})
    finally:
        conn.close()


@app.route('/api/orders/<int:order_id>', methods=['GET'])
@login_required
def get_order_detail(order_id):
    """获取订单详情（含明细和关联资产/耗材）"""
    conn = get_conn()
    try:
        order = conn.execute("SELECT * FROM purchase_orders WHERE id = ?", (order_id,)).fetchone()
        if not order:
            return jsonify({'success': False, 'error': '订单不存在'}), 404

        order_dict = dict(order)

        # 订单明细
        items = conn.execute(
            "SELECT * FROM order_items WHERE order_id = ? ORDER BY id", (order_id,)
        ).fetchall()
        order_dict['items'] = [dict(i) for i in items]

        # 关联的资产列表
        assets = conn.execute(
            """SELECT id, asset_no, category, brand, model, sn, status, assignee, order_item_id
               FROM assets WHERE order_id = ? ORDER BY id""",
            (order_id,)
        ).fetchall()
        order_dict['assets'] = [dict(a) for a in assets]

        # 关联的耗材列表
        consumables = conn.execute(
            """SELECT id, item_no, name, category, brand, model, quantity, order_item_id
               FROM consumables WHERE order_id = ? ORDER BY id""",
            (order_id,)
        ).fetchall()
        order_dict['consumables'] = [dict(c) for c in consumables]

        return jsonify({'success': True, 'data': order_dict})
    finally:
        conn.close()


@app.route('/api/orders', methods=['POST'])
@login_required
def create_order():
    """创建采购订单"""
    data = request.json or {}
    title = data.get('title', '').strip()
    if not title:
        return jsonify({'success': False, 'error': '订单标题不能为空'}), 400

    order_no = data.get('order_no') or gen_order_no()
    conn = get_conn()
    try:
        cur = conn.execute('''
            INSERT INTO purchase_orders
            (order_no, title, supplier, order_date, expected_date, total_amount,
             currency, status, purchase_by, invoice_no, contract_no, notes, expense_order_no, created_by)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        ''', (
            order_no,
            title,
            data.get('supplier', ''),
            data.get('order_date', ''),
            data.get('expected_date', ''),
            data.get('total_amount', 0),
            data.get('currency', 'CNY'),
            data.get('status', '待到货'),
            data.get('purchase_by', ''),
            data.get('invoice_no', ''),
            data.get('contract_no', ''),
            data.get('notes', ''),
            data.get('expense_order_no', ''),
            data.get('created_by', ''),
        ))
        order_id = cur.lastrowid

        # 创建订单明细
        for item in data.get('items', []):
            conn.execute('''
                INSERT INTO order_items (order_id, item_type, item_name, category, brand, model, quantity, unit_price, notes)
                VALUES (?,?,?,?,?,?,?,?,?)
            ''', (
                order_id,
                item.get('item_type', 'asset'),
                item.get('item_name', ''),
                item.get('category', ''),
                item.get('brand', ''),
                item.get('model', ''),
                item.get('quantity', 1),
                item.get('unit_price', 0),
                item.get('notes', ''),
            ))

        conn.commit()
        try:
            log_action('create', 'order', 'purchase_order', order_id, order_no, title, {'items': len(data.get('items', []))})
        except Exception:
            pass
        return jsonify({'success': True, 'id': order_id, 'order_no': order_no, 'message': '订单创建成功'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/orders/<int:order_id>', methods=['PUT'])
@login_required
def update_order(order_id):
    """更新订单基本信息"""
    data = request.json or {}
    conn = get_conn()
    try:
        order = conn.execute("SELECT * FROM purchase_orders WHERE id = ?", (order_id,)).fetchone()
        if not order:
            return jsonify({'success': False, 'error': '订单不存在'}), 404

        fields, values = [], []
        for f in ['title', 'supplier', 'order_date', 'expected_date', 'actual_date',
                  'total_amount', 'currency', 'status', 'purchase_by',
                  'invoice_no', 'contract_no', 'notes', 'expense_order_no']:
            if f in data:
                fields.append(f'{f} = ?')
                values.append(data[f])

        if fields:
            fields.append("updated_at = datetime('now','localtime')")
            values.append(order_id)
            conn.execute(f"UPDATE purchase_orders SET {', '.join(fields)} WHERE id = ?", values)

        conn.commit()
        return jsonify({'success': True, 'message': '订单更新成功'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/orders/<int:order_id>/arrive', methods=['POST'])
@login_required
def confirm_order_arrival(order_id):
    """确认订单到货：更新状态 + 自动创建资产/耗材记录
    资产编号格式: YYYYMMDD001-主体-地区-类型
    """
    data = request.json or {}
    operator = get_current_name()
    conn = get_conn()
    try:
        order = conn.execute("SELECT * FROM purchase_orders WHERE id = ?", (order_id,)).fetchone()
        if not order:
            return jsonify({'success': False, 'error': '订单不存在'}), 404
        
        order_dict = dict(order)
        if order_dict['status'] in ('已到货', '已取消'):
            return jsonify({'success': False, 'error': f'订单状态已是「{order_dict["status"]}」，无法重复确认到货'}), 400

        # 获取主体和地区信息
        subject_code = order_dict.get('subject_code', '') or data.get('subject_code', '')
        region_code = order_dict.get('region_code', '') or data.get('region_code', '')

        # 缓存资产类型映射
        type_map = {}
        rows = conn.execute("SELECT code, name FROM asset_types WHERE active=1").fetchall()
        for r in rows:
            type_map[r['name']] = r['code']
            type_map[r['code']] = r['code']  # 也支持直接用code

        # 获取订单明细
        items = conn.execute(
            "SELECT * FROM order_items WHERE order_id = ?", (order_id,)
        ).fetchall()

        created_assets = []
        created_consumables = []

        today = datetime.now().strftime('%Y-%m-%d')

        for item in items:
            item_dict = dict(item)
            qty = item_dict['quantity']
            
            if item_dict['item_type'] == 'asset':
                # 映射资产类型代码
                category_name = item_dict.get('category', '')
                asset_type_code = type_map.get(category_name, 'OTH')
                
                # 为每个数量创建资产记录
                for i in range(qty):
                    asset_no = gen_asset_no(subject_code, region_code, asset_type_code)
                    cur = conn.execute('''
                        INSERT INTO assets
                        (asset_no, category, brand, model, status, purchase_date,
                         purchase_price, notes, order_id, order_item_id,
                         subject_code, region_code, asset_type_code)
                        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
                    ''', (
                        asset_no,
                        category_name,
                        item_dict.get('brand', ''),
                        item_dict.get('model', ''),
                        '在库',
                        today,
                        item_dict.get('unit_price', 0),
                        f"采购订单 {order_dict['order_no']} - {item_dict['item_name']}",
                        order_id,
                        item_dict['id'],
                        subject_code,
                        region_code,
                        asset_type_code,
                    ))
                    # 记录入库日志
                    conn.execute('''
                        INSERT INTO stock_logs
                        (asset_id, asset_no, log_type, operator, remark, log_time, order_id)
                        VALUES (?,?,?,?,?,?,?)
                    ''', (
                        cur.lastrowid, asset_no, '入库', operator,
                        f"采购订单到货入库 [{order_dict['order_no']}] {item_dict['item_name']}",
                        datetime.now().strftime('%Y-%m-%d %H:%M:%S'), order_id
                    ))
                    created_assets.append({'asset_no': asset_no, 'name': item_dict['item_name']})
                
                # 更新明细已到货数量
                conn.execute(
                    "UPDATE order_items SET arrived_qty = ? WHERE id = ?",
                    (qty, item_dict['id'])
                )
            else:
                # 耗材：创建或更新耗材记录
                existing = conn.execute(
                    """SELECT id, quantity FROM consumables 
                       WHERE name=? AND (brand=? OR brand IS NULL) AND (model=? OR model IS NULL)
                       LIMIT 1""",
                    (item_dict['item_name'], item_dict.get('brand', ''), item_dict.get('model', ''))
                ).fetchone()

                if existing:
                    # 累加库存
                    new_qty = existing['quantity'] + qty
                    conn.execute(
                        "UPDATE consumables SET quantity = quantity + ? WHERE id = ?",
                        (qty, existing['id'])
                    )
                    consumable_id = existing['id']
                else:
                    # 新建耗材记录
                    item_no = gen_consumable_no(item_dict.get('category', ''))
                    cur = conn.execute('''
                        INSERT INTO consumables
                        (item_no, name, category, brand, model, quantity, order_id, order_item_id)
                        VALUES (?,?,?,?,?,?,?,?)
                    ''', (
                        item_no,
                        item_dict['item_name'],
                        item_dict.get('category', ''),
                        item_dict.get('brand', ''),
                        item_dict.get('model', ''),
                        qty,
                        order_id,
                        item_dict['id']
                    ))
                    consumable_id = cur.lastrowid
                    created_consumables.append({'item_no': item_no, 'name': item_dict['item_name']})
                    new_qty = qty

                # 记录耗材入库日志
                conn.execute('''
                    INSERT INTO consumable_logs
                    (item_id, item_no, log_type, quantity, operator, remark, log_time, order_id)
                    VALUES (?,?,?,?,?,?,?,?)
                ''', (
                    consumable_id,
                    item_dict.get('item_name', ''),
                    '入库',
                    qty,
                    operator,
                    f"采购订单到货入库 [{order_dict['order_no']}] {item_dict['item_name']}",
                    datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                    order_id
                ))
                
                # 更新明细已到货数量
                conn.execute(
                    "UPDATE order_items SET arrived_qty = ? WHERE id = ?",
                    (qty, item_dict['id'])
                )

        # 更新订单状态为已到货
        conn.execute(
            """UPDATE purchase_orders 
               SET status='已到货', actual_date=?, updated_at=datetime('now','localtime')
               WHERE id=?""",
            (today, order_id)
        )

        conn.commit()

        return jsonify({
            'success': True,
            'message': f'确认到货成功！已创建 {len(created_assets)} 个资产、{len(created_consumables)} 个耗材',
            'created_assets': created_assets,
            'created_consumables': created_consumables
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


# ── 订单发票 & 附件上传 ──────────────────────────────────────────────

ALLOWED_DOC_EXTS = {'.pdf', '.jpg', '.jpeg', '.png', '.gif', '.webp', '.xlsx', '.xls', '.docx', '.doc', '.zip'}

@app.route('/api/orders/<int:order_id>/invoice', methods=['POST'])
@login_required
def upload_order_invoice(order_id):
    """上传订单发票文件（图片或PDF）"""
    conn = get_conn()
    try:
        if not conn.execute("SELECT id FROM purchase_orders WHERE id=?", (order_id,)).fetchone():
            return jsonify({'success': False, 'error': '订单不存在'}), 404
    finally:
        conn.close()

    if 'file' not in request.files:
        return jsonify({'success': False, 'error': '没有文件字段'}), 400
    file = request.files['file']
    if not file or file.filename == '':
        return jsonify({'success': False, 'error': '未选择文件'}), 400

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ALLOWED_DOC_EXTS:
        return jsonify({'success': False, 'error': f'不支持的文件类型 {ext}'}), 400

    filename = secure_filename(file.filename)
    timestamp = _rand_suffix()  # 安全加固 P2-4：随机段替代可推算时间戳，防附件URL枚举
    new_filename = f"invoice_order{order_id}_{timestamp}{ext}"
    file_path = os.path.join(UPLOAD_DIR, new_filename)
    file.save(file_path)

    conn = get_conn()
    try:
        # 写入 photos 表（ref_type = 'order_invoice'）
        cur = conn.execute('''
            INSERT INTO photos(ref_type, ref_id, file_name, file_path, file_size, mime_type, uploaded_by)
            VALUES(?,?,?,?,?,?,?)
        ''', ('order_invoice', order_id, new_filename, file_path,
              os.path.getsize(file_path), file.content_type,
              request.form.get('uploaded_by', '系统')))
        photo_id = cur.lastrowid
        # 同步更新订单的 invoice_file 字段（存最新一份文件名）
        conn.execute("UPDATE purchase_orders SET invoice_file=? WHERE id=?", (new_filename, order_id))
        conn.commit()
        return jsonify({
            'success': True, 'photo_id': photo_id,
            'file_name': new_filename,
            'url': f'/api/uploads/{new_filename}'
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/orders/<int:order_id>/invoices', methods=['GET'])
@login_required
def get_order_invoices(order_id):
    """获取订单所有发票/附件文件列表"""
    conn = get_conn()
    try:
        rows = conn.execute(
            "SELECT * FROM photos WHERE ref_type='order_invoice' AND ref_id=? ORDER BY uploaded_at DESC",
            (order_id,)
        ).fetchall()
        data = rows_to_list(rows)
        for r in data:
            r['url'] = f"/api/uploads/{r['file_name']}"
        return jsonify({'success': True, 'data': data})
    finally:
        conn.close()


@app.route('/api/orders/<int:order_id>/invoices/<int:photo_id>', methods=['DELETE'])
@login_required
def delete_order_invoice(order_id, photo_id):
    """删除发票附件"""
    conn = get_conn()
    try:
        row = conn.execute("SELECT * FROM photos WHERE id=? AND ref_type='order_invoice' AND ref_id=?",
                           (photo_id, order_id)).fetchone()
        if not row:
            return jsonify({'success': False, 'error': '文件不存在'}), 404
        # 删除磁盘文件
        try:
            os.remove(row['file_path'])
        except Exception:
            pass
        conn.execute("DELETE FROM photos WHERE id=?", (photo_id,))
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


# ── 资产/耗材实物图上传 ──────────────────────────────────────────────

ALLOWED_IMG_EXTS = {'.jpg', '.jpeg', '.png', '.gif', '.webp', '.heic'}

@app.route('/api/assets/<int:asset_id>/photos', methods=['POST'])
@login_required
def upload_asset_photo(asset_id):
    """上传资产实物图"""
    conn = get_conn()
    try:
        if not conn.execute("SELECT id FROM assets WHERE id=?", (asset_id,)).fetchone():
            return jsonify({'success': False, 'error': '资产不存在'}), 404
    finally:
        conn.close()

    if 'file' not in request.files:
        return jsonify({'success': False, 'error': '没有文件字段'}), 400
    file = request.files['file']
    if not file or file.filename == '':
        return jsonify({'success': False, 'error': '未选择文件'}), 400

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ALLOWED_IMG_EXTS:
        return jsonify({'success': False, 'error': f'请上传图片文件（支持 jpg/png/gif/webp/heic）'}), 400

    filename = secure_filename(file.filename)
    timestamp = _rand_suffix()  # 安全加固 P2-4：随机段替代可推算时间戳，防附件URL枚举
    new_filename = f"asset_{asset_id}_{timestamp}{ext}"
    file_path = os.path.join(UPLOAD_DIR, new_filename)
    file.save(file_path)

    # 防御性大小限制（前端已校验，后端兜底）
    if os.path.getsize(file_path) > 10 * 1024 * 1024:
        try: os.remove(file_path)
        except OSError: pass
        return jsonify({'success': False, 'error': '文件超过 10MB 限制'}), 400

    conn = get_conn()
    try:
        cur = conn.execute('''
            INSERT INTO photos(ref_type, ref_id, file_name, file_path, file_size, mime_type, uploaded_by)
            VALUES(?,?,?,?,?,?,?)
        ''', ('asset', asset_id, new_filename, file_path,
              os.path.getsize(file_path), file.content_type,
              request.form.get('uploaded_by', '系统')))
        photo_id = cur.lastrowid
        conn.commit()
        return jsonify({'success': True, 'photo_id': photo_id,
                        'file_name': new_filename, 'url': f'/api/uploads/{new_filename}'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/assets/<int:asset_id>/photos', methods=['GET'])
@login_required
def get_asset_photos(asset_id):
    """获取资产实物图列表"""
    conn = get_conn()
    try:
        rows = conn.execute(
            "SELECT * FROM photos WHERE ref_type='asset' AND ref_id=? ORDER BY uploaded_at DESC",
            (asset_id,)
        ).fetchall()
        data = rows_to_list(rows)
        for r in data:
            r['url'] = f"/api/uploads/{r['file_name']}"
        return jsonify({'success': True, 'data': data})
    finally:
        conn.close()


@app.route('/api/assets/<int:asset_id>/photos/<int:photo_id>', methods=['DELETE'])
@login_required
def delete_asset_photo(asset_id, photo_id):
    """删除资产实物图"""
    conn = get_conn()
    try:
        row = conn.execute("SELECT * FROM photos WHERE id=? AND ref_type='asset' AND ref_id=?",
                           (photo_id, asset_id)).fetchone()
        if not row:
            return jsonify({'success': False, 'error': '图片不存在'}), 404
        try:
            os.remove(row['file_path'])
        except Exception:
            pass
        conn.execute("DELETE FROM photos WHERE id=?", (photo_id,))
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/consumables/<int:item_id>/photos', methods=['POST'])
@login_required
def upload_consumable_photo(item_id):
    """上传耗材实物图"""
    conn = get_conn()
    try:
        if not conn.execute("SELECT id FROM consumables WHERE id=?", (item_id,)).fetchone():
            return jsonify({'success': False, 'error': '耗材不存在'}), 404
    finally:
        conn.close()

    if 'file' not in request.files:
        return jsonify({'success': False, 'error': '没有文件字段'}), 400
    file = request.files['file']
    if not file or file.filename == '':
        return jsonify({'success': False, 'error': '未选择文件'}), 400

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ALLOWED_IMG_EXTS:
        return jsonify({'success': False, 'error': '请上传图片文件（支持 jpg/png/gif/webp/heic）'}), 400

    filename = secure_filename(file.filename)
    timestamp = _rand_suffix()  # 安全加固 P2-4：随机段替代可推算时间戳，防附件URL枚举
    new_filename = f"consumable_{item_id}_{timestamp}{ext}"
    file_path = os.path.join(UPLOAD_DIR, new_filename)
    file.save(file_path)

    conn = get_conn()
    try:
        cur = conn.execute('''
            INSERT INTO photos(ref_type, ref_id, file_name, file_path, file_size, mime_type, uploaded_by)
            VALUES(?,?,?,?,?,?,?)
        ''', ('consumable', item_id, new_filename, file_path,
              os.path.getsize(file_path), file.content_type,
              request.form.get('uploaded_by', '系统')))
        photo_id = cur.lastrowid
        conn.commit()
        return jsonify({'success': True, 'photo_id': photo_id,
                        'file_name': new_filename, 'url': f'/api/uploads/{new_filename}'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/consumables/<int:item_id>/photos', methods=['GET'])
@login_required
def get_consumable_photos(item_id):
    """获取耗材实物图列表"""
    conn = get_conn()
    try:
        rows = conn.execute(
            "SELECT * FROM photos WHERE ref_type='consumable' AND ref_id=? ORDER BY uploaded_at DESC",
            (item_id,)
        ).fetchall()
        data = rows_to_list(rows)
        for r in data:
            r['url'] = f"/api/uploads/{r['file_name']}"
        return jsonify({'success': True, 'data': data})
    finally:
        conn.close()


@app.route('/api/orders/<int:order_id>/items', methods=['POST'])
@login_required
def add_order_item(order_id):
    """添加订单明细"""
    data = request.json or {}
    conn = get_conn()
    try:
        if not conn.execute("SELECT id FROM purchase_orders WHERE id = ?", (order_id,)).fetchone():
            return jsonify({'success': False, 'error': '订单不存在'}), 404

        cur = conn.execute('''
            INSERT INTO order_items (order_id, item_type, item_name, category, brand, model, quantity, unit_price, notes)
            VALUES (?,?,?,?,?,?,?,?,?)
        ''', (
            order_id,
            data.get('item_type', 'asset'),
            data.get('item_name', ''),
            data.get('category', ''),
            data.get('brand', ''),
            data.get('model', ''),
            data.get('quantity', 1),
            data.get('unit_price', 0),
            data.get('notes', ''),
        ))
        conn.commit()
        return jsonify({'success': True, 'id': cur.lastrowid})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/orders/<int:order_id>/items/<int:item_id>', methods=['DELETE'])
@login_required
def delete_order_item(order_id, item_id):
    """删除订单明细（仅当未关联资产时）"""
    conn = get_conn()
    try:
        linked = conn.execute(
            "SELECT COUNT(*) FROM assets WHERE order_item_id = ?", (item_id,)
        ).fetchone()[0]
        linked += conn.execute(
            "SELECT COUNT(*) FROM consumables WHERE order_item_id = ?", (item_id,)
        ).fetchone()[0]
        if linked:
            return jsonify({'success': False, 'error': f'该明细已关联 {linked} 个资产/耗材，不能删除'}), 400

        conn.execute("DELETE FROM order_items WHERE id = ? AND order_id = ?", (item_id, order_id))
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/orders/search', methods=['GET'])
@login_required
def search_orders():
    """搜索订单（入库表单下拉用）"""
    keyword = request.args.get('q', '')
    conn = get_conn()
    try:
        clause, params = fuzzy_match_clause(keyword, ['order_no', 'title', 'supplier'])
        sql = """SELECT id, order_no, title, supplier, status
               FROM purchase_orders
               WHERE status != '已取消'"""
        if clause:
            sql += f" AND {clause}"
        sql += " ORDER BY id DESC LIMIT 20"
        rows = conn.execute(sql, params).fetchall()
        return jsonify({'success': True, 'data': rows_to_list(rows)})
    finally:
        conn.close()


# ═══════════════════════════════════════════════════════════════
# 个人中心 API
# ═══════════════════════════════════════════════════════════════

@app.route('/api/profile/assets', methods=['GET'])
@login_required
def get_my_assets():
    """获取当前用户名下的资产"""
    conn = get_conn()
    user = conn.execute(
        "SELECT name FROM employees WHERE id = ?", (request.current_user_id,)
    ).fetchone()
    
    if not user:
        conn.close()
        return jsonify({'success': False, 'error': '用户不存在'}), 404
    
    user_name = user['name']
    
    # 查询该用户名下的资产
    assets = conn.execute(
        """SELECT * FROM assets 
           WHERE assignee = ? 
           ORDER BY updated_at DESC""",
        (user_name,)
    ).fetchall()
    
    conn.close()
    
    return jsonify({
        'success': True,
        'assets': assets_to_safe_list(assets),
        'total': len(assets)
    })


@app.route('/api/profile/consumables', methods=['GET'])
@login_required
def get_my_consumables():
    """获取当前用户的耗材领用记录"""
    conn = get_conn()
    user = conn.execute(
        "SELECT name FROM employees WHERE id = ?", (request.current_user_id,)
    ).fetchone()
    
    if not user:
        conn.close()
        return jsonify({'success': False, 'error': '用户不存在'}), 404
    
    user_name = user['name']
    
    # 查询该用户的耗材领用统计
    consumables = conn.execute(
        """SELECT c.*, SUM(l.quantity) as total_out,
                  MAX(l.log_time) as last_out_time
           FROM consumables c
           JOIN consumable_logs l ON c.id = l.item_id
           WHERE l.target_user = ? AND l.log_type = '出库'
           GROUP BY c.id
           ORDER BY last_out_time DESC""",
        (user_name,)
    ).fetchall()
    
    conn.close()
    
    return jsonify({
        'success': True,
        'consumables': rows_to_list(consumables),
        'total': len(consumables)
    })


@app.route('/api/profile/inventory-pending', methods=['GET'])
@login_required
def list_my_inventory_pending():
    """当前用户需自助确认的盘点明细"""
    conn = get_conn()
    user = conn.execute(
        "SELECT name FROM employees WHERE id = ?", (request.current_user_id,)
    ).fetchone()
    if not user:
        conn.close()
        return jsonify({'success': False, 'error': '用户不存在'}), 404
    rows = conn.execute(
        """SELECT ii.id, ii.asset_id, ii.asset_no, ii.asset_info, ii.before_status,
                  ii.before_location, ii.before_assignee, ii.task_id, it.task_name, it.task_no
           FROM inventory_items ii
           JOIN inventory_tasks it ON ii.task_id = it.id
           WHERE ii.before_assignee = ?
             AND ii.confirm_required = 1
             AND ii.check_status = 'pending'
             AND it.status IN ('ongoing', '进行中')
           ORDER BY ii.id""",
        (user['name'],)
    ).fetchall()
    conn.close()
    return jsonify({'success': True, 'data': rows_to_list(rows), 'total': len(rows)})


@app.route('/api/profile/todos', methods=['GET'])
@login_required
def get_my_todos():
    """获取当前用户的待办事项"""
    conn = get_conn()
    user = conn.execute(
        "SELECT name, role FROM employees WHERE id = ?", (request.current_user_id,)
    ).fetchone()
    
    if not user:
        conn.close()
        return jsonify({'success': False, 'error': '用户不存在'}), 404
    
    user_name = user['name']
    user_role = user['role']
    
    todos = []
    
    # 1. 盘点确认（如果用户名下有分发资产正在盘点）
    inventory_items = conn.execute(
        """SELECT ii.*, it.task_name 
           FROM inventory_items ii
           JOIN inventory_tasks it ON ii.task_id = it.id
           WHERE ii.before_assignee = ? 
             AND ii.confirm_required = 1 
             AND ii.check_status = 'pending'
             AND it.status IN ('ongoing', '进行中')""",
        (user_name,)
    ).fetchall()
    
    if inventory_items:
        todos.append({
            'id': 'inventory_' + str(inventory_items[0]['id']),
            'type': 'inventory',
            'icon': '📋',
            'title': '资产盘点确认',
            'desc': f'您有 {len(inventory_items)} 项资产需要确认盘点',
            'count': len(inventory_items)
        })
    
    # 2. 审批待办（如果是管理员）
    if user_role in ['超管', '普通管理员']:
        pending_approvals = conn.execute(
            """SELECT COUNT(*) as cnt FROM delete_applications 
               WHERE status = 'pending'"""
        ).fetchone()
        
        if pending_approvals['cnt'] > 0:
            todos.append({
                'id': 'approval_delete',
                'type': 'approval',
                'icon': '✅',
                'title': '删除申请审批',
                'desc': f'有 {pending_approvals["cnt"]} 条删除申请待审批',
                'count': pending_approvals['cnt']
            })
    
    # 3. 维修中资产
    repairing_assets = conn.execute(
        """SELECT COUNT(*) as cnt FROM assets 
           WHERE assignee = ? AND status = '维修中'""",
        (user_name,)
    ).fetchone()
    
    if repairing_assets['cnt'] > 0:
        todos.append({
            'id': 'repairing',
            'type': 'repair',
            'icon': '🔧',
            'title': '维修中资产',
            'desc': f'您有 {repairing_assets["cnt"]} 项资产正在维修',
            'count': repairing_assets['cnt']
        })
    
    conn.close()
    
    return jsonify({
        'success': True,
        'todos': todos,
        'total': len(todos)
    })


@app.route('/api/profile/apply/asset', methods=['POST'])
@login_required
def apply_asset():
    """申请资产：写入 asset_applications，待审批后执行（新资产类申请只结案不自动建台账）"""
    data = request.json or {}
    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (request.current_user_id,)).fetchone()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404
        reason = data.get('reason', '')
        detail = {
            'type': data.get('type'),
            'category': data.get('category'),
            'brand_model': data.get('brand_model'),
            'reason': reason,
        }
        app_id, app_no = insert_asset_application(
            conn, user, 'new_asset', reason=reason, detail=detail)
        conn.commit()
        log_action('apply', 'asset_application', 'asset_applications', app_id, app_no, user['name'])
        return jsonify({
            'success': True,
            'message': '申请已提交，请等待审批',
            'app_no': app_no,
            'id': app_id,
            'apply_info': {
                'applicant': user['name'],
                'department': user['dept'],
                'type': data.get('type'),
                'category': data.get('category'),
                'reason': reason,
            }
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/profile/apply/consumable', methods=['POST'])
@login_required
def apply_consumable():
    """申请领用耗材：只写申请单，不扣库存"""
    data = request.json or {}
    item_id = data.get('item_id')
    quantity = data.get('quantity', 1)
    purpose = data.get('purpose', '')
    if not item_id:
        return jsonify({'success': False, 'error': '请选择耗材'}), 400
    try:
        quantity = int(quantity)
    except (TypeError, ValueError):
        return jsonify({'success': False, 'error': '数量无效'}), 400
    if quantity < 1:
        return jsonify({'success': False, 'error': '数量至少为 1'}), 400

    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (request.current_user_id,)).fetchone()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404
        item = conn.execute("SELECT * FROM consumables WHERE id = ?", (item_id,)).fetchone()
        if not item:
            return jsonify({'success': False, 'error': '耗材不存在'}), 404
        if item['quantity'] < quantity:
            return jsonify({'success': False, 'error': f'库存不足，当前库存: {item["quantity"]}'}), 400
        app_id, app_no = insert_asset_application(
            conn, user, 'consumable',
            reason=purpose,
            detail={'item_id': item['id'], 'quantity': quantity, 'qty': quantity,
                    'item_no': item['item_no'], 'item_name': item['name'], 'purpose': purpose})
        conn.commit()
        log_action('apply', 'asset_application', 'asset_applications', app_id, app_no, item['name'])
        return jsonify({
            'success': True,
            'message': '申请已提交，请等待审批',
            'app_no': app_no,
            'id': app_id,
            'item_name': item['name'],
            'quantity': quantity,
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/profile/apply/repair', methods=['POST'])
@login_required
def apply_repair():
    """申请维修：只写申请单，不改资产状态"""
    data = request.json or {}
    asset_id = data.get('asset_id')
    fault_desc = data.get('fault_desc', '')
    if not asset_id:
        return jsonify({'success': False, 'error': '请选择资产'}), 400
    if not fault_desc:
        return jsonify({'success': False, 'error': '请描述故障'}), 400

    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (request.current_user_id,)).fetchone()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404
        asset = conn.execute("SELECT * FROM assets WHERE id = ?", (asset_id,)).fetchone()
        if not asset:
            return jsonify({'success': False, 'error': '资产不存在'}), 404
        if asset['assignee'] != user['name']:
            return jsonify({'success': False, 'error': '只能报修自己名下的资产'}), 403
        app_id, app_no = insert_asset_application(
            conn, user, 'repair',
            asset_id=asset['id'], asset_no=asset['asset_no'],
            reason=fault_desc, detail={'fault_desc': fault_desc})
        conn.commit()
        log_action('apply', 'asset_application', 'asset_applications', app_id, app_no, asset['asset_no'])
        return jsonify({
            'success': True,
            'message': '报修申请已提交，请等待审批',
            'app_no': app_no,
            'id': app_id,
            'asset_no': asset['asset_no'],
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/profile/apply/return', methods=['POST'])
@login_required
def apply_return():
    """申请归还名下在用资产：只写申请单，审批通过后才入库"""
    data = request.json or {}
    asset_id = data.get('asset_id')
    reason = data.get('reason', '')
    if not asset_id:
        return jsonify({'success': False, 'error': '请选择资产'}), 400
    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (request.current_user_id,)).fetchone()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404
        asset = conn.execute("SELECT * FROM assets WHERE id = ?", (asset_id,)).fetchone()
        if not asset:
            return jsonify({'success': False, 'error': '资产不存在'}), 404
        if (asset['assignee'] or '') != user['name']:
            return jsonify({'success': False, 'error': '只能归还自己名下的资产'}), 403
        if asset['status'] != '在用':
            return jsonify({'success': False, 'error': f'资产当前状态为【{asset["status"]}】，无法归还'}), 400
        app_id, app_no = insert_asset_application(
            conn, user, 'return',
            asset_id=asset['id'], asset_no=asset['asset_no'],
            reason=reason, detail={'reason': reason})
        conn.commit()
        log_action('apply', 'asset_application', 'asset_applications', app_id, app_no, asset['asset_no'])
        return jsonify({
            'success': True,
            'message': '归还申请已提交，请等待审批',
            'app_no': app_no,
            'id': app_id,
            'asset_no': asset['asset_no'],
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/profile/update', methods=['PUT', 'POST'])
@login_required
def update_own_profile():
    """员工自助修改资料：仅部门、邮箱、所在地"""
    data = request.json or {}
    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (request.current_user_id,)).fetchone()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404
        dept = data.get('dept', data.get('department'))
        email = data.get('email')
        location = data.get('location')
        updates, params = [], []
        if dept is not None:
            updates.append('dept=?')
            params.append(str(dept).strip())
        if location is not None:
            updates.append('location=?')
            params.append(str(location).strip())
        if email is not None:
            email = str(email).strip()
            if email:
                clash = conn.execute(
                    "SELECT id FROM employees WHERE email=? AND id<>?",
                    (email, request.current_user_id)).fetchone()
                if clash:
                    return jsonify({'success': False, 'error': '邮箱已被使用'}), 400
            updates.append('email=?')
            params.append(email)
        if not updates:
            return jsonify({'success': False, 'error': '没有可更新的字段'}), 400
        params.append(request.current_user_id)
        conn.execute(f"UPDATE employees SET {', '.join(updates)} WHERE id=?", params)
        conn.commit()
        row = conn.execute(
            "SELECT id, name, role, dept, location, email, avatar, status FROM employees WHERE id=?",
            (request.current_user_id,)).fetchone()
        return jsonify({'success': True, 'message': '资料已更新', 'user': dict(row)})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()

def gen_inventory_no() -> str:
    """生成盘点单号: INV-YYYYMMDD-序号"""
    conn = get_conn()
    today = datetime.now().strftime('%Y%m%d')
    like = f'INV-{today}-%'
    row = conn.execute(
        "SELECT COUNT(*) as cnt FROM inventory_tasks WHERE task_no LIKE ?", (like,)
    ).fetchone()
    seq = (row['cnt'] or 0) + 1
    conn.close()
    return f"INV-{today}-{seq:03d}"


@app.route('/api/inventory/tasks', methods=['POST'])
def create_inventory_task():
    """创建盘点任务"""
    data = request.json or {}
    task_name = data.get('task_name', '').strip()
    scope_type = data.get('scope_type', 'in_stock')  # all/in_stock/location
    scope_value = data.get('scope_value', '').strip()
    created_by = data.get('created_by', '系统')
    notes = data.get('notes', '')
    
    if not task_name:
        return jsonify({'success': False, 'error': '请输入盘点任务名称'}), 400
    
    conn = get_conn()
    try:
        # 1. 创建盘点任务
        task_no = gen_inventory_no()
        cursor = conn.execute(
            """INSERT INTO inventory_tasks 
               (task_no, task_name, scope_type, scope_value, created_by, notes)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (task_no, task_name, scope_type, scope_value, created_by, notes)
        )
        task_id = cursor.lastrowid
        
        # 2. 根据范围筛选资产
        where_clause = "WHERE status != '已删除' AND status != '待删除'"
        params = []
        
        if scope_type == 'in_stock':
            where_clause += " AND (status = '在库' OR assignee IS NULL OR assignee = '')"
        elif scope_type == 'location' and scope_value:
            where_clause += " AND location = ?"
            params.append(scope_value)
        # scope_type == 'all' 则不添加额外筛选
        
        assets = conn.execute(
            f"SELECT * FROM assets {where_clause}", params
        ).fetchall()
        
        total_count = len(assets)
        in_stock_count = 0
        assigned_count = 0
        
        # 3. 生成盘点明细
        for asset in assets:
            asset_dict = dict(asset)
            asset_info = json.dumps({
                'category': asset_dict.get('category'),
                'brand': asset_dict.get('brand'),
                'model': asset_dict.get('model'),
                'spec': asset_dict.get('spec'),
                'sn': asset_dict.get('sn'),
                'location': asset_dict.get('location'),
                'assignee': asset_dict.get('assignee'),
                'status': asset_dict.get('status')
            }, ensure_ascii=False)
            
            # 判断是否需要责任人确认（分发资产）
            is_assigned = asset_dict.get('assignee') and asset_dict.get('status') == '在用'
            confirm_required = 1 if is_assigned else 0
            if is_assigned:
                assigned_count += 1
            else:
                in_stock_count += 1
            
            conn.execute(
                """INSERT INTO inventory_items 
                   (task_id, asset_id, asset_no, asset_info, 
                    before_status, before_location, before_assignee,
                    confirm_required)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                (task_id, asset_dict['id'], asset_dict['asset_no'], asset_info,
                 asset_dict.get('status'), asset_dict.get('location'), asset_dict.get('assignee'),
                 confirm_required)
            )
        
        # 4. 更新任务统计
        conn.execute(
            """UPDATE inventory_tasks SET 
               total_count = ?, started_at = datetime('now','localtime'), status = 'ongoing'
               WHERE id = ?""",
            (total_count, task_id)
        )
        conn.commit()
        
        return jsonify({
            'success': True,
            'task': {
                'id': task_id,
                'task_no': task_no,
                'task_name': task_name,
                'total_count': total_count,
                'in_stock_count': in_stock_count,
                'assigned_count': assigned_count
            }
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/inventory/tasks', methods=['GET'])
def list_inventory_tasks():
    """盘点任务列表"""
    conn = get_conn()
    status = request.args.get('status', '')
    
    where_clause = ""
    params = []
    if status:
        where_clause = "WHERE status = ?"
        params.append(status)
    
    rows = conn.execute(
        f"SELECT * FROM inventory_tasks {where_clause} ORDER BY created_at DESC",
        params
    ).fetchall()
    conn.close()
    return jsonify({'success': True, 'data': rows_to_list(rows)})


@app.route('/api/inventory/tasks/<int:task_id>', methods=['GET'])
def get_inventory_task(task_id):
    """获取盘点任务详情"""
    conn = get_conn()
    task = conn.execute(
        "SELECT * FROM inventory_tasks WHERE id = ?", (task_id,)
    ).fetchone()
    
    if not task:
        conn.close()
        return jsonify({'success': False, 'error': '盘点任务不存在'}), 404
    
    # 获取盘点明细统计
    stats = conn.execute(
        """SELECT 
            COUNT(*) as total,
            SUM(CASE WHEN check_status = 'pending' THEN 1 ELSE 0 END) as pending,
            SUM(CASE WHEN check_status = 'normal' THEN 1 ELSE 0 END) as normal,
            SUM(CASE WHEN check_status = 'abnormal' THEN 1 ELSE 0 END) as abnormal,
            SUM(CASE WHEN check_status = 'missing' THEN 1 ELSE 0 END) as missing,
            SUM(CASE WHEN confirm_required = 1 THEN 1 ELSE 0 END) as need_confirm,
            SUM(CASE WHEN confirm_required = 1 AND check_status != 'pending' THEN 1 ELSE 0 END) as confirmed
        FROM inventory_items WHERE task_id = ?""",
        (task_id,)
    ).fetchone()
    
    conn.close()
    
    task_dict = dict(task)
    task_dict['stats'] = dict(stats) if stats else {}
    return jsonify({'success': True, 'task': task_dict})


@app.route('/api/inventory/tasks/<int:task_id>/items', methods=['GET'])
def list_inventory_items(task_id):
    """获取盘点明细列表"""
    conn = get_conn()
    check_status = request.args.get('check_status', '')
    confirm_required = request.args.get('confirm_required', '')
    
    where_clause = "WHERE task_id = ?"
    params = [task_id]
    
    if check_status:
        where_clause += " AND check_status = ?"
        params.append(check_status)
    if confirm_required:
        where_clause += " AND confirm_required = ?"
        params.append(1 if confirm_required == '1' else 0)
    
    rows = conn.execute(
        f"SELECT * FROM inventory_items {where_clause} ORDER BY id",
        params
    ).fetchall()
    conn.close()
    return jsonify({'success': True, 'data': rows_to_list(rows)})


@app.route('/api/inventory/items/<int:item_id>/check', methods=['POST'])
def check_inventory_item(item_id):
    """盘点单个资产（扫码或手动盘点）"""
    data = request.json or {}
    check_method = data.get('check_method', 'manual')  # manual/scan
    checked_by = data.get('checked_by', '')
    notes = data.get('notes', '')
    
    conn = get_conn()
    try:
        item = conn.execute(
            "SELECT * FROM inventory_items WHERE id = ?", (item_id,)
        ).fetchone()
        
        if not item:
            conn.close()
            return jsonify({'success': False, 'error': '盘点项不存在'}), 404
        
        # 检查资产当前状态
        asset = conn.execute(
            "SELECT * FROM assets WHERE id = ?", (item['asset_id'],)
        ).fetchone()
        
        if not asset:
            # 资产已被删除，标记为盘亏
            conn.execute(
                """UPDATE inventory_items SET 
                   check_status = 'missing', abnormal_type = 'missing', abnormal_desc = '资产已删除',
                   check_method = ?, checked_by = ?, checked_at = datetime('now','localtime'), notes = ?
                   WHERE id = ?""",
                (check_method, checked_by, notes, item_id)
            )
        else:
            # 比对资产信息
            abnormal_type = None
            abnormal_desc = None
            
            if asset['status'] != item['before_status']:
                abnormal_type = 'status_mismatch'
                abnormal_desc = f"状态不符: 盘点前[{item['before_status']}] 当前[{asset['status']}]"
            elif asset['location'] != item['before_location']:
                abnormal_type = 'location_mismatch'
                abnormal_desc = f"位置不符: 盘点前[{item['before_location']}] 当前[{asset['location']}]"
            
            if abnormal_type:
                conn.execute(
                    """UPDATE inventory_items SET 
                       check_status = 'abnormal', abnormal_type = ?, abnormal_desc = ?,
                       check_method = ?, checked_by = ?, checked_at = datetime('now','localtime'), notes = ?
                       WHERE id = ?""",
                    (abnormal_type, abnormal_desc, check_method, checked_by, notes, item_id)
                )
            else:
                conn.execute(
                    """UPDATE inventory_items SET 
                       check_status = 'normal',
                       check_method = ?, checked_by = ?, checked_at = datetime('now','localtime'), notes = ?
                       WHERE id = ?""",
                    (check_method, checked_by, notes, item_id)
                )
        
        # 更新任务统计
        conn.execute(
            """UPDATE inventory_tasks SET 
               checked_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id = ? AND check_status != 'pending'),
               normal_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id = ? AND check_status = 'normal'),
               abnormal_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id = ? AND check_status IN ('abnormal', 'missing'))
               WHERE id = ?""",
            (item['task_id'], item['task_id'], item['task_id'], item['task_id'])
        )
        
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/inventory/items/<int:item_id>/abnormal', methods=['POST'])
@login_required
def mark_inventory_abnormal(item_id):
    """管理员主动标记盘点异常或盘亏"""
    data = request.json or {}
    abnormal_type = (data.get('abnormal_type') or 'other').strip()
    abnormal_desc = (data.get('abnormal_desc') or '').strip()
    user = get_current_user() or {}
    checked_by = (data.get('checked_by') or user.get('name') or '').strip() or '管理员'
    if abnormal_type not in (
        'location_mismatch', 'status_mismatch', 'damaged', 'missing', 'other'
    ):
        return jsonify({'success': False, 'error': '无效的异常类型'}), 400
    if not abnormal_desc:
        return jsonify({'success': False, 'error': '请填写异常描述'}), 400

    check_status = 'missing' if abnormal_type == 'missing' else 'abnormal'
    conn = get_conn()
    try:
        item = conn.execute(
            "SELECT * FROM inventory_items WHERE id=?", (item_id,)
        ).fetchone()
        if not item:
            return jsonify({'success': False, 'error': '盘点项不存在'}), 404
        task = conn.execute(
            "SELECT status FROM inventory_tasks WHERE id=?", (item['task_id'],)
        ).fetchone()
        if not task or task['status'] not in ('ongoing', '进行中'):
            return jsonify({'success': False, 'error': '任务已结束，无法标记'}), 400
        if item['check_status'] not in (None, '', 'pending'):
            return jsonify({'success': False, 'error': '该资产已盘点'}), 400

        conn.execute(
            """UPDATE inventory_items SET
               check_status=?, abnormal_type=?, abnormal_desc=?,
               check_method='manual', checked_by=?, checked_at=datetime('now','localtime'), notes=?
               WHERE id=?""",
            (check_status, abnormal_type, abnormal_desc, checked_by, abnormal_desc, item_id)
        )
        conn.execute(
            """UPDATE inventory_tasks SET
               checked_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id = ? AND check_status != 'pending'),
               normal_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id = ? AND check_status = 'normal'),
               abnormal_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id = ? AND check_status IN ('abnormal', 'missing'))
               WHERE id = ?""",
            (item['task_id'], item['task_id'], item['task_id'], item['task_id'])
        )
        conn.commit()
        return jsonify({'success': True, 'check_status': check_status})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/inventory/items/<int:item_id>/confirm', methods=['POST'])
@login_required
def confirm_inventory_item(item_id):
    """责任人确认资产（自助盘点）"""
    data = request.json or {}
    user = get_current_user() or {}
    confirmed_by = (data.get('confirmed_by') or user.get('name') or '').strip()
    notes = data.get('notes', '')
    
    conn = get_conn()
    try:
        item = conn.execute(
            "SELECT * FROM inventory_items WHERE id = ? AND confirm_required = 1", (item_id,)
        ).fetchone()
        
        if not item:
            conn.close()
            return jsonify({'success': False, 'error': '盘点项不存在或不需要确认'}), 404
        
        # 验证确认人是否为责任人
        if confirmed_by != item['before_assignee']:
            conn.close()
            return jsonify({'success': False, 'error': '只有责任人可以确认'}), 403
        
        # 检查资产当前状态
        asset = conn.execute(
            "SELECT * FROM assets WHERE id = ?", (item['asset_id'],)
        ).fetchone()
        
        if not asset:
            conn.execute(
                """UPDATE inventory_items SET 
                   check_status = 'missing', abnormal_type = 'missing', abnormal_desc = '资产已删除',
                   check_method = 'self_confirm', checked_by = ?, checked_at = datetime('now','localtime'), notes = ?
                   WHERE id = ?""",
                (confirmed_by, notes, item_id)
            )
        else:
            abnormal_type = None
            abnormal_desc = None
            
            if asset['status'] != item['before_status']:
                abnormal_type = 'status_mismatch'
                abnormal_desc = f"状态不符: 盘点前[{item['before_status']}] 当前[{asset['status']}]"
            elif asset['location'] != item['before_location']:
                abnormal_type = 'location_mismatch'
                abnormal_desc = f"位置不符: 盘点前[{item['before_location']}] 当前[{asset['location']}]"
            
            if abnormal_type:
                conn.execute(
                    """UPDATE inventory_items SET 
                       check_status = 'abnormal', abnormal_type = ?, abnormal_desc = ?,
                       check_method = 'self_confirm', checked_by = ?, checked_at = datetime('now','localtime'), notes = ?
                       WHERE id = ?""",
                    (abnormal_type, abnormal_desc, confirmed_by, notes, item_id)
                )
            else:
                conn.execute(
                    """UPDATE inventory_items SET 
                       check_status = 'normal',
                       check_method = 'self_confirm', checked_by = ?, checked_at = datetime('now','localtime'), notes = ?
                       WHERE id = ?""",
                    (confirmed_by, notes, item_id)
                )
        
        # 更新任务统计
        conn.execute(
            """UPDATE inventory_tasks SET 
               checked_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id = ? AND check_status != 'pending'),
               normal_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id = ? AND check_status = 'normal'),
               abnormal_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id = ? AND check_status IN ('abnormal', 'missing'))
               WHERE id = ?""",
            (item['task_id'], item['task_id'], item['task_id'], item['task_id'])
        )
        
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/inventory/tasks/<int:task_id>/complete', methods=['POST'])
def complete_inventory_task(task_id):
    """完成盘点任务：未盘明细标为异常后结案"""
    conn = get_conn()
    try:
        task = conn.execute(
            "SELECT * FROM inventory_tasks WHERE id=?", (task_id,)
        ).fetchone()
        if not task:
            return jsonify({'success': False, 'error': '盘点任务不存在'}), 404
        if task['status'] in ('completed', '已完成'):
            return jsonify({'success': False, 'error': '任务已完成'}), 400

        # 与独立盘点页文案一致：未盘点的资产标记为异常
        conn.execute(
            """UPDATE inventory_items SET
               check_status='abnormal',
               abnormal_type='other',
               abnormal_desc=CASE
                 WHEN IFNULL(abnormal_desc,'')='' THEN '完成盘点时仍未盘点'
                 ELSE abnormal_desc
               END,
               check_method=COALESCE(NULLIF(check_method,''), 'auto'),
               checked_by=COALESCE(NULLIF(checked_by,''), '系统'),
               checked_at=COALESCE(checked_at, datetime('now','localtime'))
               WHERE task_id=? AND IFNULL(check_status,'pending')='pending'""",
            (task_id,)
        )
        conn.execute(
            """UPDATE inventory_tasks SET
               checked_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id=? AND check_status!='pending'),
               normal_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id=? AND check_status='normal'),
               abnormal_count = (SELECT COUNT(*) FROM inventory_items WHERE task_id=? AND check_status IN ('abnormal','missing')),
               status='completed',
               completed_at=datetime('now','localtime')
               WHERE id=?""",
            (task_id, task_id, task_id, task_id)
        )
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/inventory/tasks/<int:task_id>/report', methods=['GET'])
def get_inventory_report(task_id):
    """获取盘点报告"""
    conn = get_conn()
    task = conn.execute(
        "SELECT * FROM inventory_tasks WHERE id = ?", (task_id,)
    ).fetchone()
    
    if not task:
        conn.close()
        return jsonify({'success': False, 'error': '盘点任务不存在'}), 404
    
    # 统计信息
    stats = conn.execute(
        """SELECT 
            COUNT(*) as total,
            SUM(CASE WHEN check_status = 'pending' THEN 1 ELSE 0 END) as pending,
            SUM(CASE WHEN check_status = 'normal' THEN 1 ELSE 0 END) as normal,
            SUM(CASE WHEN check_status = 'abnormal' THEN 1 ELSE 0 END) as abnormal,
            SUM(CASE WHEN check_status = 'missing' THEN 1 ELSE 0 END) as missing
        FROM inventory_items WHERE task_id = ?""",
        (task_id,)
    ).fetchone()
    
    # 异常明细
    abnormal_items = conn.execute(
        """SELECT * FROM inventory_items 
        WHERE task_id = ? AND check_status IN ('abnormal', 'missing')
        ORDER BY abnormal_type""",
        (task_id,)
    ).fetchall()
    
    # 未盘点明细
    pending_items = conn.execute(
        """SELECT * FROM inventory_items 
        WHERE task_id = ? AND check_status = 'pending'
        ORDER BY confirm_required, asset_no""",
        (task_id,)
    ).fetchall()
    
    conn.close()
    
    return jsonify({
        'success': True,
        'task': dict(task),
        'stats': dict(stats) if stats else {},
        'abnormal_items': rows_to_list(abnormal_items),
        'pending_items': rows_to_list(pending_items)
    })


# ═══════════════════════════════════════════════════════════════
# 资产 CRUD
# ═══════════════════════════════════════════════════════════════

@app.route('/api/assets', methods=['GET'])
@login_required
def list_assets():
    conn = get_conn()
    # 权限隔离：普通用户只能查看「名下资产」（按 assignee 过滤）
    me = conn.execute(
        "SELECT name, role FROM employees WHERE id = ?", (request.current_user_id,)
    ).fetchone()
    is_admin = me and me['role'] in ADMIN_ROLES

    q = request.args.get('q', '').strip()
    category    = request.args.get('category', '')
    status      = request.args.get('status', '')
    location    = request.args.get('location', '')
    assignee    = request.args.get('assignee', '')
    asset_class = request.args.get('asset_class', '')
    license_type = request.args.get('license_type', '')
    page      = max(1, int(request.args.get('page', 1)))
    page_size = int(request.args.get('page_size', 20))
    scope     = request.args.get('scope', '')

    sql    = "SELECT * FROM assets WHERE 1=1"
    params = []
    # 出入库/盘点等操作场景需看到全量资产，允许前端通过 scope=all 跳过权限隔离
    if scope != 'all' and not is_admin and me:
        # 强制只能看自己名下的资产，忽略前端传入的 assignee 参数
        # 防御重名：用 assignee_id 强匹配 + assignee 姓名匹配，覆盖老数据中只填 id 或只填姓名的资产
        sql += " AND (assignee = ? OR assignee_id = ?)"; params.extend([me['name'], request.current_user_id])
    if q:
        clause, p = fuzzy_match_clause(q, ['asset_no', 'brand', 'model', 'sn', 'assignee'])
        if clause:
            sql += " AND " + clause
            params += p
    if category:
        clause, cparams = category_filter_clause(conn, category)
        if clause:
            sql += " AND " + clause
            params += cparams
    if status:
        sql += " AND status=?";   params.append(status)
    if location:
        sql += " AND location=?"; params.append(location)
    if assignee and is_admin:
        sql += " AND assignee LIKE ?"; params.append(f'%{assignee}%')
    if asset_class:
        sql += " AND asset_class=?"; params.append(asset_class)
    if license_type:
        sql += " AND license_type=?"; params.append(license_type)

    total = conn.execute(
        sql.replace("SELECT *", "SELECT COUNT(*) as cnt"), params
    ).fetchone()['cnt']

    sql += " ORDER BY id DESC LIMIT ? OFFSET ?"
    params += [page_size, (page-1)*page_size]
    rows = conn.execute(sql, params).fetchall()
    data = enrich_assets_book_value(assets_to_safe_list(rows), conn)
    conn.close()
    return jsonify({'total': total, 'page': page, 'page_size': page_size,
                    'data': data})


@app.route('/api/assets/<int:aid>', methods=['GET'])
def get_asset(aid):
    conn = get_conn()
    row = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'error': '资产不存在'}), 404
    data = enrich_assets_book_value([strip_sensitive_asset(dict(row))], conn)[0]
    conn.close()
    return jsonify(data)


@app.route('/api/assets/category-stats', methods=['GET'])
@login_required
def assets_category_stats():
    """按资产类型统计在库数量（普通用户仅统计自己名下）"""
    status = request.args.get('status', '在库')
    conn = get_conn()

    # 权限隔离：普通用户仅统计自己名下的资产；前端通过 scope=all 跳过隔离（出入库等场景需要全量）
    me = conn.execute(
        "SELECT name, role FROM employees WHERE id = ?", (request.current_user_id,)
    ).fetchone()
    is_admin = bool(me and me['role'] in ADMIN_ROLES)
    scope = request.args.get('scope', '')
    apply_user_scope = (not is_admin) and (scope != 'all')
    me_name = me['name'] if me else ''

    # 获取所有资产类型
    categories = conn.execute(
        'SELECT field_code, field_name FROM field_options WHERE field_type="category" AND active=1 ORDER BY sort_order'
    ).fetchall()

    # 统计每种类型的数量（兼容 category 存中文名、asset_type_code 存编码）
    stats = []
    for cat in categories:
        clause, cparams = category_filter_clause(conn, cat['field_code'])
        if apply_user_scope:
            count = conn.execute(
                f'SELECT COUNT(*) as cnt FROM assets WHERE status=? AND assignee=? AND {clause}',
                [status, me_name] + cparams
            ).fetchone()['cnt']
        else:
            count = conn.execute(
                f'SELECT COUNT(*) as cnt FROM assets WHERE status=? AND {clause}',
                [status] + cparams
            ).fetchone()['cnt']
        stats.append({
            'code': cat['field_code'],
            'name': cat['field_name'],
            'count': count
        })

    conn.close()
    return jsonify({'data': stats})


@app.route('/api/assets', methods=['POST'])
def create_asset():
    data = request.json or {}
    if not data.get('category'):
        return jsonify({'error': '类目不能为空'}), 400

    # 新资产编号规则: YYYYMMDD001-主体-地区-类型
    subject_code = data.get('subject_code', '')
    region_code = data.get('region_code', '')
    conn = get_conn()
    asset_type_code = resolve_asset_type_code(
        conn, data.get('category', ''), data.get('asset_type_code', ''))
    
    if data.get('asset_no'):
        # 手动指定编号时，保持原有逻辑
        asset_no = data['asset_no']
    elif subject_code and region_code and asset_type_code:
        # 新规则
        asset_no = gen_asset_no(subject_code, region_code, asset_type_code)
    else:
        # 兼容旧逻辑
        asset_no = gen_asset_no(data['category'], '', '')
    
    order_id = data.get('order_id') or None
    order_item_id = data.get('order_item_id') or None

    # 关联账号密码（可选项）
    account_enabled = bool(data.get('account_enabled', False))
    account_name = (data.get('account_name','') or '') if account_enabled else ''
    account_password = (data.get('account_password','') or '') if account_enabled else ''
    account_pwd = (account_name + ' / ' + account_password) if (account_name or account_password) else ''

    try:
        cur = conn.execute('''
            INSERT INTO assets
            (asset_no,category,brand,model,spec,sn,mac_wired,mac_wireless,
             status,location,assignee,key_info,account_pwd,account_name,account_password,purchase_date,
             warranty_expire,purchase_price,notes,order_id,order_item_id,
             subject_code,region_code,asset_type_code,
             asset_class,license_type,expire_date,seat_count,renew_remind_days)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        ''', (
            asset_no,
            data.get('category',''),
            data.get('brand',''),
            data.get('model',''),
            data.get('spec',''),
            data.get('sn',''),
            data.get('mac_wired',''),
            data.get('mac_wireless',''),
            data.get('status','在库'),
            data.get('location',''),
            data.get('assignee',''),
            data.get('key_info',''),
            account_pwd,
            account_name,
            account_password,
            data.get('purchase_date',''),
            data.get('warranty_expire',''),
            data.get('purchase_price') or None,
            data.get('notes',''),
            order_id,
            order_item_id,
            subject_code,
            region_code,
            asset_type_code,
            data.get('asset_class', 'hardware'),
            data.get('license_type') or None,
            data.get('expire_date') or None,
            data.get('seat_count') or None,
            data.get('renew_remind_days', 30),
        ))
        new_id = cur.lastrowid

        # 自动写入库记录
        conn.execute('''
            INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,
                                   to_location,target_user,remark,order_id)
            VALUES(?,?,?,?,?,?,?,?)
        ''', (new_id, asset_no, '入库',
              get_current_name(),
              data.get('location',''),
              data.get('assignee',''),
              data.get('notes',''),
              order_id))

        # 如果关联了订单明细，更新已到货数量
        if order_item_id:
            conn.execute(
                "UPDATE order_items SET arrived_qty = arrived_qty + 1 WHERE id = ?",
                (order_item_id,)
            )
            # 更新订单状态
            _refresh_order_status(conn, order_id)

        conn.commit()

        log_action('create', 'asset', 'asset', new_id, asset_no, data.get('category',''),
                   get_current_name() or '系统')

        notify_lark('stock',
            f"📦 【资产入库】\n编号：{asset_no}\n类目：{data.get('category')}\n品牌型号：{data.get('brand','')} {data.get('model','')}\n存放地：{data.get('location','')}\n关联人：{data.get('assignee','')}\n订单：{data.get('order_no','无')}\n时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({'error': f'资产编号 {asset_no} 已存在'}), 409
    conn.close()
    return jsonify({'id': new_id, 'asset_no': asset_no}), 201


@app.route('/api/assets/<int:aid>', methods=['PUT'])
def update_asset(aid):
    data = request.json or {}
    conn = get_conn()
    old = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not old:
        conn.close()
        return jsonify({'error': '资产不存在'}), 404

    fields = ['category','brand','model','spec','sn','mac_wired','mac_wireless',
              'status','location','assignee','key_info','account_pwd','account_name','account_password',
              'purchase_date','warranty_expire','purchase_price','notes',
              'asset_class','license_type','expire_date','seat_count','renew_remind_days',
              'asset_type_code','subject_code','region_code']
    updates = []
    params  = []
    for f in fields:
        if f in data:
            updates.append(f'{f}=?')
            params.append(data[f])
    # 类目变更或未传类型码时，按类目回填
    new_category = data.get('category', old['category'] if 'category' in old.keys() else '')
    if 'asset_type_code' not in data or not data.get('asset_type_code'):
        resolved = resolve_asset_type_code(conn, new_category, data.get('asset_type_code', '') or (old['asset_type_code'] if 'asset_type_code' in old.keys() else ''))
        if resolved and resolved != (old['asset_type_code'] if 'asset_type_code' in old.keys() else None):
            if 'asset_type_code=?' not in updates:
                updates.append('asset_type_code=?')
                params.append(resolved)
            else:
                # 已在上面追加了空值，替换最后一个对应值
                idx = updates.index('asset_type_code=?')
                # params 与 updates 顺序一致，找到该字段对应的 params 位置
                # updates 中在 asset_type_code 之前的字段数 = idx
                params[idx] = resolved
    # 同步 account_pwd 合并字段
    if 'account_name' in data or 'account_password' in data:
        new_name = data.get('account_name', old['account_name'] or '') or ''
        new_pwd = data.get('account_password', old['account_password'] or '') or ''
        account_pwd_combined = (new_name + ' / ' + new_pwd) if (new_name or new_pwd) else ''
        updates.append('account_pwd=?')
        params.append(account_pwd_combined)
    updates.append("updated_at=datetime('now','localtime')")
    params.append(aid)

    conn.execute(f"UPDATE assets SET {','.join(updates)} WHERE id=?", params)
    conn.commit()
    log_action('update', 'asset', 'asset', aid, old['asset_no'], old['category'] or '',
               get_current_user()['name'] if get_current_user() else '系统')
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/assets/<int:aid>', methods=['DELETE'])
@superadmin_required
def delete_asset(aid):
    """直接删除资产（仅超管）。
    安全加固（三轮审计#2）：原为无鉴权接口，操作员/财务可绕过 delete-applications 审批直接删资产；
    现收紧为仅超管可直删，其余管理角色的删除须走 delete-applications 审批流。"""
    conn = get_conn()
    old = conn.execute("SELECT asset_no, category FROM assets WHERE id=?", (aid,)).fetchone()
    conn.execute("DELETE FROM assets WHERE id=?", (aid,))
    conn.commit()
    if old:
        log_action('delete', 'asset', 'asset', aid, old['asset_no'], old['category'] or '',
                   get_current_user()['name'] if get_current_user() else '系统')
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/expire-reminder', methods=['GET'])
@login_required
def get_expire_reminder():
    """续期提醒：返回订阅制资产的到期状态（普通用户仅看自己名下）"""
    conn = get_conn()
    me = conn.execute(
        "SELECT name, role FROM employees WHERE id = ?", (request.current_user_id,)
    ).fetchone()
    is_admin = bool(me and me['role'] in ADMIN_ROLES)
    me_name = me['name'] if me else ''

    today = datetime.now().strftime('%Y-%m-%d')

    # 查出所有订阅制且有到期日的资产（普通用户只看自己名下）
    if is_admin:
        rows = conn.execute("""
            SELECT id, asset_no, category, brand, model, assignee,
                   expire_date, seat_count, renew_remind_days, status, notes
            FROM assets
            WHERE license_type = 'subscription'
              AND expire_date IS NOT NULL AND expire_date != ''
            ORDER BY expire_date ASC
        """).fetchall()
    else:
        rows = conn.execute("""
            SELECT id, asset_no, category, brand, model, assignee,
                   expire_date, seat_count, renew_remind_days, status, notes
            FROM assets
            WHERE license_type = 'subscription'
              AND expire_date IS NOT NULL AND expire_date != ''
              AND assignee = ?
            ORDER BY expire_date ASC
        """, (me_name,)).fetchall()
    conn.close()

    result = []
    for r in rows:
        r = dict(r)
        expire = r['expire_date']
        remind_days = r.get('renew_remind_days') or 30
        # 计算剩余天数
        try:
            from datetime import date
            delta = (date.fromisoformat(expire) - date.fromisoformat(today)).days
            r['days_left'] = delta
            if delta < 0:
                r['expire_status'] = 'expired'     # 已过期
            elif delta <= remind_days:
                r['expire_status'] = 'warning'     # 即将到期
            else:
                r['expire_status'] = 'normal'      # 正常
        except Exception:
            r['days_left'] = None
            r['expire_status'] = 'unknown'
        result.append(r)

    # 分组统计
    expired = [x for x in result if x['expire_status'] == 'expired']
    warning = [x for x in result if x['expire_status'] == 'warning']
    normal  = [x for x in result if x['expire_status'] == 'normal']

    return jsonify({
        'today': today,
        'total': len(result),
        'expired_count': len(expired),
        'warning_count': len(warning),
        'normal_count': len(normal),
        'data': result
    })


@app.route('/api/alerts/export', methods=['GET'])
@login_required
def export_alerts():
    """导出到期订阅与低库存耗材（管理角色）"""
    user = get_current_user()
    if not user or user.get('role') not in ADMIN_ROLES:
        return jsonify({'success': False, 'error': '需要管理权限', 'code': 403}), 403
    from openpyxl import Workbook
    from datetime import date as _date
    conn = get_conn()
    try:
        today = datetime.now().strftime('%Y-%m-%d')
        rows = conn.execute("""
            SELECT id, asset_no, category, brand, model, assignee,
                   expire_date, seat_count, renew_remind_days, status, notes
            FROM assets
            WHERE license_type = 'subscription'
              AND expire_date IS NOT NULL AND expire_date != ''
            ORDER BY expire_date ASC
        """).fetchall()
        wb = Workbook()
        ws1 = wb.active
        ws1.title = '订阅到期'
        ws1.append(['资产编号', '类目', '品牌', '型号', '关联人', '到期日', '剩余天数', '状态'])
        for r in rows:
            expire = r['expire_date']
            remind_days = r['renew_remind_days'] or 30
            try:
                delta = (_date.fromisoformat(str(expire)[:10]) - _date.fromisoformat(today)).days
                if delta < 0:
                    st = '已过期'
                elif delta <= remind_days:
                    st = '即将到期'
                else:
                    st = '正常'
            except Exception:
                delta = None
                st = '未知'
            if st == '正常':
                continue
            ws1.append([r['asset_no'], r['category'], r['brand'], r['model'], r['assignee'],
                        expire, delta, st])
        ws2 = wb.create_sheet('低库存耗材')
        ws2.append(['编号', '名称', '类目', '库存', '预警阈值', '状态', '存放地'])
        low = conn.execute('''
            SELECT * FROM consumables
            WHERE status = '低库存' OR quantity < min_stock
            ORDER BY (min_stock - quantity) DESC
        ''').fetchall()
        for r in low:
            ws2.append([r['item_no'], r['name'], r['category'], r['quantity'],
                        r['min_stock'], r['status'], r['location']])
        ws3 = wb.create_sheet('保修到期')
        ws3.append(['资产编号', '类目', '品牌', '型号', '关联人', '保修到期日', '剩余天数', '状态'])
        warranty_rows = conn.execute("""
            SELECT asset_no, category, brand, model, assignee, warranty_expire, status
            FROM assets
            WHERE warranty_expire IS NOT NULL AND warranty_expire != ''
            ORDER BY warranty_expire ASC
        """).fetchall()
        for r in warranty_rows:
            wdate = r['warranty_expire']
            try:
                delta = (_date.fromisoformat(str(wdate)[:10]) - _date.fromisoformat(today)).days
                if delta < 0:
                    st = '已过期'
                elif delta <= 30:
                    st = '即将到期'
                else:
                    st = '正常'
            except Exception:
                delta = None
                st = '未知'
            if st == '正常':
                continue
            ws3.append([r['asset_no'], r['category'], r['brand'], r['model'], r['assignee'],
                        wdate, delta, st])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name='到期与低库存提醒.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    finally:
        conn.close()


# ─────────────────────────────────────────────
# 删除申请与审批
# ─────────────────────────────────────────────

@app.route('/api/delete-applications', methods=['POST'])
def create_delete_application():
    """提交删除申请"""
    data = request.json or {}
    target_type = data.get('target_type')  # asset/consumable
    target_id = data.get('target_id')
    reason = data.get('reason', '')
    apply_by = data.get('apply_by', '未知用户')
    
    if not target_type or not target_id:
        return jsonify({'error': '缺少必要参数'}), 400
    
    conn = get_conn()
    
    # 获取目标信息
    if target_type == 'asset':
        target = conn.execute("SELECT * FROM assets WHERE id=?", (target_id,)).fetchone()
        if not target:
            conn.close()
            return jsonify({'error': '资产不存在'}), 404
        target_no = target['asset_no']
        target_info = json.dumps({
            'category': target['category'],
            'brand': target['brand'],
            'model': target['model'],
            'status': target['status'],
            'assignee': target['assignee']
        })
        # 更新资产状态为待删除
        conn.execute("UPDATE assets SET status='待删除', updated_at=datetime('now','localtime') WHERE id=?", (target_id,))
    elif target_type == 'consumable':
        target = conn.execute("SELECT * FROM consumables WHERE id=?", (target_id,)).fetchone()
        if not target:
            conn.close()
            return jsonify({'error': '耗材不存在'}), 404
        target_no = target['item_no']
        target_info = json.dumps({
            'name': target['name'],
            'category': target['category'],
            'quantity': target['quantity']
        })
    else:
        conn.close()
        return jsonify({'error': '不支持的目标类型'}), 400
    
    # 创建申请记录
    cursor = conn.execute('''
        INSERT INTO delete_applications (target_type, target_id, target_no, target_info, apply_by, reason)
        VALUES (?, ?, ?, ?, ?, ?)
    ''', (target_type, target_id, target_no, target_info, apply_by, reason))
    app_id = cursor.lastrowid
    conn.commit()
    conn.close()

    try:
        notify_lark_all(
            f"🗑️ 【删除申请待审批】\n"
            f"类型：{'固定资产' if target_type == 'asset' else '耗材'}\n"
            f"编号：{target_no}\n"
            f"申请人：{apply_by}\n"
            f"原因：{reason or '-'}\n"
            f"时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
        )
    except Exception as _e:
        print(f"[Lark] 删除申请通知失败(已隔离): {_e}")

    return jsonify({'ok': True, 'id': app_id})


@app.route('/api/delete-applications', methods=['GET'])
def list_delete_applications():
    """获取删除申请列表"""
    status = request.args.get('status', '')  # pending/approved/rejected
    conn = get_conn()
    sql = "SELECT * FROM delete_applications WHERE 1=1"
    params = []
    if status:
        sql += " AND status=?"
        params.append(status)
    sql += " ORDER BY apply_time DESC"
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return jsonify({'data': rows_to_list(rows)})


@app.route('/api/delete-applications/<int:app_id>/approve', methods=['POST'])
@approver_required
def approve_delete_application(app_id):
    """审批删除申请（超管 / 管理员 / 普通管理员）"""
    data = request.json or {}
    action = data.get('action')  # approve/reject
    # 审批人取当前登录用户，拒绝前端传值伪造
    approver = request.current_user['name']
    note = data.get('note', '')
    
    if action not in ('approve', 'reject'):
        return jsonify({'error': '无效的审批操作'}), 400
    
    conn = get_conn()
    app = conn.execute("SELECT * FROM delete_applications WHERE id=?", (app_id,)).fetchone()
    if not app:
        conn.close()
        return jsonify({'error': '申请不存在'}), 404
    if app['status'] != 'pending':
        conn.close()
        return jsonify({'error': '该申请已处理'}), 400
    
    new_status = 'approved' if action == 'approve' else 'rejected'
    conn.execute('''
        UPDATE delete_applications 
        SET status=?, approver=?, approve_time=datetime('now','localtime'), approve_note=?
        WHERE id=?
    ''', (new_status, approver, note, app_id))
    
    if action == 'approve':
        # 审批通过，执行删除
        if app['target_type'] == 'asset':
            conn.execute("DELETE FROM assets WHERE id=?", (app['target_id'],))
        elif app['target_type'] == 'consumable':
            conn.execute("DELETE FROM consumables WHERE id=?", (app['target_id'],))
    else:
        # 审批拒绝，恢复资产状态
        if app['target_type'] == 'asset':
            conn.execute("UPDATE assets SET status='在库', updated_at=datetime('now','localtime') WHERE id=?", (app['target_id'],))
    
    conn.commit()
    conn.close()

    try:
        action_zh = '已通过' if action == 'approve' else '已驳回'
        notify_lark_all(
            f"{'✅' if action == 'approve' else '❌'} 【删除申请{action_zh}】\n"
            f"编号：{app['target_no']}\n"
            f"申请人：{app['apply_by']}\n"
            f"审批人：{approver}\n"
            f"备注：{note or '-'}\n"
            f"时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
        )
    except Exception as _e:
        print(f"[Lark] 删除审批通知失败(已隔离): {_e}")

    return jsonify({'ok': True})


# ════════════════════════════════════════════════════════
# 采购入库审批 / Excel 模板链路
# ════════════════════════════════════════════════════════
import json as _json
from openpyxl import Workbook, load_workbook
import os as _os

_EXPORT_DIR = _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), 'export_tmp')
_os.makedirs(_EXPORT_DIR, exist_ok=True)


def _map_status_zh(s):
    return {'pending': '待审批', 'approved': '已通过', 'rejected': '已驳回',
            'reviewing': '审核中', 'processing': '审核中'}.get(s, s)


def _normalize_approval_status(status):
    """前端可能传中文或英文；库内统一用英文过滤。"""
    s = (status or '').strip()
    if not s:
        return ''
    return {
        '待审批': 'pending', '已通过': 'approved', '已驳回': 'rejected',
        '已拒绝': 'rejected', '审核中': 'processing',
        'pending': 'pending', 'approved': 'approved', 'rejected': 'rejected',
        'processing': 'processing', 'reviewing': 'reviewing',
    }.get(s, s)


def _count_pending_approvals(conn):
    """与审批中心列表一致：资产申请 + 操作单（处置/出入库/维修等）+ 删除 + 采购入库。"""
    total = 0
    for sql in (
        "SELECT COUNT(*) AS c FROM asset_applications WHERE status='pending'",
        "SELECT COUNT(*) AS c FROM operation_applications WHERE status='pending'",
        "SELECT COUNT(*) AS c FROM delete_applications WHERE status='pending'",
        "SELECT COUNT(*) AS c FROM stock_in_applications WHERE status='pending'",
    ):
        total += conn.execute(sql).fetchone()['c'] or 0
    return total


@app.route('/api/approvals', methods=['GET'])
@login_required
def list_approvals():
    """聚合审批列表：采购入库 / 资产删除 / 资产申请"""
    status = _normalize_approval_status(request.args.get('status', ''))
    apply_type = request.args.get('apply_type', '')
    page = int(request.args.get('page', 1))
    page_size = int(request.args.get('page_size', 20))
    offset = (page - 1) * page_size
    conn = get_conn()
    try:
        rows = []
        if not apply_type or apply_type == '采购入库':
            q = "SELECT * FROM stock_in_applications WHERE 1=1"
            p = []
            if status: q += " AND status=?"; p.append(status)
            q += " ORDER BY id DESC"
            for r in conn.execute(q, p).fetchall():
                d = dict(r)
                rows.append({
                    'id': d['id'], 'apply_type': '采购入库',
                    'title': f"采购入库 - 订单 {d['order_no']}",
                    'apply_by': d['apply_by'], 'applicant': d['apply_by'],
                    'status': _map_status_zh(d['status']),
                    'raw_status': d['status'],
                    'created_at': d['apply_time'], 'apply_time': d['apply_time'],
                    'apply_no': d.get('apply_no') or str(d['id']),
                    'asset_info': d['order_no'],
                    'approver': d.get('reviewed_by') or '',
                    'info': {'order_id': d['order_id'], 'order_no': d['order_no'],
                             'asset_ids': _json.loads(d['asset_ids']) if d['asset_ids'] else []}
                })
        if not apply_type or apply_type == '资产删除':
            q = "SELECT * FROM delete_applications WHERE 1=1"
            p = []
            if status: q += " AND status=?"; p.append(status)
            q += " ORDER BY id DESC"
            for r in conn.execute(q, p).fetchall():
                d = dict(r)
                rows.append({
                    'id': d['id'], 'apply_type': '资产删除',
                    'title': f"删除 {d['target_no']}",
                    'apply_by': d['apply_by'], 'applicant': d['apply_by'],
                    'status': _map_status_zh(d['status']),
                    'raw_status': d['status'],
                    'created_at': d['apply_time'], 'apply_time': d['apply_time'],
                    'apply_no': str(d['id']),
                    'asset_info': d['target_no'],
                    'approver': d.get('approver') or '',
                    'info': {'target_type': d['target_type'], 'target_no': d['target_no'], 'reason': d['reason']}
                })
        if not apply_type or apply_type == '资产申请':
            q = "SELECT * FROM asset_applications WHERE 1=1"
            p = []
            if status: q += " AND status=?"; p.append(status)
            q += " ORDER BY id DESC"
            for r in conn.execute(q, p).fetchall():
                d = dict(r)
                type_zh = _APP_TYPE_LABELS.get(d.get('app_type'), d.get('app_type') or '申请')
                detail_raw = d.get('detail')
                try:
                    detail_obj = _json.loads(detail_raw) if isinstance(detail_raw, str) and detail_raw else (detail_raw or {})
                except Exception:
                    detail_obj = {}
                if not isinstance(detail_obj, dict):
                    detail_obj = {}
                info = dict(d)
                info['detail'] = detail_obj
                info['app_type_zh'] = type_zh
                rows.append({
                    'id': d['id'], 'apply_type': '资产申请',
                    'title': f"{type_zh} - {d.get('applicant_name','')}",
                    'apply_by': d.get('applicant_name', ''), 'applicant': d.get('applicant_name', ''),
                    'status': _map_status_zh(d['status']),
                    'raw_status': d['status'],
                    'created_at': d.get('created_at') or d.get('apply_time') or '',
                    'apply_time': d.get('apply_time') or d.get('created_at') or '',
                    'apply_no': d.get('app_no') or str(d['id']),
                    'asset_info': d.get('asset_no') or d.get('title') or type_zh,
                    'asset_no': d.get('asset_no') or '',
                    'approver': d.get('reviewer_name') or '',
                    'info': info
                })
        if not apply_type or apply_type in OP_TYPE_ZH.values():
            q = "SELECT * FROM operation_applications WHERE 1=1"
            p = []
            if status: q += " AND status=?"; p.append(status)
            if apply_type:
                target_op = [k for k, v in OP_TYPE_ZH.items() if v == apply_type]
                if target_op:
                    q += " AND op_type=?"; p.append(target_op[0])
            q += " ORDER BY id DESC"
            for r in conn.execute(q, p).fetchall():
                d = dict(r)
                op_type = d['op_type']
                asset_nos = _json.loads(d['asset_nos']) if d['asset_nos'] else []
                payload = _json.loads(d['payload']) if d['payload'] else {}
                proof = rows_to_list(conn.execute(
                    "SELECT id, file_name FROM photos WHERE ref_type='operation_proof' AND ref_id=?", (d['id'],)))
                type_zh = OP_TYPE_ZH.get(op_type, op_type)
                nos_text = '/'.join(asset_nos) if asset_nos else '资产'
                rows.append({
                    'id': d['id'], 'apply_type': type_zh,
                    'title': f"{type_zh} - {nos_text}",
                    'apply_by': d['applicant'], 'applicant': d['applicant'],
                    'status': _map_status_zh(d['status']),
                    'raw_status': d['status'],
                    'created_at': d['apply_time'], 'apply_time': d['apply_time'],
                    'apply_no': str(d['id']),
                    'asset_info': nos_text,
                    'asset_no': asset_nos[0] if asset_nos else '',
                    'approver': d.get('reviewer') or '',
                    'info': {'op_type': op_type, 'asset_nos': asset_nos, 'payload': payload, 'proof_photos': proof}
                })
        total = len(rows)
        return jsonify({'success': True, 'data': rows[offset:offset + page_size], 'total': total})
    finally:
        conn.close()


@app.route('/api/approvals/<int:app_id>/approve', methods=['POST'])
@approver_required
def approve_aggregated(app_id):
    data = request.json or {}
    apply_type = data.get('apply_type')
    action = 'approve' if data.get('status') == '已通过' else 'reject'
    # 审批角色：超管 / 管理员 / 普通管理员；审批人取当前登录用户，拒绝前端传值伪造
    approver = request.current_user['name']
    note = data.get('remark', '') or data.get('approve_note', '')
    conn = get_conn()
    try:
        if apply_type == '资产删除':
            appr = conn.execute("SELECT * FROM delete_applications WHERE id=?", (app_id,)).fetchone()
            if not appr or appr['status'] != 'pending':
                return jsonify({'success': False, 'error': '申请不存在或已处理'}), 400
            new_status = 'approved' if action == 'approve' else 'rejected'
            conn.execute("UPDATE delete_applications SET status=?, approver=?, approve_time=datetime('now','localtime'), approve_note=? WHERE id=?",
                         (new_status, approver, note, app_id))
            if action == 'approve':
                if appr['target_type'] == 'asset':
                    conn.execute("DELETE FROM assets WHERE id=?", (appr['target_id'],))
                elif appr['target_type'] == 'consumable':
                    conn.execute("DELETE FROM consumables WHERE id=?", (appr['target_id'],))
            else:
                if appr['target_type'] == 'asset':
                    conn.execute("UPDATE assets SET status='在库', updated_at=datetime('now','localtime') WHERE id=?", (appr['target_id'],))
        elif apply_type == '资产申请':
            # 原子抢占 pending → processing，避免并发重复执行借用/扣库存
            cur = conn.execute(
                "UPDATE asset_applications SET status='processing' WHERE id=? AND status='pending'",
                (app_id,))
            if cur.rowcount == 0:
                return jsonify({'success': False, 'error': '申请不存在或已处理'}), 400
            appr = conn.execute("SELECT * FROM asset_applications WHERE id=?", (app_id,)).fetchone()
            new_status = 'approved' if action == 'approve' else 'rejected'
            if action == 'approve':
                try:
                    execute_asset_application(conn, dict(appr))
                except AssetAppExecError as e:
                    conn.execute("UPDATE asset_applications SET status='pending' WHERE id=?", (app_id,))
                    conn.commit()
                    return jsonify({'success': False, 'error': str(e)}), 400
            conn.execute("UPDATE asset_applications SET status=?, reviewer_id=?, reviewer_name=?, review_time=datetime('now','localtime'), review_comment=? WHERE id=?",
                         (new_status, request.current_user['id'], approver, note, app_id))
        elif apply_type == '采购入库':
            # 安全加固（P1-2）：原子抢占，避免并发下重复插入入库日志/重复置已到货
            cur = conn.execute(
                "UPDATE stock_in_applications SET status='processing' WHERE id=? AND status='pending'",
                (app_id,))
            if cur.rowcount == 0:
                return jsonify({'success': False, 'error': '申请不存在或已处理'}), 400
            appr = conn.execute("SELECT * FROM stock_in_applications WHERE id=?", (app_id,)).fetchone()
            ids = _json.loads(appr['asset_ids']) if appr['asset_ids'] else []
            if action == 'approve':
                now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                for aid in ids:
                    conn.execute("UPDATE assets SET status='在库', updated_at=datetime('now','localtime') WHERE id=?", (aid,))
                    a = conn.execute("SELECT asset_no FROM assets WHERE id=?", (aid,)).fetchone()
                    if a:
                        conn.execute('''INSERT INTO stock_logs (asset_id, asset_no, log_type, operator, remark, log_time, order_id)
                                        VALUES (?,?,?,?,?,?,?)''',
                                     (aid, a['asset_no'], '入库', approver, f"采购入库审批通过 [{appr['order_no']}]", now, appr['order_id']))
                conn.execute("UPDATE purchase_orders SET status='已到货', updated_at=datetime('now','localtime') WHERE id=?", (appr['order_id'],))
                new_status = 'approved'
            else:
                for aid in ids:
                    conn.execute("DELETE FROM assets WHERE id=? AND status='待审批入库'", (aid,))
                new_status = 'rejected'
            conn.execute("UPDATE stock_in_applications SET status=?, reviewed_by=?, reviewed_at=datetime('now','localtime'), remark=? WHERE id=?",
                         (new_status, approver, note, app_id))
        elif apply_type in OP_TYPE_ZH.values():
            # 安全加固（P1-2）：原子抢占——先原子地把 pending 置为 processing 抢占处理权，
            # 抢占失败（rowcount=0）说明已被其他并发请求处理，直接拒绝，
            # 避免「先查 pending 后执行」的 TOCTOU 窗口导致出库/维修/处置被并发重复执行。
            cur = conn.execute(
                "UPDATE operation_applications SET status='processing' WHERE id=? AND status='pending'",
                (app_id,))
            if cur.rowcount == 0:
                return jsonify({'success': False, 'error': '申请不存在或已处理'}), 400
            op = conn.execute("SELECT * FROM operation_applications WHERE id=?", (app_id,)).fetchone()
            if action == 'approve':
                ok = execute_operation(dict(op))
                if not ok:
                    # 执行失败：回退为 pending，允许后续重试，不遗留 processing 僵死状态
                    conn.execute("UPDATE operation_applications SET status='pending' WHERE id=?", (app_id,))
                    conn.commit()
                    return jsonify({'success': False, 'error': '执行操作失败，请查看服务端日志'}), 500
                new_status = 'approved'
            else:
                new_status = 'rejected'
            conn.execute("UPDATE operation_applications SET status=?, reviewer=?, review_time=datetime('now','localtime'), remark=? WHERE id=?",
                         (new_status, approver, note, app_id))
            try:
                log_action('approve' if action == 'approve' else 'reject', 'approval', 'operation_application', app_id,
                           op['applicant'], f"{OP_TYPE_ZH.get(op['op_type'], op['op_type'])}申请", {'action': action})
            except Exception:
                pass
        else:
            return jsonify({'success': False, 'error': '未知审批类型'}), 400
        conn.commit()
        try:
            if apply_type == '采购入库':
                log_action('approve' if action == 'approve' else 'reject', 'approval', 'stock_in_application', app_id, appr['order_no'], f"采购入库-订单{appr['order_no']}", {'count': len(ids)})
            elif apply_type == '资产删除':
                log_action('approve' if action == 'approve' else 'reject', 'approval', 'delete_application', app_id, appr['target_no'], appr['target_no'], {'action': action})
                action_zh = '已通过' if action == 'approve' else '已驳回'
                notify_lark_all(
                    f"{'✅' if action == 'approve' else '❌'} 【删除申请{action_zh}】\n"
                    f"编号：{appr['target_no']}\n"
                    f"申请人：{appr['apply_by']}\n"
                    f"审批人：{approver}\n"
                    f"备注：{note or '-'}\n"
                    f"时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
                )
            elif apply_type == '资产申请':
                log_action('approve' if action == 'approve' else 'reject', 'approval', 'asset_application', app_id, '', appr.get('applicant_name', ''), {'action': action})
        except Exception:
            pass
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/orders/<int:order_id>/complete', methods=['POST'])
@login_required
def order_complete_gen_template(order_id):
    """采购完成：按订单 asset 明细数量生成 Excel 入库模板（预填资产编号）"""
    conn = get_conn()
    try:
        order = conn.execute("SELECT * FROM purchase_orders WHERE id=?", (order_id,)).fetchone()
        if not order:
            return jsonify({'success': False, 'error': '订单不存在'}), 404
        order = dict(order)
        if order['status'] in ('已取消',):
            return jsonify({'success': False, 'error': '已取消订单无法操作'}), 400
        items = conn.execute("SELECT * FROM order_items WHERE order_id=? AND item_type='asset'", (order_id,)).fetchall()
        type_map = {}
        for r in conn.execute("SELECT code, name FROM asset_types WHERE active=1").fetchall():
            type_map[r['name']] = r['code']; type_map[r['code']] = r['code']
        subject_code = order.get('subject_code', '')
        region_code = order.get('region_code', '')
        # 先更新订单状态（必须在 gen_asset_no 关闭线程局部连接之前执行）
        if order['status'] not in ('已到货',):
            conn.execute("UPDATE purchase_orders SET status='采购完成待入库', updated_at=datetime('now','localtime') WHERE id=?", (order_id,))
            conn.commit()
        # 生成模板（本地维护递增序号，避免 gen_asset_no 关闭连接且批量时不递增）
        sc = subject_code or 'NA'
        rc = region_code or 'NA'
        today = datetime.now().strftime('%Y%m%d')
        wb = Workbook(); ws = wb.active; ws.title = '资产入库'
        headers = ['资产编号', '资产类目', '品牌', '型号', '规格', 'SN', 'MAC(有线)', 'MAC(无线)',
                   '存放位置', '关联人', 'KEY', '关联账号', '关联密码', '备注']
        ws.append(headers)
        gen_list = []
        for it in items:
            it = dict(it); qty = it['quantity']
            cat = it.get('category', '')
            atc = type_map.get(cat, 'OTH')
            like = f"{today}%{sc}%{rc}%{atc}%"
            row = conn.execute("SELECT COUNT(*) as cnt FROM assets WHERE asset_no LIKE ?", (like,)).fetchone()
            seq = row['cnt'] or 0
            for _ in range(qty):
                seq += 1
                no = f"{today}{seq:03d}-{sc}-{rc}-{atc}"
                ws.append([no, cat, it.get('brand', ''), it.get('model', ''), it.get('spec', ''),
                           '', '', '', '', '', '', '', '', f"采购订单{order['order_no']}"])
                gen_list.append(no)
        fname = f"stock_in_{order['order_no']}.xlsx"
        fpath = _os.path.join(_EXPORT_DIR, fname)
        wb.save(fpath)
        try:
            log_action('complete', 'order', 'purchase_order', order_id, order['order_no'], order['title'], {'count': len(gen_list)})
        except Exception:
            pass
        return jsonify({'success': True, 'file_url': f'/api/orders/{order_id}/template-download',
                        'file_name': fname, 'asset_nos': gen_list, 'count': len(gen_list)})
    finally:
        try:
            conn.close()
        except Exception:
            pass


@app.route('/api/orders/<int:order_id>/template-download', methods=['GET'])
@login_required
def order_template_download(order_id):
    conn = get_conn()
    try:
        order = conn.execute("SELECT order_no FROM purchase_orders WHERE id=?", (order_id,)).fetchone()
        if not order:
            return jsonify({'success': False}), 404
        fname = f"stock_in_{order['order_no']}.xlsx"
        fpath = _os.path.join(_EXPORT_DIR, fname)
        if not _os.path.exists(fpath):
            return jsonify({'success': False, 'error': '模板未生成，请先点击采购完成'}), 404
        return send_file(fpath, as_attachment=True, download_name=fname)
    finally:
        conn.close()


@app.route('/api/orders/<int:order_id>/import-template', methods=['POST'])
@login_required
def order_import_template(order_id):
    """回传 Excel 模板：批量建档为'待审批入库'，并生成采购入库审批单"""
    file = request.files.get('file')
    if not file:
        return jsonify({'success': False, 'error': '未收到文件'}), 400
    apply_by = request.form.get('apply_by') or '操作员'
    conn = get_conn()
    try:
        order = conn.execute("SELECT * FROM purchase_orders WHERE id=?", (order_id,)).fetchone()
        if not order:
            return jsonify({'success': False, 'error': '订单不存在'}), 404
        order = dict(order)
        tmp = _os.path.join(_EXPORT_DIR, f'upload_{order_id}.xlsx')
        file.save(tmp)
        wb = load_workbook(tmp); ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        created = []
        today = datetime.now().strftime('%Y-%m-%d')
        for r in rows[1:]:
            if not r or not r[0]:
                continue
            asset_no = str(r[0]).strip()
            if not asset_no:
                continue
            cur = conn.execute('''INSERT INTO assets
                (asset_no,category,brand,model,spec,sn,mac_wired,mac_wireless,status,location,
                 assignee,key_info,account_name,account_password,notes,purchase_date,order_id)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (asset_no, r[1] or '', r[2] or '', r[3] or '', r[4] or '', r[5] or '', r[6] or '', r[7] or '',
                 '待审批入库', r[8] or '', r[9] or '', r[10] or '', r[11] or '', r[12] or '', r[13] or '', today, order_id))
            created.append(cur.lastrowid)
        if created:
            conn.execute('''INSERT INTO stock_in_applications (order_id, order_no, asset_ids, apply_by, status)
                            VALUES (?,?,?,?,?)''',
                         (order_id, order['order_no'], _json.dumps(created), apply_by, 'pending'))
        conn.execute("UPDATE purchase_orders SET status='采购完成待入库', updated_at=datetime('now','localtime') WHERE id=?", (order_id,))
        conn.commit()
        try:
            log_action('import', 'asset', 'purchase_order', order_id, order['order_no'], order['title'], {'count': len(created), 'asset_ids': created})
        except Exception:
            pass
        return jsonify({'success': True, 'count': len(created), 'message': f'已生成 {len(created)} 条待审批资产，请在审批中心确认入库'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


# ─────────────────────────────────────────────
# 出入库操作
# ─────────────────────────────────────────────

@app.route('/api/stock/out', methods=['POST'])
def stock_out():
    """出库（分配给人员）—— 提交审批，超管确认后执行"""
    data = request.json or {}
    aid = data.get('asset_id')
    if not aid:
        return jsonify({'error': 'asset_id 不能为空'}), 400
    conn = get_conn()
    asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not asset:
        return jsonify({'error': '资产不存在'}), 404
    if asset['status'] not in ('在库',):
        return jsonify({'error': f'资产当前状态为【{asset["status"]}】，无法出库'}), 400
    try:
        create_operation_application('stock_out', [aid], data, applicant=get_current_name())
        return jsonify({'success': True, 'message': '已提交出库申请，请等待超管审批'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/stock/return', methods=['POST'])
def stock_return():
    """归还入库 —— 提交审批，超管确认后执行"""
    data = request.json or {}
    aid = data.get('asset_id')
    if not aid:
        return jsonify({'error': 'asset_id 不能为空'}), 400
    conn = get_conn()
    asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not asset:
        return jsonify({'error': '资产不存在'}), 404
    try:
        create_operation_application('stock_return', [aid], data, applicant=get_current_name())
        return jsonify({'success': True, 'message': '已提交归还申请，请等待超管审批'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/stock/transfer', methods=['POST'])
def stock_transfer():
    """调拨（换人/换地点，主体不变）—— 提交审批，超管确认后执行"""
    data = request.json or {}
    aid = data.get('asset_id')
    if not aid:
        return jsonify({'error': 'asset_id 不能为空'}), 400
    conn = get_conn()
    asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not asset:
        return jsonify({'error': '资产不存在'}), 404
    try:
        create_operation_application('transfer', [aid], data, applicant=get_current_name())
        return jsonify({'success': True, 'message': '已提交调拨申请，请等待超管审批'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/stock/asset-transfer', methods=['POST'])
def asset_transfer():
    """资产转移（跨主体，变更资产编号）—— 提交审批，超管确认后执行"""
    data = request.json or {}
    aid        = data.get('asset_id')
    to_subject = data.get('to_subject', '')
    to_region  = data.get('to_region', '')
    to_user    = data.get('target_user', '')
    to_loc     = data.get('to_location', '')
    operator   = get_current_name()  # 修复(四轮N4)：operator 强制取服务端登录身份，拒绝前端传值伪造
    remark     = data.get('remark', '')

    if not to_subject:
        return jsonify({'error': '请选择目标主体'}), 400

    conn = get_conn()
    asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not asset:
        return jsonify({'error': '资产不存在'}), 404

    subj = conn.execute("SELECT * FROM asset_subjects WHERE code=? AND active=1", (to_subject,)).fetchone()
    if not subj:
        return jsonify({'error': f'主体 {to_subject} 不存在或已停用'}), 400

    if not to_region:
        reg_row = conn.execute("SELECT code FROM asset_regions WHERE active=1 ORDER BY sort_order LIMIT 1").fetchone()
        to_region = reg_row['code'] if reg_row else ''
    if not to_loc:
        to_loc = asset['location']
    if not to_user:
        to_user = asset['assignee'] or ''

    payload = {
        'to_subject': to_subject, 'to_region': to_region,
        'target_user': to_user, 'to_location': to_loc,
        'operator': operator, 'remark': remark
    }
    try:
        create_operation_application('asset_transfer', [aid], payload, applicant=operator)
        return jsonify({'success': True, 'message': '已提交资产转移申请，请等待超管审批'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/stock/borrow-out', methods=['POST'])
def stock_borrow_out():
    """借用出库 —— 提交审批，超管确认后执行（资产状态变为"在用"，log_type=借用出库）"""
    data = request.json or {}
    aid = data.get('asset_id')
    if not aid:
        return jsonify({'error': 'asset_id 不能为空'}), 400
    conn = get_conn()
    asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not asset:
        return jsonify({'error': '资产不存在'}), 404
    if not asset.get('is_borrowable'):
        return jsonify({'error': '该资产未标记为可借，无法借用出库'}), 400
    if asset['status'] not in ('在库',):
        return jsonify({'error': f'资产当前状态为【{asset["status"]}】，无法借用出库'}), 400
    try:
        payload = {
            'target_user': data.get('target_user', ''),
            'remark': data.get('remark', '借用出库')
        }
        create_operation_application('borrow_out', [aid], payload, applicant=get_current_name())
        return jsonify({'success': True, 'message': '已提交借用出库申请，请等待超管审批'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/stock/borrow-return', methods=['POST'])
def stock_borrow_return():
    """借用归还 —— 提交审批，超管确认后执行（资产状态恢复为"在库"，log_type=借用归还）"""
    data = request.json or {}
    aid = data.get('asset_id')
    if not aid:
        return jsonify({'error': 'asset_id 不能为空'}), 400
    conn = get_conn()
    asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not asset:
        return jsonify({'error': '资产不存在'}), 404
    if asset['status'] != '在用':
        return jsonify({'error': f'资产当前状态为【{asset["status"]}】，只有"在用"状态才能归还'}), 400
    try:
        payload = {
            'target_user': data.get('target_user', ''),
            'remark': data.get('remark', '借用归还')
        }
        create_operation_application('borrow_return', [aid], payload, applicant=get_current_name())
        return jsonify({'success': True, 'message': '已提交借用归还申请，请等待超管审批'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/assets/borrowable', methods=['GET'])
def list_borrowable_assets():
    """获取可用于借用的资产列表（is_borrowable=1 且 status=在库）"""
    conn = get_conn()
    try:
        rows = conn.execute("""
            SELECT id, asset_no, category, brand, model, status, location, assignee, is_borrowable
            FROM assets WHERE is_borrowable = 1 AND status = '在库'
            ORDER BY category, brand, model
        """).fetchall()
        return jsonify({'success': True, 'data': rows_to_list(rows), 'total': len(rows)})
    finally:
        conn.close()


@app.route('/api/stock/logs', methods=['GET'])
def list_stock_logs():
    conn = get_conn()
    asset_id = request.args.get('asset_id')
    log_type = request.args.get('log_type', '')
    page     = max(1, int(request.args.get('page', 1)))
    page_size = int(request.args.get('page_size', 20))

    sql    = "SELECT * FROM stock_logs WHERE 1=1"
    params = []
    if asset_id:
        sql += " AND asset_id=?"; params.append(int(asset_id))
    if log_type:
        sql += " AND log_type=?"; params.append(log_type)

    total = conn.execute(
        sql.replace("SELECT *", "SELECT COUNT(*) as cnt"), params
    ).fetchone()['cnt']
    sql += " ORDER BY id DESC LIMIT ? OFFSET ?"
    params += [page_size, (page-1)*page_size]
    rows = conn.execute(sql, params).fetchall()
    # 获取每张照片
    result = rows_to_list(rows)
    for row in result:
        photos = conn.execute("SELECT * FROM photos WHERE ref_type='stock' AND ref_id=?", (row['id'],)).fetchall()
        row['photos'] = rows_to_list(photos)
    conn.close()
    return jsonify({'total': total, 'data': result})


@app.route('/api/stock/logs/<int:log_id>/photos', methods=['POST'])
def upload_stock_photo(log_id):
    """上传出入库照片"""
    if 'file' not in request.files:
        return jsonify({'error': '没有文件'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': '没有选择文件'}), 400

    filename = secure_filename(file.filename)
    timestamp = _rand_suffix()  # 安全加固 P2-4：随机段替代可推算时间戳，防附件URL枚举
    new_filename = f"stock_{log_id}_{timestamp}_{filename}"
    file_path = os.path.join(UPLOAD_DIR, new_filename)
    file.save(file_path)

    conn = get_conn()
    cur = conn.execute('''
        INSERT INTO photos(ref_type, ref_id, file_name, file_path, file_size, mime_type, uploaded_by)
        VALUES(?,?,?,?,?,?,?)
    ''', ('stock', log_id, new_filename, file_path, os.path.getsize(file_path),
          file.content_type, request.form.get('uploaded_by', '系统')))
    photo_id = cur.lastrowid
    conn.commit()
    conn.close()
    return jsonify({'id': photo_id, 'file_name': new_filename}), 201


@app.route('/api/photos/<int:photo_id>', methods=['GET'])
def get_photo(photo_id):
    """获取照片文件"""
    conn = get_conn()
    photo = conn.execute("SELECT * FROM photos WHERE id=?", (photo_id,)).fetchone()
    conn.close()
    if not photo:
        return jsonify({'error': '照片不存在'}), 404
    return send_file(photo['file_path'], mimetype=photo['mime_type'])


# ─────────────────────────────────────────────
# 维修管理
# ─────────────────────────────────────────────

@app.route('/api/repair', methods=['POST'])
def submit_repair():
    """报修 —— 提交审批，超管确认后执行（创建维修记录并置为维修中）"""
    data = request.json or {}
    aid = data.get('asset_id')
    if not aid:
        return jsonify({'error': 'asset_id 不能为空'}), 400
    conn = get_conn()
    asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not asset:
        return jsonify({'error': '资产不存在'}), 404
    payload = {
        'fault_desc': data.get('fault_desc', ''),
        'repair_vendor': data.get('repair_vendor', ''),
        'submit_by': data.get('submit_by', ''),
        'repair_notes': data.get('repair_notes', '')
    }
    try:
        app_id = create_operation_application('repair', [aid], payload, applicant=data.get('submit_by'))
        return jsonify({'success': True, 'app_id': app_id, 'message': '已提交报修申请，请等待超管审批'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/repair/<int:rid>', methods=['PUT'])
def update_repair(rid):
    data = request.json or {}
    conn = get_conn()
    rep = conn.execute("SELECT * FROM repair_logs WHERE id=?", (rid,)).fetchone()
    if not rep:
        conn.close()
        return jsonify({'error': '维修记录不存在'}), 404

    new_status = data.get('repair_status', rep['repair_status'])
    fields  = ['repair_status','repair_vendor','repair_cost','repair_notes']
    updates = []
    params  = []
    for f in fields:
        if f in data:
            updates.append(f'{f}=?')
            params.append(data[f])
    if new_status == '已完成' and not rep['finish_time']:
        updates.append("finish_time=datetime('now','localtime')")
    params.append(rid)
    if updates:
        conn.execute(f"UPDATE repair_logs SET {','.join(updates)} WHERE id=?", params)

    # 同步更新资产状态
    if new_status in ('已完成', '无法修复'):
        restore_status = '在库' if new_status == '已完成' else '已报废'
        conn.execute('''UPDATE assets SET status=?,
                        updated_at=datetime('now','localtime') WHERE id=?''',
                     (restore_status, rep['asset_id']))
        conn.execute('''INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,remark)
                        VALUES(?,?,?,?,?)''',
                     (rep['asset_id'], rep['asset_no'],
                      '归还' if new_status == '已完成' else '报废',
                      get_current_name(),
                      f"维修完成: {data.get('repair_notes','')}"))

        notify_lark('repair',
            f"✅ 【维修完成】\n编号：{rep['asset_no']}\n状态：{new_status}\n费用：{data.get('repair_cost','未填写')}\n说明：{data.get('repair_notes','')}\n时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")

    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/repair', methods=['GET'])
def list_repairs():
    conn = get_conn()
    status   = request.args.get('status', '')
    asset_id = request.args.get('asset_id', '')
    page     = max(1, int(request.args.get('page', 1)))
    page_size = int(request.args.get('page_size', 20))

    sql    = "SELECT r.*, a.category, a.brand, a.model FROM repair_logs r LEFT JOIN assets a ON r.asset_id=a.id WHERE 1=1"
    params = []
    if status:
        sql += " AND r.repair_status=?"; params.append(status)
    if asset_id:
        sql += " AND r.asset_id=?"; params.append(int(asset_id))

    total = conn.execute(
        sql.replace("SELECT r.*, a.category, a.brand, a.model", "SELECT COUNT(*) as cnt"), params
    ).fetchone()['cnt']
    sql += " ORDER BY r.id DESC LIMIT ? OFFSET ?"
    params += [page_size, (page-1)*page_size]
    rows = conn.execute(sql, params).fetchall()
    # 获取每张照片
    result = rows_to_list(rows)
    for row in result:
        photos = conn.execute("SELECT * FROM photos WHERE ref_type='repair' AND ref_id=?", (row['id'],)).fetchall()
        row['photos'] = rows_to_list(photos)
    conn.close()
    return jsonify({'total': total, 'data': result})


@app.route('/api/repair/<int:repair_id>/photos', methods=['POST'])
def upload_repair_photo(repair_id):
    """上传维修照片"""
    if 'file' not in request.files:
        return jsonify({'error': '没有文件'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': '没有选择文件'}), 400

    filename = secure_filename(file.filename)
    timestamp = _rand_suffix()  # 安全加固 P2-4：随机段替代可推算时间戳，防附件URL枚举
    new_filename = f"repair_{repair_id}_{timestamp}_{filename}"
    file_path = os.path.join(UPLOAD_DIR, new_filename)
    file.save(file_path)

    conn = get_conn()
    cur = conn.execute('''
        INSERT INTO photos(ref_type, ref_id, file_name, file_path, file_size, mime_type, uploaded_by)
        VALUES(?,?,?,?,?,?,?)
    ''', ('repair', repair_id, new_filename, file_path, os.path.getsize(file_path),
          file.content_type, request.form.get('uploaded_by', '系统')))
    photo_id = cur.lastrowid
    conn.commit()
    conn.close()
    return jsonify({'id': photo_id, 'file_name': new_filename}), 201


# ─────────────────────────────────────────────
# 统计仪表盘
# ─────────────────────────────────────────────

@app.route('/api/stats/overview', methods=['GET'])
@login_required
def stats_overview():
    conn = get_conn()
    # 权限隔离：普通用户只统计「自己名下」的资产 / 耗材；其它角色（超管/管理员/普通管理员/操作员/财务）仍看全量
    me = conn.execute(
        "SELECT name, role FROM employees WHERE id = ?", (request.current_user_id,)
    ).fetchone()
    is_admin = bool(me and me['role'] in ADMIN_ROLES)
    me_name = me['name'] if me else ''

    purchase_value = 0.0
    book_value_total = 0.0
    pending_approvals = 0
    fulfillable_count = 0
    warranty_alert_count = 0
    warranty_alerts = []
    inventory_ongoing = 0
    idle_count = 0
    idle_assets = []
    value_by_category = []
    pending_po_amount = 0.0
    pending_po_count = 0
    uninvoiced_po_count = 0

    if is_admin:
        total   = conn.execute("SELECT COUNT(*) as c FROM assets").fetchone()['c']
        in_use  = conn.execute("SELECT COUNT(*) as c FROM assets WHERE status='在用'").fetchone()['c']
        in_stock= conn.execute("SELECT COUNT(*) as c FROM assets WHERE status='在库'").fetchone()['c']
        repairing=conn.execute("SELECT COUNT(*) as c FROM assets WHERE status='维修中'").fetchone()['c']
        scrapped= conn.execute("SELECT COUNT(*) as c FROM assets WHERE status='已报废'").fetchone()['c']
        by_cat  = rows_to_list(conn.execute(
            "SELECT category, COUNT(*) as cnt FROM assets GROUP BY category ORDER BY cnt DESC"
        ).fetchall())
        by_loc  = rows_to_list(conn.execute(
            "SELECT location, COUNT(*) as cnt FROM assets GROUP BY location ORDER BY cnt DESC"
        ).fetchall())
        top_assignees = rows_to_list(conn.execute(
            "SELECT assignee, COUNT(*) as cnt FROM assets WHERE assignee!='' GROUP BY assignee ORDER BY cnt DESC LIMIT 10"
        ).fetchall())
        recent_logs_raw = conn.execute(
            "SELECT * FROM stock_logs ORDER BY id DESC LIMIT 10"
        ).fetchall()
        cons_total = conn.execute("SELECT COUNT(*) as c FROM consumables").fetchone()['c']
        cons_normal = conn.execute("SELECT COUNT(*) as c FROM consumables WHERE status='正常'").fetchone()['c']
        cons_low = conn.execute("SELECT COUNT(*) as c FROM consumables WHERE status='低库存'").fetchone()['c']
        cons_by_cat = rows_to_list(conn.execute(
            "SELECT category, COUNT(*) as cnt FROM consumables GROUP BY category ORDER BY cnt DESC"
        ).fetchall())

        # 原值 / 账面价值（未报废）；待审批 / 待发放；保修临期
        pv_row = conn.execute(
            "SELECT COALESCE(SUM(purchase_price),0) as s FROM assets "
            "WHERE status!='已报废' AND purchase_price IS NOT NULL AND purchase_price!=''"
        ).fetchone()
        purchase_value = round(float(pv_row['s'] or 0), 2)
        pmap = load_type_params_map(conn)
        valued = conn.execute(
            "SELECT purchase_price, purchase_date, asset_type_code, category FROM assets WHERE status!='已报废'"
        ).fetchall()
        bv_sum = 0.0
        cat_bv = {}
        for a in valued:
            code = a['asset_type_code'] or resolve_asset_type_code(
                conn, a['category'], a['asset_type_code'] or '')
            d = attach_book_value({
                'purchase_price': a['purchase_price'],
                'purchase_date': a['purchase_date'],
                'asset_type_code': code,
            }, pmap)
            if d.get('book_value') is not None:
                bv = float(d['book_value'])
                bv_sum += bv
                cat = a['category'] or '未分类'
                cat_bv[cat] = cat_bv.get(cat, 0.0) + bv
        book_value_total = round(bv_sum, 2)
        value_by_category = [
            {'category': k, 'book_value': round(v, 2)}
            for k, v in sorted(cat_bv.items(), key=lambda x: -x[1])
        ]

        pending_approvals = _count_pending_approvals(conn)
        fulfillable_count = conn.execute(
            "SELECT COUNT(*) as c FROM asset_applications "
            "WHERE status='approved' AND IFNULL(fulfilled,0)=0 "
            "AND app_type IN ('new_asset','replace','special')"
        ).fetchone()['c']

        today = datetime.now().date()
        wrows = conn.execute("""
            SELECT id, asset_no, category, brand, model, assignee, warranty_expire, status
            FROM assets
            WHERE warranty_expire IS NOT NULL AND warranty_expire != ''
            ORDER BY warranty_expire ASC
        """).fetchall()
        for r in wrows:
            try:
                wdate = datetime.strptime(str(r['warranty_expire'])[:10], '%Y-%m-%d').date()
            except (TypeError, ValueError):
                continue
            days_left = (wdate - today).days
            if days_left > 30:
                continue
            st = 'expired' if days_left < 0 else 'warning'
            warranty_alerts.append({
                'id': r['id'],
                'asset_no': r['asset_no'],
                'category': r['category'],
                'brand': r['brand'],
                'model': r['model'],
                'assignee': r['assignee'],
                'warranty_expire': r['warranty_expire'],
                'days_left': days_left,
                'expire_status': st,
                'status': r['status'],
            })
        warranty_alert_count = len(warranty_alerts)

        inventory_ongoing = conn.execute(
            "SELECT COUNT(*) as c FROM inventory_tasks WHERE status IN ('ongoing','进行中')"
        ).fetchone()['c']

        idle_cutoff = (today - timedelta(days=90)).isoformat()
        irows = conn.execute("""
            SELECT id, asset_no, category, brand, model, location, purchase_date, status
            FROM assets
            WHERE status='在库'
              AND purchase_date IS NOT NULL AND purchase_date != ''
              AND purchase_date <= ?
            ORDER BY purchase_date ASC
            LIMIT 100
        """, (idle_cutoff,)).fetchall()
        for r in irows:
            try:
                pdate = datetime.strptime(str(r['purchase_date'])[:10], '%Y-%m-%d').date()
            except (TypeError, ValueError):
                continue
            idle_days = (today - pdate).days
            idle_assets.append({
                'id': r['id'],
                'asset_no': r['asset_no'],
                'category': r['category'],
                'brand': r['brand'],
                'model': r['model'],
                'location': r['location'],
                'purchase_date': r['purchase_date'],
                'idle_days': idle_days,
                'status': r['status'],
            })
        idle_count = conn.execute("""
            SELECT COUNT(*) as c FROM assets
            WHERE status='在库'
              AND purchase_date IS NOT NULL AND purchase_date != ''
              AND purchase_date <= ?
        """, (idle_cutoff,)).fetchone()['c']

        # 采购待办：未全部到货（非取消）金额与笔数；未到票（无发票号且无附件）
        po_open = conn.execute("""
            SELECT COUNT(*) as c, COALESCE(SUM(total_amount), 0) as s
            FROM purchase_orders
            WHERE status NOT IN ('已取消', '已到货')
        """).fetchone()
        pending_po_count = po_open['c'] or 0
        pending_po_amount = round(float(po_open['s'] or 0), 2)
        uninvoiced_po_count = conn.execute("""
            SELECT COUNT(*) as c FROM purchase_orders
            WHERE status != '已取消'
              AND TRIM(IFNULL(invoice_no,'')) = ''
              AND TRIM(IFNULL(invoice_file,'')) = ''
        """).fetchone()['c']
    else:
        # 普通用户只统计自己名下的资产 / 自己领用的耗材
        total   = conn.execute("SELECT COUNT(*) as c FROM assets WHERE assignee=?", (me_name,)).fetchone()['c']
        in_use  = conn.execute("SELECT COUNT(*) as c FROM assets WHERE assignee=? AND status='在用'", (me_name,)).fetchone()['c']
        in_stock= conn.execute("SELECT COUNT(*) as c FROM assets WHERE assignee=? AND status='在库'", (me_name,)).fetchone()['c']
        repairing=conn.execute("SELECT COUNT(*) as c FROM assets WHERE assignee=? AND status='维修中'", (me_name,)).fetchone()['c']
        scrapped= conn.execute("SELECT COUNT(*) as c FROM assets WHERE assignee=? AND status='已报废'", (me_name,)).fetchone()['c']
        by_cat  = rows_to_list(conn.execute(
            "SELECT category, COUNT(*) as cnt FROM assets WHERE assignee=? GROUP BY category ORDER BY cnt DESC", (me_name,)
        ).fetchall())
        by_loc  = rows_to_list(conn.execute(
            "SELECT location, COUNT(*) as cnt FROM assets WHERE assignee=? GROUP BY location ORDER BY cnt DESC", (me_name,)
        ).fetchall())
        top_assignees = []
        recent_logs_raw = conn.execute(
            "SELECT * FROM stock_logs WHERE target_user=? ORDER BY id DESC LIMIT 10", (me_name,)
        ).fetchall()
        cons_total = conn.execute(
            "SELECT COUNT(DISTINCT item_id) as c FROM consumable_logs WHERE target_user=?", (me_name,)
        ).fetchone()['c']
        cons_normal = 0  # 普通用户视角下不展示全公司库存健康度，避免误读
        cons_low = 0
        cons_by_cat = rows_to_list(conn.execute(
            "SELECT c.category as category, COUNT(*) as cnt FROM consumable_logs l JOIN consumables c ON c.id=l.item_id WHERE l.target_user=? GROUP BY c.category ORDER BY cnt DESC", (me_name,)
        ).fetchall())

    by_status = rows_to_list(conn.execute(
        "SELECT status, COUNT(*) as cnt FROM assets GROUP BY status"
    ).fetchall())

    recent_logs = rows_to_list(recent_logs_raw)
    for log in recent_logs:
        photos = conn.execute(
            "SELECT * FROM photos WHERE ref_type='stock' AND ref_id=?", (log['id'],)
        ).fetchall()
        log['photos'] = rows_to_list(photos)
    repair_stats = rows_to_list(conn.execute(
        "SELECT repair_status, COUNT(*) as cnt FROM repair_logs GROUP BY repair_status"
    ).fetchall())
    monthly_in = rows_to_list(conn.execute(
        "SELECT substr(log_time,1,7) as month, COUNT(*) as cnt FROM stock_logs WHERE log_type='入库' GROUP BY month ORDER BY month DESC LIMIT 12"
    ).fetchall())

    conn.close()
    return jsonify({
        'total': total, 'in_use': in_use, 'in_stock': in_stock,
        'repairing': repairing, 'scrapped': scrapped,
        'by_category': by_cat,
        'by_location': by_loc,
        'by_status': by_status,
        'top_assignees': top_assignees,
        'recent_logs': recent_logs,
        'repair_stats': repair_stats,
        'monthly_in': monthly_in,
        # 耗材统计
        'consumable_total': cons_total,
        'consumable_normal': cons_normal,
        'consumable_low': cons_low,
        'consumable_by_category': cons_by_cat,
        # 管理端一眼可见：原值 / 账面 / 待办 / 保修 / 盘点 / 闲置 / 类目价值
        'purchase_value': purchase_value,
        'book_value_total': book_value_total,
        'pending_approvals': pending_approvals,
        'fulfillable_count': fulfillable_count,
        'warranty_alert_count': warranty_alert_count,
        'warranty_alerts': warranty_alerts,
        'inventory_ongoing': inventory_ongoing,
        'idle_count': idle_count,
        'idle_assets': idle_assets,
        'value_by_category': value_by_category,
        'pending_po_amount': pending_po_amount,
        'pending_po_count': pending_po_count,
        'uninvoiced_po_count': uninvoiced_po_count,
        # 提示前端是否为受限制视图（普通用户），用于 UI 提示
        'restricted_view': not is_admin,
    })


# ─────────────────────────────────────────────
# 人员管理
# ─────────────────────────────────────────────

@app.route('/api/employees', methods=['GET'])
def list_employees():
    conn = get_conn()
    # 安全加固（三轮审计#4）：SELECT * → 显式列名，剔除 password 哈希（防泄露后被离线爆破）。
    rows = conn.execute('''SELECT id, name, role, dept, location, email, active, status, avatar,
                           lark_user_id, lark_open_id, lark_union_id, lark_only, last_login, created_at, custom_fields
                           FROM employees WHERE active=1 ORDER BY name''').fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))


@app.route('/api/employees/search', methods=['GET'])
@login_required
def search_employees():
    """搜索用户（支持姓名、邮箱、部门模糊搜索）"""
    query = request.args.get('q', '').strip()
    if not query:
        return jsonify({'success': True, 'data': []})
    
    conn = get_conn()
    try:
        # 模糊搜索姓名、邮箱、部门
        clause, params = fuzzy_match_clause(query, ['name', 'email', 'dept'])
        sql = f"""SELECT id, name, email, dept, role, status 
               FROM employees 
               WHERE 1=1 AND status = '在职'"""
        if clause:
            sql += f" AND {clause}"
        sql += " ORDER BY name LIMIT 20"
        rows = conn.execute(sql, params).fetchall()
        return jsonify({'success': True, 'data': rows_to_list(rows)})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/employees', methods=['POST'])
def create_employee():
    data = request.json or {}
    conn = get_conn()
    # 安全加固（三轮审计#1）：role 属敏感权限字段，仅超管可指定非普通角色；
    # 非超管强制为「普通用户」，防止操作员/财务经此接口创建超管账号越权提权。
    _cu = getattr(request, 'current_user', None) or {}
    is_super = _cu.get('role') == '超管'
    role = data.get('role', '普通用户') if is_super else '普通用户'
    cur = conn.execute(
        "INSERT INTO employees(name,dept,location,email,role) VALUES(?,?,?,?,?)",
        (data.get('name',''), data.get('dept',''), data.get('location',''), data.get('email',''), role)
    )
    conn.commit()
    conn.close()
    return jsonify({'id': cur.lastrowid}), 201


@app.route('/api/employees/<int:eid>', methods=['PUT'])
def update_employee(eid):
    data = request.json or {}
    conn = get_conn()
    # 安全加固（三轮审计#1）：role 仅超管可修改，防止操作员/财务把自己或他人改为超管（越权提权）。
    # 修复（四轮N3）：active 与 role 同权限——非超管也不能改 active，
    # 否则操作员/财务可将任意员工（含超管）active=0 禁用账号。
    _cu = getattr(request, 'current_user', None) or {}
    is_super = _cu.get('role') == '超管'
    if is_super:
        conn.execute(
            "UPDATE employees SET name=?,dept=?,location=?,email=?,role=?,active=? WHERE id=?",
            (data.get('name',''), data.get('dept',''), data.get('location',''),
             data.get('email',''), data.get('role','普通用户'), data.get('active',1), eid)
        )
    else:
        # 非超管：不更新 role / active（均保持原值），仅可维护基础信息
        conn.execute(
            "UPDATE employees SET name=?,dept=?,location=?,email=? WHERE id=?",
            (data.get('name',''), data.get('dept',''), data.get('location',''),
             data.get('email',''), eid)
        )
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


# ─────────────────────────────────────────────
# 飞书配置
# ─────────────────────────────────────────────

@app.route('/api/lark/config', methods=['GET'])
def lark_list():
    conn = get_conn()
    rows = conn.execute("SELECT * FROM lark_config ORDER BY id").fetchall()
    conn.close()
    data = rows_to_list(rows)
    # 安全加固（P2-3）：secret/webhook 掩码返回，避免密钥明文泄露给前端
    for r in data:
        if r.get('secret'):
            r['secret'] = '********'
        url = r.get('webhook_url') or ''
        if len(url) > 40:
            r['webhook_url'] = url[:40] + '...****'
    return jsonify(data)


@app.route('/api/lark/config', methods=['POST'])
def lark_add():
    data = request.json or {}
    webhook_url = data.get('webhook_url', '')
    # 安全加固（三轮审计#3）：webhook 强制飞书官方域名校验，防存储型SSRF。
    # （此前仅 lark_test 校验，遗漏了 lark_add/lark_update 两个写入口——配置后 notify_lark 会从服务端 POST 到该 URL）
    if webhook_url and not _is_valid_lark_webhook(webhook_url):
        return jsonify({'ok': False, 'msg': '仅支持飞书官方 Webhook 地址（https://open.feishu.cn 或 https://open.larksuite.com）'}), 400
    conn = get_conn()
    cur = conn.execute('''
        INSERT INTO lark_config(name,webhook_url,secret,enabled,notify_stock,notify_repair,notify_expire)
        VALUES(?,?,?,?,?,?,?)
    ''', (data.get('name',''), webhook_url, data.get('secret',''),
          data.get('enabled',1), data.get('notify_stock',1), data.get('notify_repair',1),
          data.get('notify_expire',1)))
    conn.commit()
    conn.close()
    return jsonify({'id': cur.lastrowid}), 201


@app.route('/api/lark/config/<int:cid>', methods=['PUT'])
def lark_update(cid):
    data = request.json or {}
    conn = get_conn()
    # 安全加固（P2-3）：secret/webhook 提交掩码值或为空时保留旧值（不覆盖真实配置），
    # 管理员仅在显式输入新值时才更新，配合 GET 脱敏返回。
    old = conn.execute("SELECT secret, webhook_url FROM lark_config WHERE id=?", (cid,)).fetchone()
    secret = data.get('secret', '')
    if not secret or secret == '********':
        secret = old['secret'] if old else ''
    webhook_url = data.get('webhook_url', '')
    if not webhook_url or webhook_url.endswith('...****'):
        webhook_url = old['webhook_url'] if old else ''
    # 安全加固（三轮审计#3）：webhook 强制飞书官方域名校验，防存储型SSRF（非保留旧值的新地址必须校验）
    if webhook_url and not _is_valid_lark_webhook(webhook_url):
        return jsonify({'ok': False, 'msg': '仅支持飞书官方 Webhook 地址（https://open.feishu.cn 或 https://open.larksuite.com）'}), 400
    conn.execute('''
        UPDATE lark_config SET name=?,webhook_url=?,secret=?,
        enabled=?,notify_stock=?,notify_repair=?,notify_expire=? WHERE id=?
    ''', (data.get('name',''), webhook_url, secret,
          data.get('enabled',1), data.get('notify_stock',1), data.get('notify_repair',1),
          data.get('notify_expire',1), cid))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/lark/config/<int:cid>', methods=['DELETE'])
def lark_delete(cid):
    conn = get_conn()
    conn.execute("DELETE FROM lark_config WHERE id=?", (cid,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/lark/test', methods=['POST'])
def lark_test():
    data = request.json or {}
    webhook_url = data.get('webhook_url', '')
    # 安全加固（P2-1）：SSRF 防护——仅允许飞书官方 webhook 域名，拒绝内网/回环/任意 URL
    if not _is_valid_lark_webhook(webhook_url):
        return jsonify({'ok': False, 'msg': '仅支持飞书官方 Webhook 地址（https://open.feishu.cn 或 https://open.larksuite.com）'}), 400
    ok = lark_send(webhook_url, data.get('secret'), '🔔 资产管理系统飞书通知测试消息，连接成功！')
    return jsonify({'ok': ok, 'msg': '发送成功' if ok else '发送失败，请检查 Webhook 地址'})


# ─────────────────────────────────────────────
# 系统配置
# ─────────────────────────────────────────────

@app.route('/api/config', methods=['GET'])
def get_config():
    conn = get_conn()
    rows = conn.execute("SELECT * FROM sys_config").fetchall()
    conn.close()
    config = {r['key']: r['value'] for r in rows}
    # 注入环境变量配置（Lark 免登需要）
    config['LARK_APP_ID'] = os.environ.get('LARK_APP_ID', '')
    return jsonify(config)


@app.route('/api/config', methods=['PUT'])
def update_config():
    data = request.json or {}
    conn = get_conn()
    for k, v in data.items():
        conn.execute("INSERT OR REPLACE INTO sys_config(key,value) VALUES(?,?)", (k, v))
    conn.commit()
    log_action('update', 'config', 'sys_config', 0,
          "sys_config", "系统配置",
          (get_current_user() or {}).get('name','系统'))
    conn.close()
    return jsonify({'ok': True})


# ─────────────────────────────────────────────
# 资产编号配置管理
# ─────────────────────────────────────────────

@app.route('/api/asset-config/subjects', methods=['GET'])
def get_asset_subjects():
    """获取主体列表"""
    conn = get_conn()
    rows = conn.execute("SELECT * FROM asset_subjects ORDER BY sort_order, id").fetchall()
    conn.close()
    return jsonify({'success': True, 'data': rows_to_list(rows)})

@app.route('/api/asset-config/subjects', methods=['POST'])
def create_asset_subject():
    """创建主体"""
    data = request.json or {}
    if not data.get('code') or not data.get('name'):
        return jsonify({'success': False, 'error': '代码和名称不能为空'}), 400
    conn = get_conn()
    try:
        cur = conn.execute(
            "INSERT INTO asset_subjects(code, name, region, sort_order) VALUES(?,?,?,?)",
            (data['code'], data['name'], data.get('region', ''), data.get('sort_order', 0))
        )
        conn.commit()
        return jsonify({'success': True, 'id': cur.lastrowid})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 400
    finally:
        conn.close()

@app.route('/api/asset-config/subjects/<int:sid>', methods=['PUT'])
def update_asset_subject(sid):
    """更新主体"""
    data = request.json or {}
    conn = get_conn()
    try:
        conn.execute(
            "UPDATE asset_subjects SET code=?, name=?, region=?, sort_order=?, active=? WHERE id=?",
            (data['code'], data['name'], data.get('region', ''), data.get('sort_order', 0), data.get('active', 1), sid)
        )
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 400
    finally:
        conn.close()

@app.route('/api/asset-config/subjects/<int:sid>', methods=['DELETE'])
def delete_asset_subject(sid):
    """删除主体"""
    conn = get_conn()
    conn.execute("DELETE FROM asset_subjects WHERE id=?", (sid,))
    conn.commit()
    conn.close()
    return jsonify({'success': True})


@app.route('/api/asset-config/types', methods=['GET'])
def get_asset_types():
    """获取资产类型列表"""
    conn = get_conn()
    rows = conn.execute("SELECT * FROM asset_types ORDER BY sort_order, id").fetchall()
    conn.close()
    return jsonify({'success': True, 'data': rows_to_list(rows)})

@app.route('/api/asset-config/types', methods=['POST'])
def create_asset_type():
    """创建资产类型"""
    data = request.json or {}
    if not data.get('code') or not data.get('name'):
        return jsonify({'success': False, 'error': '代码和名称不能为空'}), 400
    conn = get_conn()
    try:
        cur = conn.execute(
            "INSERT INTO asset_types(code, name, sort_order) VALUES(?,?,?)",
            (data['code'], data['name'], data.get('sort_order', 0))
        )
        # 自动为新类型创建默认折旧参数
        existing = conn.execute("SELECT id FROM asset_type_params WHERE type_code=?", (data['code'],)).fetchone()
        if not existing:
            conn.execute(
                "INSERT INTO asset_type_params (type_code, type_name, depreciation_months, residual_rate, max_issue_years) VALUES (?,?,36,0.05,2)",
                (data['code'], data['name'])
            )
        conn.commit()
        return jsonify({'success': True, 'id': cur.lastrowid})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 400
    finally:
        conn.close()

@app.route('/api/asset-config/types/<int:tid>', methods=['PUT'])
def update_asset_type(tid):
    """更新资产类型"""
    data = request.json or {}
    conn = get_conn()
    try:
        conn.execute(
            "UPDATE asset_types SET code=?, name=?, sort_order=?, active=? WHERE id=?",
            (data['code'], data['name'], data.get('sort_order', 0), data.get('active', 1), tid)
        )
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 400
    finally:
        conn.close()

@app.route('/api/asset-config/types/<int:tid>', methods=['DELETE'])
def delete_asset_type(tid):
    """删除资产类型"""
    conn = get_conn()
    conn.execute("DELETE FROM asset_types WHERE id=?", (tid,))
    conn.commit()
    conn.close()
    return jsonify({'success': True})


@app.route('/api/asset-config/regions', methods=['GET'])
def get_asset_regions():
    """获取存放地列表"""
    conn = get_conn()
    rows = conn.execute("SELECT * FROM asset_regions ORDER BY sort_order, id").fetchall()
    conn.close()
    return jsonify({'success': True, 'data': rows_to_list(rows)})

@app.route('/api/asset-config/regions', methods=['POST'])
def create_asset_region():
    """创建存放地"""
    data = request.json or {}
    if not data.get('code') or not data.get('name'):
        return jsonify({'success': False, 'error': '代码和名称不能为空'}), 400
    conn = get_conn()
    try:
        cur = conn.execute(
            "INSERT INTO asset_regions(code, name, sort_order) VALUES(?,?,?)",
            (data['code'], data['name'], data.get('sort_order', 0))
        )
        conn.commit()
        return jsonify({'success': True, 'id': cur.lastrowid})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 400
    finally:
        conn.close()

@app.route('/api/asset-config/regions/<int:rid>', methods=['PUT'])
def update_asset_region(rid):
    """更新存放地"""
    data = request.json or {}
    conn = get_conn()
    try:
        conn.execute(
            "UPDATE asset_regions SET code=?, name=?, sort_order=?, active=? WHERE id=?",
            (data['code'], data['name'], data.get('sort_order', 0), data.get('active', 1), rid)
        )
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 400
    finally:
        conn.close()

@app.route('/api/asset-config/regions/<int:rid>', methods=['DELETE'])
def delete_asset_region(rid):
    """删除存放地"""
    conn = get_conn()
    conn.execute("DELETE FROM asset_regions WHERE id=?", (rid,))
    conn.commit()
    conn.close()
    return jsonify({'success': True})


@app.route('/api/asset-config/all', methods=['GET'])
def get_all_asset_config():
    """获取所有资产配置（一次性获取）"""
    conn = get_conn()
    subjects = rows_to_list(conn.execute("SELECT * FROM asset_subjects ORDER BY sort_order").fetchall())
    types = rows_to_list(conn.execute("SELECT * FROM asset_types ORDER BY sort_order").fetchall())
    regions = rows_to_list(conn.execute("SELECT * FROM asset_regions ORDER BY sort_order").fetchall())
    conn.close()
    return jsonify({'success': True, 'data': {'subjects': subjects, 'types': types, 'regions': regions}})


# ─────────────────────────────────────────────
# 资产类型参数管理（折旧、可发放年限等）
# ─────────────────────────────────────────────

@app.route('/api/asset-type-params', methods=['GET'])
def get_asset_type_params():
    """获取所有资产类型参数"""
    conn = get_conn()
    rows = rows_to_list(conn.execute("SELECT * FROM asset_type_params ORDER BY type_code").fetchall())
    conn.close()
    return jsonify({'success': True, 'data': rows})


@app.route('/api/asset-type-params/<type_code>', methods=['GET'])
def get_asset_type_param(type_code):
    """获取单个资产类型参数"""
    conn = get_conn()
    row = conn.execute("SELECT * FROM asset_type_params WHERE type_code=?", (type_code,)).fetchone()
    conn.close()
    if row:
        return jsonify({'success': True, 'data': dict(row)})
    return jsonify({'success': False, 'error': '未找到该类型参数'}), 404


@app.route('/api/asset-type-params/<type_code>', methods=['PUT'])
def update_asset_type_param(type_code):
    """更新资产类型参数"""
    data = request.json or {}
    conn = get_conn()
    try:
        # 检查是否存在
        existing = conn.execute("SELECT id FROM asset_type_params WHERE type_code=?", (type_code,)).fetchone()
        if not existing:
            return jsonify({'success': False, 'error': '未找到该类型参数'}), 404
        
        # 更新字段
        depreciation_months = data.get('depreciation_months')
        residual_rate = data.get('residual_rate')
        max_issue_years = data.get('max_issue_years')
        
        conn.execute("""UPDATE asset_type_params 
                        SET depreciation_months=?, residual_rate=?, max_issue_years=?, updated_at=datetime('now','localtime')
                        WHERE type_code=?""",
                     (depreciation_months, residual_rate, max_issue_years, type_code))
        conn.commit()
        return jsonify({'success': True})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/asset-type-params/sync', methods=['POST'])
def sync_asset_type_params():
    """同步资产类型参数（根据asset_types表自动创建缺失的参数记录）"""
    conn = get_conn()
    try:
        # 获取所有资产类型
        types = conn.execute("SELECT code, name FROM asset_types WHERE active=1").fetchall()
        created = 0
        for t in types:
            # 检查是否已存在
            existing = conn.execute("SELECT id FROM asset_type_params WHERE type_code=?", (t['code'],)).fetchone()
            if not existing:
                # 创建默认参数
                conn.execute("""INSERT INTO asset_type_params (type_code, type_name, depreciation_months, residual_rate, max_issue_years)
                                VALUES (?, ?, 36, 0.05, 2)""",
                             (t['code'], t['name']))
                created += 1
        conn.commit()
        return jsonify({'success': True, 'created': created})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


# ─────────────────────────────────────────────
# 导出
# ─────────────────────────────────────────────

@app.route('/api/assets/export', methods=['GET'])
@login_required
def export_assets():
    """导出资产台账为 Excel(.xlsx)"""
    conn = get_conn()
    try:
        user = get_current_user()
        is_admin = user and user.get('role') in ('超管', '普通管理员', '管理员')
        if is_admin:
            rows = conn.execute("SELECT * FROM assets ORDER BY id").fetchall()
        else:
            # 普通用户仅能导出自己名下的资产（权限隔离）
            rows = conn.execute("SELECT * FROM assets WHERE assignee=? ORDER BY id",
                                (user.get('name'),)).fetchall()
        wb = Workbook(); ws = wb.active; ws.title = '资产台账'
        headers = ['资产编号','类目','品牌','型号','配置信息','SN','有线MAC','无线MAC',
                   '状态','存放地','关联人','KEY','关联账号密码','采购日期','保修到期','采购价格',
                   '账面价值','折旧月数','备注','入库时间','更新时间']
        ws.append(headers)
        # 安全加固（复审P2-4）：关联账号密码仅超管导出明文，其余角色脱敏为 ***
        is_superadmin = bool(user and user.get('role') == '超管')
        pmap = load_type_params_map(conn)
        for r in rows:
            d = attach_book_value(dict(r), pmap)
            pwd_cell = r['account_pwd'] if is_superadmin else ('***' if r['account_pwd'] else '')
            # 安全加固（复审2-⑥）：KEY 同属敏感凭据，与账号密码一致，仅超管导出明文，其余角色脱敏
            key_cell = r['key_info'] if is_superadmin else ('***' if r['key_info'] else '')
            ws.append([r['asset_no'], r['category'], r['brand'], r['model'], r['spec'], r['sn'],
                       r['mac_wired'], r['mac_wireless'], r['status'], r['location'], r['assignee'],
                       key_cell, pwd_cell, r['purchase_date'], r['warranty_expire'],
                       r['purchase_price'], d.get('book_value'), d.get('depreciation_months'),
                       r['notes'], r['created_at'], r['updated_at']])
        # 修复(四轮N5)：导出改内存 buffer 直接发送，不再落盘 export_tmp（含明文凭据的表格曾残留磁盘）
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name='资产台账.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    finally:
        conn.close()


ASSET_IMPORT_HEADERS = (
    '资产编号', '类目', '品牌', '型号', '配置信息', 'SN', '有线MAC', '无线MAC',
    '状态', '存放地', '关联人', 'KEY', '关联账号密码', '采购日期', '保修到期', '采购价格', '备注',
)
CONSUMABLE_IMPORT_HEADERS = (
    '耗材编号', '名称', '类目', '品牌', '型号', '单位', '库存数量', '预警阈值',
    '存放位置', '状态', '备注', '订单号',
)
ASSET_IMPORT_STATUSES = ('在库', '在用', '维修中', '已报废', '已出库')
CONSUMABLE_IMPORT_STATUSES = ('正常', '低库存', '停用')
IMPORT_ROLES = ('超管', '管理员', '普通管理员')


def _xlsx_import_template(sheet_title, headers, notes, download_name, status_col=None, status_choices=None):
    """仅含标准表头的空模板；第二张表为填写说明。"""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(list(headers))
    ws.freeze_panes = 'A2'
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(20, len(h) * 2 + 2))
    if status_col and status_choices:
        formula = '"' + ','.join(status_choices) + '"'
        dv = DataValidation(type='list', formula1=formula, allow_blank=False)
        dv.error = '请选择模板允许的取值'
        letter = get_column_letter(status_col)
        dv.add('%s2:%s1000' % (letter, letter))
        ws.add_data_validation(dv)
    ws2 = wb.create_sheet('填写说明')
    for i, line in enumerate(notes, 1):
        ws2.cell(row=i, column=1, value=line)
    ws2.column_dimensions['A'].width = 92
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True, download_name=download_name,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.route('/api/assets/template', methods=['GET'])
@login_required
def asset_template_download():
    """下载资产导入 Excel 模板：只有标准表头，无示例行。"""
    notes = [
        '请用本文件填写后，在资产台账点击「导入 Excel」。表头必须与模板完全一致，不可增删列或改列名。',
        '必填：资产编号、类目、状态。资产编号须唯一。',
        '状态只能是：在库 / 在用 / 维修中 / 已报废 / 已出库（请用下拉）。',
        '日期（采购日期、保修到期）格式 YYYY-MM-DD，如 2026-07-01；可留空。',
        '采购价格只填数字，不要货币符号或千分位逗号；可留空。',
        '第一张表不要改名；不要填满空行。',
    ]
    return _xlsx_import_template(
        '资产导入', ASSET_IMPORT_HEADERS, notes, '资产导入模板.xlsx',
        status_col=9, status_choices=ASSET_IMPORT_STATUSES,
    )


@app.route('/api/consumables/export', methods=['GET'])
@login_required
def export_consumables():
    """耗材台账导出为 Excel(.xlsx)"""
    conn = get_conn()
    try:
        rows = conn.execute("SELECT c.*, o.order_no as linked_order_no, o.order_date as linked_order_date FROM consumables c LEFT JOIN purchase_orders o ON c.order_id = o.id ORDER BY c.id").fetchall()
        wb = Workbook(); ws = wb.active; ws.title = '耗材台账'
        headers = ['耗材编号','名称','类目','品牌','型号','单位','库存数量','预警阈值','状态','存放位置','关联订单号','采购日期','备注','创建时间','更新时间']
        ws.append(headers)
        for r in rows:
            ws.append([r['item_no'], r['name'], r['category'], r['brand'], r['model'], r['unit'],
                       r['quantity'], r['min_stock'], r['status'], r['location'],
                       (r['linked_order_no'] or ''), (r['linked_order_date'] or ''),
                       r['notes'], r['created_at'], r['updated_at']])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name='耗材台账.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    finally:
        conn.close()


@app.route('/api/consumables/template', methods=['GET'])
@login_required
def consumable_template_download():
    """下载耗材导入 Excel 模板：只有标准表头，无示例行。"""
    notes = [
        '请用本文件填写后，在耗材台账点击「导入 Excel」。表头必须与模板完全一致，不可增删列或改列名。',
        '必填：耗材编号、名称。耗材编号须唯一。',
        '状态只能是：正常 / 低库存 / 停用；留空则按正常。',
        '库存数量、预警阈值为整数；可留空（库存默认 0，预警默认 10）。',
        '订单号须与系统中已有采购订单号一致，没有则留空。',
        '第一张表不要改名；不要填满空行。',
    ]
    return _xlsx_import_template(
        '耗材导入', CONSUMABLE_IMPORT_HEADERS, notes, '耗材导入模板.xlsx',
        status_col=10, status_choices=CONSUMABLE_IMPORT_STATUSES,
    )


# ─────────────────────────────────────────────
# 耗材管理 API
# ─────────────────────────────────────────────

@app.route('/api/consumables', methods=['GET'])
def list_consumables():
    conn = get_conn()
    q = request.args.get('q', '').strip()
    category = request.args.get('category', '')
    status   = request.args.get('status', '')
    page     = max(1, int(request.args.get('page', 1)))
    page_size = int(request.args.get('page_size', 20))

    sql = "SELECT c.*, o.order_no as linked_order_no, o.order_date as linked_order_date FROM consumables c LEFT JOIN purchase_orders o ON c.order_id = o.id WHERE 1=1"
    params = []
    if q:
        clause, p = fuzzy_match_clause(q, ['item_no', 'name', 'brand', 'model'])
        if clause:
            sql += " AND " + clause
            params += p
    if category:
        sql += " AND category=?"; params.append(category)
    if status:
        sql += " AND c.status=?"; params.append(status)

    total = conn.execute("SELECT COUNT(*) as cnt FROM " + sql.split("FROM",1)[1], params).fetchone()['cnt']
    sql += " ORDER BY id DESC LIMIT ? OFFSET ?"
    params += [page_size, (page-1)*page_size]
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return jsonify({'total': total, 'page': page, 'page_size': page_size, 'data': rows_to_list(rows)})


@app.route('/api/consumables', methods=['POST'])
def create_consumable():
    data = request.json or {}
    if not data.get('name'):
        return jsonify({'error': '耗材名称不能为空'}), 400

    item_no = data.get('item_no') or gen_consumable_no(data.get('category', ''))
    conn = get_conn()
    try:
        cur = conn.execute('''
            INSERT INTO consumables (item_no, name, category, brand, model, unit, quantity, min_stock, location, status, notes, order_id)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
        ''', (
            item_no, data.get('name',''), data.get('category',''), data.get('brand',''),
            data.get('model',''), data.get('unit','个'), data.get('quantity',0),
            data.get('min_stock',10), data.get('location',''), '正常', data.get('notes',''),
            data.get('order_id') or None
        ))
        new_id = cur.lastrowid
        # 自动写入库记录
        if data.get('quantity', 0) > 0:
            conn.execute('''
                INSERT INTO consumable_logs(item_id,item_no,log_type,quantity,operator,target_user,location,remark)
                VALUES(?,?,?,?,?,?,?,?)
            ''', (new_id, item_no, '入库', data.get('quantity',0),
                  get_current_name(), '', data.get('location',''), '初始入库'))
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({'error': f'耗材编号 {item_no} 已存在'}), 409
    conn.close()
    return jsonify({'id': new_id, 'item_no': item_no}), 201


@app.route('/api/consumables/<int:cid>', methods=['PUT'])
def update_consumable(cid):
    data = request.json or {}
    conn = get_conn()
    old = conn.execute("SELECT * FROM consumables WHERE id=?", (cid,)).fetchone()
    if not old:
        conn.close()
        return jsonify({'error': '耗材不存在'}), 404

    fields = ['name','category','brand','model','unit','min_stock','location','status','notes','order_id']
    updates = []
    params = []
    for f in fields:
        if f in data:
            updates.append(f'{f}=?')
            params.append(data[f])
    updates.append("updated_at=datetime('now','localtime')")
    params.append(cid)

    conn.execute(f"UPDATE consumables SET {','.join(updates)} WHERE id=?", params)
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/consumables/<int:cid>', methods=['DELETE'])
def delete_consumable(cid):
    conn = get_conn()
    conn.execute("DELETE FROM consumables WHERE id=?", (cid,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/consumables/stock/in', methods=['POST'])
def consumable_stock_in():
    """耗材入库"""
    data = request.json or {}
    cid = data.get('item_id')
    qty = int(data.get('quantity', 0))
    if not cid or qty <= 0:
        return jsonify({'error': '耗材ID和数量不能为空'}), 400

    order_id = data.get('order_id') or None
    order_item_id = data.get('order_item_id') or None

    conn = get_conn()
    item = conn.execute("SELECT * FROM consumables WHERE id=?", (cid,)).fetchone()
    if not item:
        conn.close()
        return jsonify({'error': '耗材不存在'}), 404

    # 更新库存
    new_qty = item['quantity'] + qty
    new_status = '正常' if new_qty >= item['min_stock'] else '低库存'
    conn.execute("UPDATE consumables SET quantity=?, status=?, updated_at=datetime('now','localtime') WHERE id=?",
                 (new_qty, new_status, cid))
    # 记录日志
    conn.execute('''
        INSERT INTO consumable_logs(item_id,item_no,log_type,quantity,operator,target_user,location,remark,order_id)
        VALUES(?,?,?,?,?,?,?,?,?)
    ''', (cid, item['item_no'], '入库', qty,
          get_current_name(), '', item['location'], data.get('remark',''), order_id))

    # 如果关联了订单明细，更新已到货数量
    if order_item_id:
        conn.execute(
            "UPDATE order_items SET arrived_qty = arrived_qty + ? WHERE id = ?",
            (qty, order_item_id)
        )
        _refresh_order_status(conn, order_id)

    conn.commit()
    conn.close()

    # 库存恢复正常通知（从低库存恢复到正常）
    if new_status == '正常' and item['status'] == '低库存':
        notify_lark('consumable_low_stock',
            f"✅ 【耗材库存恢复】\n耗材：{item['name']}（{item['item_no']}）\n当前库存：{new_qty} {item['unit']}\n预警阈值：{item['min_stock']} {item['unit']}\n库存已恢复正常\n经办人：{get_current_name()}\n时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")

    return jsonify({'ok': True, 'new_quantity': new_qty})


@app.route('/api/consumables/stock/out', methods=['POST'])
def consumable_stock_out():
    """耗材出库（领用）"""
    data = request.json or {}
    cid = data.get('item_id')
    qty = int(data.get('quantity', 0))
    if not cid or qty <= 0:
        return jsonify({'error': '耗材ID和数量不能为空'}), 400

    conn = get_conn()
    item = conn.execute("SELECT * FROM consumables WHERE id=?", (cid,)).fetchone()
    if not item:
        conn.close()
        return jsonify({'error': '耗材不存在'}), 404
    if item['quantity'] < qty:
        conn.close()
        return jsonify({'error': f'库存不足，当前库存 {item["quantity"]}'}), 400

    # 更新库存
    new_qty = item['quantity'] - qty
    new_status = '正常' if new_qty >= item['min_stock'] else '低库存'
    conn.execute("UPDATE consumables SET quantity=?, status=?, updated_at=datetime('now','localtime') WHERE id=?",
                 (new_qty, new_status, cid))
    # 记录日志
    conn.execute('''
        INSERT INTO consumable_logs(item_id,item_no,log_type,quantity,operator,target_user,location,remark)
        VALUES(?,?,?,?,?,?,?,?)
    ''', (cid, item['item_no'], '出库', qty,
          get_current_name(), data.get('target_user',''), item['location'], data.get('remark','')))
    conn.commit()
    conn.close()

    # 低库存通知
    if new_status == '低库存' and item['status'] != '低库存':
        notify_lark('consumable_low_stock',
            f"⚠️ 【耗材库存预警】\n耗材：{item['name']}（{item['item_no']}）\n当前库存：{new_qty} {item['unit']}\n预警阈值：{item['min_stock']} {item['unit']}\n领用人：{data.get('target_user','')}\n经办人：{get_current_name()}\n时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")

    return jsonify({'ok': True, 'new_quantity': new_qty})


@app.route('/api/consumables/logs', methods=['GET'])
def list_consumable_logs():
    conn = get_conn()
    item_id = request.args.get('item_id')
    log_type = request.args.get('log_type', '')
    target_user = request.args.get('target_user', '')
    page = max(1, int(request.args.get('page', 1)))
    page_size = int(request.args.get('page_size', 20))

    sql = "SELECT * FROM consumable_logs WHERE 1=1"
    params = []
    if item_id:
        sql += " AND item_id=?"; params.append(int(item_id))
    if log_type:
        sql += " AND log_type=?"; params.append(log_type)
    if target_user:
        sql += " AND target_user LIKE ?"; params.append(f'%{target_user}%')

    total = conn.execute(sql.replace("SELECT *", "SELECT COUNT(*) as cnt"), params).fetchone()['cnt']
    sql += " ORDER BY id DESC LIMIT ? OFFSET ?"
    params += [page_size, (page-1)*page_size]
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return jsonify({'total': total, 'data': rows_to_list(rows)})


@app.route('/api/consumables/by-assignee', methods=['GET'])
def consumables_by_assignee():
    """获取某人领用的耗材列表"""
    assignee = request.args.get('assignee', '').strip()
    if not assignee:
        return jsonify({'data': []})
    conn = get_conn()
    # 查询该人员领用的所有耗材（按耗材分组汇总）
    rows = conn.execute('''
        SELECT c.*, SUM(l.quantity) as total_out
        FROM consumables c
        JOIN consumable_logs l ON c.id = l.item_id
        WHERE l.target_user = ? AND l.log_type = '出库'
        GROUP BY c.id
        ORDER BY total_out DESC
    ''', (assignee,)).fetchall()
    conn.close()
    return jsonify({'data': rows_to_list(rows)})


@app.route('/api/consumables/by-user', methods=['GET'])
def consumables_by_user():
    """获取某人领用的耗材列表（别名接口）"""
    user = request.args.get('user', '').strip()
    if not user:
        return jsonify({'data': []})
    conn = get_conn()
    # 查询该人员领用的所有耗材（按耗材分组汇总）
    rows = conn.execute('''
        SELECT c.*, SUM(l.quantity) as total_out
        FROM consumables c
        JOIN consumable_logs l ON c.id = l.item_id
        WHERE l.target_user = ? AND l.log_type = '出库'
        GROUP BY c.id
        ORDER BY total_out DESC
    ''', (user,)).fetchall()
    conn.close()
    return jsonify({'data': rows_to_list(rows)})


@app.route('/api/consumables/low-stock', methods=['GET'])
def list_low_stock_consumables():
    """获取低库存耗材列表"""
    conn = get_conn()
    rows = conn.execute('''
        SELECT * FROM consumables
        WHERE status = '低库存' OR quantity < min_stock
        ORDER BY (min_stock - quantity) DESC
    ''').fetchall()
    conn.close()
    return jsonify({'data': rows_to_list(rows)})


# ─────────────────────────────────────────────
# 枚举/元数据
# ─────────────────────────────────────────────

# ─────────────────────────────────────────────
# 资产处理（报废、捐赠、出售等）
# ─────────────────────────────────────────────

@app.route('/api/dispose', methods=['POST'])
def create_dispose():
    """创建资产处理申请（提交审批，超管确认后执行）"""
    data = request.json or {}
    aid = data.get('asset_id')
    if not aid:
        return jsonify({'error': '资产ID不能为空'}), 400

    conn = get_conn()
    asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not asset:
        return jsonify({'error': '资产不存在'}), 404

    payload = {
        'dispose_type': data.get('dispose_type', '报废'),
        'reason': data.get('reason', ''),
        'operator': get_current_name(),
        'approved_by': data.get('approved_by', ''),
        'dispose_date': data.get('dispose_date', ''),
        'notes': data.get('notes', '')
    }
    try:
        app_id = create_operation_application('dispose', [aid], payload, applicant=get_current_name())
        return jsonify({'success': True, 'app_id': app_id, 'message': '已提交资产处置申请，请等待超管审批'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/dispose', methods=['GET'])
def list_dispose():
    """处理记录：待审批的处置申请 + 审批通过后的 dispose_logs。"""
    conn = get_conn()
    try:
        dispose_type = request.args.get('dispose_type', '')
        page = max(1, int(request.args.get('page', 1)))
        page_size = max(1, min(200, int(request.args.get('page_size', 20))))
        merged = []

        ops = conn.execute(
            "SELECT * FROM operation_applications WHERE op_type='dispose' "
            "AND status IN ('pending','processing','rejected') ORDER BY id DESC"
        ).fetchall()
        for op in ops:
            payload = {}
            try:
                payload = json.loads(op['payload'] or '{}') if op['payload'] else {}
            except (TypeError, json.JSONDecodeError):
                payload = {}
            if not isinstance(payload, dict):
                payload = {}
            dt = payload.get('dispose_type', '报废')
            if dispose_type and dt != dispose_type:
                continue
            nos, aids = [], []
            try:
                nos = json.loads(op['asset_nos'] or '[]') if op['asset_nos'] else []
            except (TypeError, json.JSONDecodeError):
                nos = []
            try:
                aids = json.loads(op['asset_ids'] or '[]') if op['asset_ids'] else []
            except (TypeError, json.JSONDecodeError):
                aids = []
            cat = brand = model = ''
            if aids:
                a = conn.execute(
                    "SELECT category, brand, model FROM assets WHERE id=?", (aids[0],)
                ).fetchone()
                if a:
                    cat, brand, model = a['category'] or '', a['brand'] or '', a['model'] or ''
            status = op['status'] or 'pending'
            status_zh = {'pending': '待审批', 'processing': '审批中', 'rejected': '已驳回'}.get(status, status)
            merged.append({
                'id': 'pending-%s' % op['id'],
                'pending': status != 'rejected',
                'record_status': status_zh,
                'asset_id': aids[0] if aids else None,
                'asset_no': nos[0] if nos else '',
                'category': cat,
                'brand': brand,
                'model': model,
                'dispose_type': dt,
                'reason': payload.get('reason', ''),
                'operator': payload.get('operator') or op['applicant'] or '',
                'approved_by': '',
                'dispose_date': payload.get('dispose_date', ''),
                'notes': payload.get('notes', ''),
                'created_at': op['apply_time'] or '',
                'photos': [],
            })

        sql = ("SELECT d.*, a.category, a.brand, a.model FROM dispose_logs d "
               "LEFT JOIN assets a ON d.asset_id=a.id WHERE 1=1")
        params = []
        if dispose_type:
            sql += " AND d.dispose_type=?"
            params.append(dispose_type)
        sql += " ORDER BY d.id DESC"
        for row in rows_to_list(conn.execute(sql, params).fetchall()):
            row['pending'] = False
            row['record_status'] = '已完成'
            photos = conn.execute(
                "SELECT * FROM photos WHERE ref_type='dispose' AND ref_id=?",
                (row['id'],)
            ).fetchall()
            row['photos'] = rows_to_list(photos)
            merged.append(row)

        total = len(merged)
        start = (page - 1) * page_size
        return jsonify({'total': total, 'data': merged[start:start + page_size]})
    finally:
        conn.close()


@app.route('/api/dispose/<int:dispose_id>/photos', methods=['POST'])
def upload_dispose_photo(dispose_id):
    """上传资产处理照片"""
    if 'file' not in request.files:
        return jsonify({'error': '没有文件'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': '没有选择文件'}), 400

    filename = secure_filename(file.filename)
    timestamp = _rand_suffix()  # 安全加固 P2-4：随机段替代可推算时间戳，防附件URL枚举
    new_filename = f"dispose_{dispose_id}_{timestamp}_{filename}"
    file_path = os.path.join(UPLOAD_DIR, new_filename)
    file.save(file_path)

    conn = get_conn()
    cur = conn.execute('''
        INSERT INTO photos(ref_type, ref_id, file_name, file_path, file_size, mime_type, uploaded_by)
        VALUES(?,?,?,?,?,?,?)
    ''', ('dispose', dispose_id, new_filename, file_path, os.path.getsize(file_path),
          file.content_type, request.form.get('uploaded_by', '系统')))
    photo_id = cur.lastrowid
    conn.commit()
    conn.close()
    return jsonify({'id': photo_id, 'file_name': new_filename}), 201


# ─────────────────────────────────────────────
# 批量导入（严格按模板表头的 .xlsx）
# ─────────────────────────────────────────────

def _cell_text(v):
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return '1' if v else '0'
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return str(v)
    return str(v).strip()


def _load_xlsx_workbook(upload):
    """从上传对象读出完整字节再打开，避免流指针不在开头导致「不是 zip」。"""
    try:
        if hasattr(upload, 'stream') and upload.stream is not None:
            try:
                upload.stream.seek(0)
            except Exception:
                pass
        if hasattr(upload, 'read'):
            raw = upload.read()
        elif hasattr(upload, 'stream'):
            raw = upload.stream.read()
        else:
            raw = upload
        if isinstance(raw, str):
            raw = raw.encode('utf-8', errors='replace')
    except Exception:
        raw = b''
    if not raw:
        raise ValueError('没有收到表格内容。请先关闭 Excel，再重新选择填好的文件上传。')
    if raw.startswith(b'\xd0\xcf\x11\xe0'):
        raise ValueError('这是旧版 Excel（.xls）。请在 Excel 里点「另存为」，选「Excel 工作簿 (*.xlsx)」后再上传。')
    if not raw.startswith(b'PK'):
        raise ValueError('这个文件不是 Excel 工作簿。请用本系统「下载模板」，用 Excel 打开填写后直接保存为 .xlsx 再上传。')
    buf = io.BytesIO(raw)
    try:
        return load_workbook(buf, data_only=True)
    except Exception:
        buf.seek(0)
        try:
            return load_workbook(buf, data_only=False)
        except Exception:
            raise ValueError(
                '表格读不出来。请用 Excel 打开后点「另存为」→「Excel 工作簿 (*.xlsx)」，'
                '不要选「严格 Open XML」，也不要加密或设打开密码。'
            )


def _parse_xlsx_strict(file, expected_headers):
    fname = (file.filename or '').lower()
    if fname and not fname.endswith('.xlsx'):
        raise ValueError('请上传 .xlsx 文件（Excel 工作簿）。请先点「下载模板」填好后再上传。')
    wb = _load_xlsx_workbook(file)
    expected = list(expected_headers)
    wsheet = None
    for name in ('资产导入', '耗材导入'):
        if name in wb.sheetnames:
            wsheet = wb[name]
            break
    if wsheet is None:
        wsheet = wb.active
    grid = list(wsheet.iter_rows(values_only=True))
    if not grid:
        raise ValueError('表格是空的。请用下载的模板填写至少一行数据。')
    headers = [_cell_text(h) for h in grid[0]]
    while headers and headers[-1] == '':
        headers.pop()
    if headers != expected:
        raise ValueError('表格第一行和模板不一样。请重新下载模板，不要改第一行的栏目名称，填好后再上传。')
    records = []
    for excel_row, row in enumerate(grid[1:], start=2):
        cells = list(row) if row is not None else []
        if all(_cell_text(c) == '' for c in cells):
            continue
        rec = {}
        for i, h in enumerate(expected):
            rec[h] = cells[i] if i < len(cells) else None
        records.append((excel_row, rec))
    return records


def _parse_ymd(v, field):
    s = _cell_text(v)
    if not s:
        return ''
    try:
        datetime.strptime(s[:10], '%Y-%m-%d')
    except ValueError:
        raise ValueError('%s须为 YYYY-MM-DD' % field)
    return s[:10]


def _parse_money(v, field):
    if v is None or _cell_text(v) == '':
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = _cell_text(v).replace(',', '')
    try:
        return float(s)
    except ValueError:
        raise ValueError('%s须为数字，不要货币符号' % field)


def _parse_int_cell(v, field, default=None):
    if v is None or _cell_text(v) == '':
        return default
    if isinstance(v, bool):
        raise ValueError('%s须为整数' % field)
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v == int(v):
        return int(v)
    s = _cell_text(v)
    if not re.fullmatch(r'-?\d+', s):
        raise ValueError('%s须为整数' % field)
    return int(s)


def _import_file_or_error():
    user = get_current_user()
    if not user or user.get('role') not in IMPORT_ROLES:
        return None, (jsonify({'success': False, 'error': '权限不足，仅超管/管理员/普通管理员可导入'}), 403)
    if 'file' not in request.files:
        return None, (jsonify({'success': False, 'error': '没有文件'}), 400)
    f = request.files['file']
    if not f or f.filename == '':
        return None, (jsonify({'success': False, 'error': '没有选择文件'}), 400)
    return f, None


@app.route('/api/assets/import', methods=['POST'])
@login_required
def import_assets():
    """按资产导入模板严格导入 .xlsx；任一档格式错误则整批不写入。"""
    file, err = _import_file_or_error()
    if err:
        return err
    try:
        records = _parse_xlsx_strict(file, ASSET_IMPORT_HEADERS)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e), 'imported': 0, 'errors': [str(e)]}), 400
    if not records:
        return jsonify({'success': False, 'error': '文件只有表头，没有数据行', 'imported': 0, 'errors': []}), 400

    conn = get_conn()
    errors = []
    prepared = []
    seen = set()
    try:
        for excel_row, rec in records:
            try:
                asset_no = _cell_text(rec.get('资产编号'))
                category = _cell_text(rec.get('类目'))
                status = _cell_text(rec.get('状态'))
                if not asset_no:
                    raise ValueError('资产编号不能为空')
                if not category:
                    raise ValueError('类目不能为空')
                if status not in ASSET_IMPORT_STATUSES:
                    raise ValueError('状态必须是：' + ' / '.join(ASSET_IMPORT_STATUSES))
                if asset_no in seen:
                    raise ValueError('文件内资产编号重复')
                seen.add(asset_no)
                exists = conn.execute("SELECT 1 FROM assets WHERE asset_no=?", (asset_no,)).fetchone()
                if exists:
                    raise ValueError('资产编号已存在')
                prepared.append((
                    asset_no, category, _cell_text(rec.get('品牌')), _cell_text(rec.get('型号')),
                    _cell_text(rec.get('配置信息')), _cell_text(rec.get('SN')),
                    _cell_text(rec.get('有线MAC')), _cell_text(rec.get('无线MAC')),
                    status, _cell_text(rec.get('存放地')), _cell_text(rec.get('关联人')),
                    _cell_text(rec.get('KEY')), _cell_text(rec.get('关联账号密码')),
                    _parse_ymd(rec.get('采购日期'), '采购日期'),
                    _parse_ymd(rec.get('保修到期'), '保修到期'),
                    _parse_money(rec.get('采购价格'), '采购价格'),
                    _cell_text(rec.get('备注')),
                ))
            except Exception as e:
                errors.append('第 %s 行：%s' % (excel_row, e))
        if errors:
            return jsonify({
                'success': False, 'imported': 0, 'errors': errors,
                'error': '共 %s 处格式问题，未导入任何记录' % len(errors),
            }), 400
        for row in prepared:
            conn.execute('''
                INSERT INTO assets(asset_no, category, brand, model, spec, sn, mac_wired, mac_wireless,
                    status, location, assignee, key_info, account_pwd, purchase_date, warranty_expire, purchase_price, notes)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ''', row)
        conn.commit()
        return jsonify({'success': True, 'imported': len(prepared), 'errors': []})
    finally:
        conn.close()


@app.route('/api/consumables/import', methods=['POST'])
@login_required
def import_consumables():
    """按耗材导入模板严格导入 .xlsx；任一档格式错误则整批不写入。"""
    file, err = _import_file_or_error()
    if err:
        return err
    try:
        records = _parse_xlsx_strict(file, CONSUMABLE_IMPORT_HEADERS)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e), 'imported': 0, 'errors': [str(e)]}), 400
    if not records:
        return jsonify({'success': False, 'error': '文件只有表头，没有数据行', 'imported': 0, 'errors': []}), 400

    conn = get_conn()
    errors = []
    prepared = []
    seen = set()
    try:
        for excel_row, rec in records:
            try:
                item_no = _cell_text(rec.get('耗材编号'))
                name = _cell_text(rec.get('名称'))
                if not item_no:
                    raise ValueError('耗材编号不能为空')
                if not name:
                    raise ValueError('名称不能为空')
                status = _cell_text(rec.get('状态')) or '正常'
                if status not in CONSUMABLE_IMPORT_STATUSES:
                    raise ValueError('状态必须是：' + ' / '.join(CONSUMABLE_IMPORT_STATUSES))
                if item_no in seen:
                    raise ValueError('文件内耗材编号重复')
                seen.add(item_no)
                exists = conn.execute("SELECT 1 FROM consumables WHERE item_no=?", (item_no,)).fetchone()
                if exists:
                    raise ValueError('耗材编号已存在')
                order_id_val = None
                order_no_val = _cell_text(rec.get('订单号'))
                if order_no_val:
                    orow = conn.execute(
                        "SELECT id FROM purchase_orders WHERE order_no=?", (order_no_val,)
                    ).fetchone()
                    if not orow:
                        raise ValueError('订单号在系统中不存在')
                    order_id_val = orow['id']
                qty = _parse_int_cell(rec.get('库存数量'), '库存数量', 0)
                min_stock = _parse_int_cell(rec.get('预警阈值'), '预警阈值', 10)
                if qty is None or qty < 0:
                    raise ValueError('库存数量须为大于等于 0 的整数')
                if min_stock is None or min_stock < 0:
                    raise ValueError('预警阈值须为大于等于 0 的整数')
                prepared.append((
                    item_no, name, _cell_text(rec.get('类目')), _cell_text(rec.get('品牌')),
                    _cell_text(rec.get('型号')), _cell_text(rec.get('单位')) or '个',
                    qty, min_stock, _cell_text(rec.get('存放位置')), status,
                    _cell_text(rec.get('备注')), order_id_val,
                ))
            except Exception as e:
                errors.append('第 %s 行：%s' % (excel_row, e))
        if errors:
            return jsonify({
                'success': False, 'imported': 0, 'errors': errors,
                'error': '共 %s 处格式问题，未导入任何记录' % len(errors),
            }), 400
        for row in prepared:
            conn.execute('''
                INSERT INTO consumables(item_no, name, category, brand, model, unit, quantity, min_stock, location, status, notes, order_id)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
            ''', row)
        conn.commit()
        return jsonify({'success': True, 'imported': len(prepared), 'errors': []})
    finally:
        conn.close()


# ─────────────────────────────────────────────
# 快速操作页面API（支持照片上传）
# ─────────────────────────────────────────────

@app.route('/api/quick/stock', methods=['POST'])
def quick_stock():
    """快速出入库（带照片）"""
    data = request.form.to_dict()
    log_type = data.get('log_type', '出库')
    aid = data.get('asset_id')
    operator = get_current_name()  # 修复(四轮N4)：operator 强制取服务端登录身份，拒绝前端传值伪造

    if not aid:
        return jsonify({'error': '资产ID不能为空'}), 400

    conn = get_conn()
    asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not asset:
        conn.close()
        return jsonify({'error': '资产不存在'}), 404

    # 执行出入库操作
    if log_type == '出库':
        if asset['status'] not in ('在库',):
            conn.close()
            return jsonify({'error': f'资产当前状态为【{asset["status"]}】，无法出库'}), 400
        conn.execute('''UPDATE assets SET status='在用', assignee=?, location=?,
                        updated_at=datetime('now','localtime') WHERE id=?''',
                     (data.get('target_user',''), data.get('to_location',asset['location']), aid))
    elif log_type == '归还':
        conn.execute('''UPDATE assets SET status='在库', assignee='', location=?,
                        updated_at=datetime('now','localtime') WHERE id=?''',
                     (data.get('to_location',asset['location']), aid))
    elif log_type == '调拨':
        conn.execute('''UPDATE assets SET assignee=?, location=?,
                        updated_at=datetime('now','localtime') WHERE id=?''',
                     (data.get('target_user',asset['assignee']), data.get('to_location',asset['location']), aid))

    # 记录日志
    cur = conn.execute('''
        INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,target_user,from_location,to_location,remark)
        VALUES(?,?,?,?,?,?,?,?)
    ''', (aid, asset['asset_no'], log_type, operator,
          data.get('target_user',''), asset['location'], data.get('to_location',asset['location']), data.get('remark','')))
    log_id = cur.lastrowid
    conn.commit()
    conn.close()

    # 处理上传的照片（照片保存失败时返回明确 JSON 错误，避免前端 res.json() 解析 500 HTML 而误报"网络错误"）
    photo_ids = []
    if 'photos' in request.files:
        photos = request.files.getlist('photos')
        for photo in photos:
            if not photo.filename:
                continue
            try:
                filename = secure_filename(photo.filename)
                timestamp = _rand_suffix()  # 安全加固 P2-4：随机段替代可推算时间戳，防附件URL枚举
                new_filename = f"stock_{log_id}_{timestamp}_{filename}"
                file_path = os.path.join(UPLOAD_DIR, new_filename)
                photo.save(file_path)

                conn = get_conn()
                cur = conn.execute('''
                    INSERT INTO photos(ref_type, ref_id, file_name, file_path, file_size, mime_type, uploaded_by)
                    VALUES(?,?,?,?,?,?,?)
                ''', ('stock', log_id, new_filename, file_path, os.path.getsize(file_path),
                      photo.content_type, operator))
                photo_ids.append(cur.lastrowid)
                conn.commit()
                conn.close()
            except Exception as e:
                return jsonify({'error': f'照片保存失败: {e}'}), 500

    try:
        notify_lark('stock',
            f"📦 【资产{log_type}】\n编号：{asset['asset_no']}\n类目：{asset['category']}\n品牌型号：{asset['brand']} {asset['model']}\n操作人：{operator}\n时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    except Exception:
        pass  # 通知失败不阻断出入库

    return jsonify({'ok': True, 'log_id': log_id, 'photo_ids': photo_ids})


@app.route('/api/quick/repair', methods=['POST'])
def quick_repair():
    """快速报修（带照片）—— 提交审批，超管确认后执行"""
    data = request.form.to_dict()
    aid = data.get('asset_id')
    if not aid:
        return jsonify({'error': '资产ID不能为空'}), 400

    conn = get_conn()
    asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not asset:
        return jsonify({'error': '资产不存在'}), 404

    payload = {
        'fault_desc': data.get('fault_desc', ''),
        'repair_vendor': data.get('repair_vendor', ''),
        'submit_by': data.get('submit_by', ''),
        'repair_notes': data.get('repair_notes', '')
    }
    try:
        app_id = create_operation_application('repair', [aid], payload, applicant=data.get('submit_by'))
        photo_ids = save_operation_proof(app_id, request.files.getlist('photos'), data.get('submit_by', '系统'))
        return jsonify({'success': True, 'app_id': app_id, 'photo_ids': photo_ids, 'message': '已提交报修申请，请等待超管审批'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/quick/dispose', methods=['POST'])
def quick_dispose():
    """快速资产处理（带照片）—— 提交审批，超管确认并验证处置证明后执行"""
    data = request.form.to_dict()
    aid = data.get('asset_id')
    if not aid:
        return jsonify({'error': '资产ID不能为空'}), 400

    conn = get_conn()
    asset = conn.execute("SELECT * FROM assets WHERE id=?", (aid,)).fetchone()
    if not asset:
        return jsonify({'error': '资产不存在'}), 404

    dispose_type = data.get('dispose_type', '报废')
    # 处置证明要求：丢失→派出所/人力/财务证明；报废→安全处理证明（二者必传）
    PROOF_REQUIRED = {'丢失', '报废'}
    asset_photos = request.files.getlist('photos')
    proof_photos = request.files.getlist('proof_photos')
    if dispose_type in PROOF_REQUIRED and not proof_photos:
        return jsonify({'error': f'【{dispose_type}】必须上传处置证明（{"派出所/人力/财务证明" if dispose_type == "丢失" else "安全处理证明"}）'}), 400

    payload = {
        'dispose_type': dispose_type,
        'reason': data.get('reason', ''),
        'operator': get_current_name(),
        'approved_by': data.get('approved_by', ''),
        'dispose_date': data.get('dispose_date', ''),
        'notes': data.get('notes', ''),
        'proof_required': dispose_type in PROOF_REQUIRED
    }
    try:
        app_id = create_operation_application('dispose', [aid], payload, applicant=get_current_name())
        # 资产现状照片 + 处置证明 均作为审批单证明暂存，审批通过后迁移到 dispose_log
        saved = save_operation_proof(app_id, asset_photos + proof_photos, get_current_name())
        return jsonify({'success': True, 'app_id': app_id, 'photo_ids': saved, 'message': '已提交资产处置申请，请等待超管审批'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/meta', methods=['GET'])
def get_meta():
    conn = get_conn()
    cats      = [r[0] for r in conn.execute("SELECT DISTINCT category FROM assets ORDER BY category").fetchall()]
    locations = [r[0] for r in conn.execute("SELECT DISTINCT location FROM assets WHERE location!='' UNION SELECT DISTINCT location FROM employees WHERE location IS NOT NULL AND location!='' ORDER BY location").fetchall()]
    statuses  = ['在库','在用','维修中','已报废','已出库']
    # 耗材类目
    cons_cats = [r[0] for r in conn.execute("SELECT DISTINCT category FROM consumables ORDER BY category").fetchall()]
    cons_statuses = ['正常','低库存','停用']
    # 资产处理类型
    dispose_types = ['报废', '捐赠', '出售', '丢失', '其他']
    
    # 资产编号配置
    subjects = rows_to_list(conn.execute("SELECT code, name, region FROM asset_subjects WHERE active=1 ORDER BY sort_order").fetchall())
    asset_types = rows_to_list(conn.execute("SELECT code, name FROM asset_types WHERE active=1 ORDER BY sort_order").fetchall())
    asset_regions = rows_to_list(conn.execute("SELECT code, name FROM asset_regions WHERE active=1 ORDER BY sort_order").fetchall())
    
    # 字段选项（用于下拉框颜色）
    field_options = rows_to_list(conn.execute("SELECT field_type, field_name, color FROM field_options WHERE active=1").fetchall())
    
    # 用户自定义字段：field_value 优先，没有则用 field_name 自身
    user_fields_raw = conn.execute("SELECT field_code, field_name, field_value FROM field_options WHERE field_type='user' AND active=1 ORDER BY sort_order, id").fetchall()
    user_fields = []
    seen_codes = set()
    for r in user_fields_raw:
        code = r['field_code']
        if code in seen_codes: continue
        seen_codes.add(code)
        opts = (r['field_value'] or r['field_name'] or '').split(',')
        opts = [x.strip() for x in opts if x.strip()]
        user_fields.append({
            'field_code': code,
            'field_name': r['field_name'],
            'field_value': ','.join(opts)
        })
    
    conn.close()
    return jsonify({
        'categories': cats or ['显示器','笔记本电脑','台式电脑','打印机','网络设备','软件','其他'],
        'locations':  locations or ['北京','深圳','广州'],
        'statuses':   statuses,
        'consumable_categories': cons_cats or ['墨盒','硒鼓','打印纸','网线','电源线','鼠标','键盘','其他'],
        'consumable_statuses': cons_statuses,
        'dispose_types': dispose_types,
        # 资产编号配置
        'asset_subjects': subjects,
        'asset_types': asset_types,
        'asset_regions': asset_regions,
        # 字段选项颜色
        'field_options': field_options,
        'user_fields': user_fields,
    })


# ═══════════════════════════════════════════════════════════════════════════
# 操作日志管理
# ═══════════════════════════════════════════════════════════════════════════

@app.route('/api/logs', methods=['GET'])
@login_required
def get_action_logs():
    """获取操作日志列表（供超管/领导审计）"""
    page = int(request.args.get('page', 1))
    page_size = int(request.args.get('page_size', 50))
    module = request.args.get('module', '')      # 筛选模块
    action_type = request.args.get('action_type', '')  # 筛选操作类型
    operator = request.args.get('operator', '')   # 筛选操作人
    keyword = request.args.get('keyword', '')      # 搜索关键词
    start_date = request.args.get('start_date', '')
    end_date = request.args.get('end_date', '')
    offset = (page - 1) * page_size
    
    conn = get_conn()
    try:
        # 构建查询条件
        conditions = []
        params = []
        
        if module:
            conditions.append("module = ?")
            params.append(module)
        
        if action_type:
            conditions.append("action_type = ?")
            params.append(action_type)
        
        if operator:
            conditions.append("operator LIKE ?")
            params.append(f"%{operator}%")
        
        if keyword:
            conditions.append("(target_name LIKE ? OR target_no LIKE ? OR detail LIKE ?)")
            params.extend([f"%{keyword}%", f"%{keyword}%", f"%{keyword}%"])
        
        if start_date:
            conditions.append("created_at >= ?")
            params.append(start_date)
        
        if end_date:
            conditions.append("created_at <= ?")
            params.append(end_date + " 23:59:59")
        
        where_clause = " AND ".join(conditions) if conditions else "1=1"
        
        # 获取总数
        total = conn.execute(
            f"SELECT COUNT(*) as cnt FROM action_logs WHERE {where_clause}",
            params
        ).fetchone()['cnt']
        
        # 获取分页数据
        rows = conn.execute(
            f"""SELECT * FROM action_logs 
                WHERE {where_clause}
                ORDER BY created_at DESC 
                LIMIT ? OFFSET ?""",
            params + [page_size, offset]
        ).fetchall()
        
        logs = []
        for row in rows:
            log = dict(row)
            # 解析 JSON 字段
            if log.get('detail'):
                try:
                    log['detail'] = json.loads(log['detail'])
                except:
                    pass
            if log.get('before_data'):
                try:
                    log['before_data'] = json.loads(log['before_data'])
                except:
                    pass
            if log.get('after_data'):
                try:
                    log['after_data'] = json.loads(log['after_data'])
                except:
                    pass
            logs.append(log)
        
        return jsonify({
            'success': True,
            'data': logs,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    finally:
        conn.close()


@app.route('/api/logs/<int:log_id>', methods=['GET'])
@login_required
def get_action_log_detail(log_id):
    """获取操作日志详情"""
    conn = get_conn()
    try:
        row = conn.execute("SELECT * FROM action_logs WHERE id = ?", (log_id,)).fetchone()
        if not row:
            return jsonify({'success': False, 'error': '日志不存在'}), 404
        
        log = dict(row)
        # 解析 JSON 字段
        for field in ['detail', 'before_data', 'after_data']:
            if log.get(field):
                try:
                    log[field] = json.loads(log[field])
                except:
                    pass
        
        return jsonify({'success': True, 'data': log})
    finally:
        conn.close()


# ═══════════════════════════════════════════════════════════════════════════
# 订单修改申请
# ═══════════════════════════════════════════════════════════════════════════

@app.route('/api/orders/<int:order_id>/modifications', methods=['GET'])
@login_required
def get_order_modifications(order_id):
    """获取订单的修改申请列表"""
    conn = get_conn()
    try:
        rows = conn.execute(
            """SELECT * FROM order_modifications 
               WHERE order_id = ? 
               ORDER BY created_at DESC""",
            (order_id,)
        ).fetchall()
        
        mods = []
        for row in rows:
            mod = dict(row)
            if mod.get('modifications'):
                try:
                    mod['modifications'] = json.loads(mod['modifications'])
                except:
                    pass
            if mod.get('attachments'):
                try:
                    mod['attachments'] = json.loads(mod['attachments'])
                except:
                    pass
            mods.append(mod)
        
        return jsonify({'success': True, 'data': mods})
    finally:
        conn.close()


@app.route('/api/orders/<int:order_id>/modification', methods=['POST'])
@login_required
def create_order_modification(order_id):
    """提交订单修改申请
    - 普通用户：必须上传附件才能提交申请
    - 超管：直接修改，无需审批
    """
    conn = get_conn()
    try:
        # 获取订单信息
        order = conn.execute("SELECT * FROM purchase_orders WHERE id = ?", (order_id,)).fetchone()
        if not order:
            return jsonify({'success': False, 'error': '订单不存在'}), 404
        
        # 获取当前用户
        user = get_current_user()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404
        
        is_admin = user.get('role') == '超管'
        
        # 处理附件（如果有文件上传）
        attachments = []
        if 'attachments' in request.files:
            for file in request.files.getlist('attachments'):
                if file.filename:
                    filename = secure_filename(file.filename)
                    timestamp = _rand_suffix()  # 安全加固 P2-4：随机段替代可推算时间戳，防附件URL枚举
                    new_filename = f"order_mod_{order_id}_{timestamp}_{filename}"
                    file_path = os.path.join(UPLOAD_DIR, new_filename)
                    file.save(file_path)
                    attachments.append({
                        'file_name': new_filename,
                        'original_name': filename,
                        'file_path': file_path,
                        'file_size': os.path.getsize(file_path)
                    })
        
        # 获取修改内容
        data = request.json or {}
        modifications = data.get('modifications', {})
        reason = data.get('reason', '')
        
        # 如果是超管，直接执行修改
        if is_admin:
            before_data = dict(order)
            fields, values = [], []
            for field, change in modifications.items():
                new_value = change.get('new', order.get(field))
                if field in ['title', 'supplier', 'order_date', 'expected_date', 'purchase_by', 'invoice_no', 'contract_no', 'notes', 'expense_order_no']:
                    fields.append(f"{field} = ?")
                    values.append(new_value)
                elif field == 'total_amount':
                    fields.append("total_amount = ?")
                    values.append(float(new_value) if new_value else 0)
            
            if fields:
                values.append(datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
                values.append(order_id)
                conn.execute(
                    f"UPDATE purchase_orders SET {', '.join(fields)}, updated_at = ? WHERE id = ?",
                    values
                )
                
                # 更新明细
                items_data = data.get('items', [])
                for item_change in items_data:
                    item_id = item_change.get('id')
                    if item_id:
                        item_fields, item_values = [], []
                        # 安全加固（P2-2）：明细字段列名白名单，杜绝 SQL 列名注入（field 取自用户 JSON key）
                        _ITEM_UPDATABLE = {'item_type','item_name','category','brand','model','quantity','unit_price','arrived_qty','notes'}
                        for field, change in item_change.items():
                            if field in _ITEM_UPDATABLE and change.get('new') != change.get('old'):
                                item_fields.append(f"{field} = ?")
                                item_values.append(change.get('new'))
                        if item_fields:
                            item_values.append(item_id)
                            conn.execute(
                                f"UPDATE order_items SET {', '.join(item_fields)} WHERE id = ?",
                                item_values
                            )
                
                # 添加附件关联
                for att in attachments:
                    conn.execute('''
                        INSERT INTO photos(ref_type, ref_id, file_name, file_path, file_size, uploaded_by)
                        VALUES(?, ?, ?, ?, ?, ?)
                    ''', ('order_modification', order_id, att['file_name'], att['file_path'], att['file_size'], user['name']))
                
                # 记录日志
                log_action('update', 'order', 'purchase_order', order_id, order['order_no'], order['title'],
                          detail={'reason': reason, 'attachments': [a['original_name'] for a in attachments]},
                          before_data=before_data,
                          after_data={k: v for k, v in dict(conn.execute("SELECT * FROM purchase_orders WHERE id = ?", (order_id,)).fetchone()).items() 
                                     if k in modifications})
                
                conn.commit()
            
            return jsonify({'success': True, 'message': '订单已直接修改', 'is_admin': True})
        
        # 普通用户：创建修改申请
        # 至少要有一个附件
        if not attachments and not data.get('attachment_urls'):
            return jsonify({'success': False, 'error': '修改订单必须上传附件'}), 400
        
        # 记录原始附件URL
        attachment_urls = data.get('attachment_urls', [])
        
        cursor = conn.execute('''
            INSERT INTO order_modifications 
            (order_id, order_no, title, modifications, reason, attachments, apply_by, apply_by_role)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            order_id, order['order_no'], order['title'],
            json.dumps(modifications, ensure_ascii=False),
            reason,
            json.dumps(attachments + [{'file_name': url} for url in attachment_urls], ensure_ascii=False),
            user['name'], user['role']
        ))
        modification_id = cursor.lastrowid
        conn.commit()
        
        # 记录日志
        log_action('apply', 'order', 'purchase_order', order_id, order['order_no'], order['title'],
                  detail={'modification_id': modification_id, 'reason': reason, 'attachments_count': len(attachments)})
        
        return jsonify({
            'success': True, 
            'message': '修改申请已提交，请等待审批',
            'modification_id': modification_id
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/order-modifications', methods=['GET'])
@login_required
def get_pending_modifications():
    """获取待审批的订单修改申请列表"""
    conn = get_conn()
    try:
        # 只有管理员能看到
        user = get_current_user()
        if not user or user.get('role') != '超管':
            return jsonify({'success': False, 'error': '需要管理员权限'}), 403
        
        page = int(request.args.get('page', 1))
        page_size = int(request.args.get('page_size', 20))
        offset = (page - 1) * page_size
        
        total = conn.execute(
            "SELECT COUNT(*) as cnt FROM order_modifications WHERE status = 'pending'"
        ).fetchone()['cnt']
        
        rows = conn.execute(
            """SELECT * FROM order_modifications 
               WHERE status = 'pending'
               ORDER BY created_at DESC 
               LIMIT ? OFFSET ?""",
            (page_size, offset)
        ).fetchall()
        
        mods = []
        for row in rows:
            mod = dict(row)
            if mod.get('modifications'):
                try:
                    mod['modifications'] = json.loads(mod['modifications'])
                except:
                    pass
            if mod.get('attachments'):
                try:
                    mod['attachments'] = json.loads(mod['attachments'])
                except:
                    pass
            mods.append(mod)
        
        return jsonify({
            'success': True,
            'data': mods,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    finally:
        conn.close()


@app.route('/api/order-modifications/<int:mod_id>/approve', methods=['POST'])
@login_required
def approve_order_modification(mod_id):
    """审批订单修改申请"""
    conn = get_conn()
    try:
        user = get_current_user()
        if not user or user.get('role') != '超管':
            return jsonify({'success': False, 'error': '需要管理员权限'}), 403
        
        mod = conn.execute("SELECT * FROM order_modifications WHERE id = ?", (mod_id,)).fetchone()
        if not mod:
            return jsonify({'success': False, 'error': '申请不存在'}), 404
        
        if mod['status'] != 'pending':
            return jsonify({'success': False, 'error': '该申请已被处理'}), 400
        
        data = request.json or {}
        approve_note = data.get('note', '')
        
        # 获取订单原始数据
        order = conn.execute("SELECT * FROM purchase_orders WHERE id = ?", (mod['order_id'],)).fetchone()
        before_data = dict(order)
        
        # 执行修改
        modifications = json.loads(mod['modifications']) if isinstance(mod['modifications'], str) else mod['modifications']
        fields, values = [], []
        for field, change in modifications.items():
            new_value = change.get('new', order.get(field))
            if field in ['title', 'supplier', 'order_date', 'expected_date', 'purchase_by', 'invoice_no', 'contract_no', 'notes', 'expense_order_no']:
                fields.append(f"{field} = ?")
                values.append(new_value)
            elif field == 'total_amount':
                fields.append("total_amount = ?")
                values.append(float(new_value) if new_value else 0)
        
        if fields:
            values.append(datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
            values.append(mod['order_id'])
            conn.execute(
                f"UPDATE purchase_orders SET {', '.join(fields)}, updated_at = ? WHERE id = ?",
                values
            )
        
        # 更新申请状态
        conn.execute('''
            UPDATE order_modifications 
            SET status = 'approved', approver = ?, approve_time = datetime('now','localtime'), approve_note = ?
            WHERE id = ?
        ''', (user['name'], approve_note, mod_id))
        
        # 记录日志
        log_action('approve', 'order', 'purchase_order', mod['order_id'], mod['order_no'], mod['title'],
                  detail={'modification_id': mod_id, 'approve_note': approve_note, 'applier': mod['apply_by']},
                  before_data=before_data)
        
        conn.commit()
        return jsonify({'success': True, 'message': '已通过修改申请'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/order-modifications/<int:mod_id>/reject', methods=['POST'])
@login_required
def reject_order_modification(mod_id):
    """拒绝订单修改申请"""
    conn = get_conn()
    try:
        user = get_current_user()
        if not user or user.get('role') != '超管':
            return jsonify({'success': False, 'error': '需要管理员权限'}), 403
        
        mod = conn.execute("SELECT * FROM order_modifications WHERE id = ?", (mod_id,)).fetchone()
        if not mod:
            return jsonify({'success': False, 'error': '申请不存在'}), 404
        
        if mod['status'] != 'pending':
            return jsonify({'success': False, 'error': '该申请已被处理'}), 400
        
        data = request.json or {}
        approve_note = data.get('note', '')
        
        # 更新申请状态
        conn.execute('''
            UPDATE order_modifications 
            SET status = 'rejected', approver = ?, approve_time = datetime('now','localtime'), approve_note = ?
            WHERE id = ?
        ''', (user['name'], approve_note, mod_id))
        
        # 记录日志
        log_action('reject', 'order', 'purchase_order', mod['order_id'], mod['order_no'], mod['title'],
                  detail={'modification_id': mod_id, 'reject_note': approve_note, 'applier': mod['apply_by']})
        
        conn.commit()
        return jsonify({'success': True, 'message': '已拒绝修改申请'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


# ─────────────────────────────────────────────
# 字段维护（可选下拉选项）
# ─────────────────────────────────────────────

@app.route('/api/field-options', methods=['GET'])
def get_field_options():
    """获取所有字段选项"""
    field_type = request.args.get('field_type', '')
    include_inactive = request.args.get('include_inactive', '0') == '1'
    active_filter = '' if include_inactive else ' AND active=1'
    conn = get_conn()
    if field_type:
        rows = conn.execute(
            f"SELECT * FROM field_options WHERE field_type=?{active_filter} ORDER BY sort_order, id",
            (field_type,)
        ).fetchall()
    else:
        rows = conn.execute(
            f"SELECT * FROM field_options WHERE 1=1{active_filter} ORDER BY field_type, sort_order, id"
        ).fetchall()
    conn.close()
    return jsonify({'success': True, 'data': rows_to_list(rows)})


@app.route('/api/field-options/<field_type>', methods=['GET'])
def get_field_options_by_type(field_type):
    """按类型获取字段选项"""
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM field_options WHERE field_type=? AND active=1 ORDER BY sort_order, id",
        (field_type,)
    ).fetchall()
    conn.close()
    return jsonify({'success': True, 'data': rows_to_list(rows)})


@app.route('/api/field-options', methods=['POST'])
@login_required
def create_field_option():
    """创建字段选项"""
    data = request.json or {}
    if not data.get('field_type') or not data.get('field_code') or not data.get('field_name'):
        return jsonify({'success': False, 'error': '类型、代码、名称不能为空'}), 400

    user = get_current_user()
    conn = get_conn()
    try:
        conn.execute('''
            INSERT INTO field_options (field_type, field_code, field_name, field_value, color, sort_order)
            VALUES (?, ?, ?, ?, ?, ?)
        ''', (
            data['field_type'],
            data['field_code'],
            data['field_name'],
            json.dumps(data.get('field_value', {})),
            data.get('color', ''),
            data.get('sort_order', 0)
        ))
        conn.commit()

        # 如果是资产类别(category)，自动同步到asset_type_params表
        if data['field_type'] == 'category':
            existing = conn.execute("SELECT id FROM asset_type_params WHERE type_code=?", (data['field_code'],)).fetchone()
            if not existing:
                conn.execute(
                    "INSERT INTO asset_type_params (type_code, type_name, depreciation_months, residual_rate, max_issue_years) VALUES (?,?,36,0.05,2)",
                    (data['field_code'], data['field_name'])
                )
                conn.commit()

        # 记录日志
        log_action('create', 'field_option', 'field_option', None, data['field_code'], data['field_name'],
                  detail={'field_type': data['field_type']})

        return jsonify({'success': True, 'message': '创建成功'})
    except Exception as e:
        if 'UNIQUE constraint' in str(e):
            return jsonify({'success': False, 'error': '该选项已存在'}), 400
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/field-options/<int:option_id>', methods=['PUT'])
@login_required
def update_field_option(option_id):
    """更新字段选项"""
    data = request.json or {}
    user = get_current_user()
    conn = get_conn()
    try:
        option = conn.execute("SELECT * FROM field_options WHERE id = ?", (option_id,)).fetchone()
        if not option:
            return jsonify({'success': False, 'error': '选项不存在'}), 404

        before_data = dict(option)

        conn.execute('''
            UPDATE field_options
            SET field_code=?, field_name=?, field_value=?, color=?, sort_order=?, active=?
            WHERE id=?
        ''', (
            data.get('field_code', option['field_code']),
            data.get('field_name', option['field_name']),
            data.get('field_value', option['field_value'] or ''),
            data.get('color', option['color']),
            data.get('sort_order', option['sort_order']),
            data.get('active', option['active']),
            option_id
        ))
        conn.commit()

        # 如果是资产类别(category)，同步更新asset_type_params表
        if option['field_type'] == 'category':
            new_active = data.get('active', option['active'])
            new_field_name = data.get('field_name', option['field_name'])
            new_field_code = data.get('field_code', option['field_code'])
            # 如果停用了，从参数表中删除
            if new_active == 0:
                conn.execute("DELETE FROM asset_type_params WHERE type_code=?", (option['field_code'],))
            else:
                # 检查是否已存在
                existing = conn.execute("SELECT id FROM asset_type_params WHERE type_code=?", (option['field_code'],)).fetchone()
                if existing:
                    # 更新名称
                    conn.execute("UPDATE asset_type_params SET type_name=? WHERE type_code=?", 
                                 (new_field_name, option['field_code']))
                else:
                    # 新增
                    conn.execute(
                        "INSERT INTO asset_type_params (type_code, type_name, depreciation_months, residual_rate, max_issue_years) VALUES (?,?,36,0.05,2)",
                        (new_field_code, new_field_name)
                    )
            conn.commit()

        # 记录日志
        log_action('update', 'field_option', 'field_option', option_id, option['field_code'], option['field_name'])

        return jsonify({'success': True, 'message': '更新成功'})
    except Exception as e:
        if 'UNIQUE constraint' in str(e):
            return jsonify({'success': False, 'error': '该选项已存在'}), 400
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/field-options/<int:option_id>', methods=['DELETE'])
@login_required
def delete_field_option(option_id):
    """删除字段选项（软删除）"""
    user = get_current_user()
    if not user or user.get('role') != '超管':
        return jsonify({'success': False, 'error': '需要管理员权限'}), 403

    conn = get_conn()
    try:
        option = conn.execute("SELECT * FROM field_options WHERE id = ?", (option_id,)).fetchone()
        if not option:
            return jsonify({'success': False, 'error': '选项不存在'}), 404

        conn.execute("UPDATE field_options SET active=0 WHERE id=?", (option_id,))
        conn.commit()

        # 如果是资产类别(category)，从asset_type_params表中删除
        if option['field_type'] == 'category':
            conn.execute("DELETE FROM asset_type_params WHERE type_code=?", (option['field_code'],))
            conn.commit()

        # 记录日志
        log_action('delete', 'field_option', 'field_option', option_id, option['field_code'], option['field_name'])

        return jsonify({'success': True, 'message': '删除成功'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/field-options/types', methods=['GET'])
def get_field_option_types():
    """获取所有字段类型列表"""
    conn = get_conn()
    rows = conn.execute(
        "SELECT DISTINCT field_type FROM field_options ORDER BY field_type"
    ).fetchall()
    conn.close()
    return jsonify({'success': True, 'data': [r['field_type'] for r in rows]})


# ═══════════════════════════════════════════════════════════════
# 普通用户：我的资产 & 申请管理
# ═══════════════════════════════════════════════════════════════

@app.route('/api/my/assets', methods=['GET'])
@login_required
def my_assets():
    """获取当前用户名下的资产"""
    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (request.current_user_id,)).fetchone()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404
        
        # 资产表以 assignee(关联人姓名) 关联用户，按姓名精确匹配
        rows = conn.execute('''
            SELECT a.*
            FROM assets a
            WHERE a.assignee = ?
            ORDER BY a.updated_at DESC
        ''', (user['name'],)).fetchall()

        return jsonify({'success': True, 'data': enrich_assets_book_value(assets_to_safe_list(rows), conn), 'total': len(rows)})
    finally:
        conn.close()


@app.route('/api/my/applications', methods=['GET'])
@login_required
def my_applications():
    """获取当前用户的所有申请记录"""
    conn = get_conn()
    try:
        app_type = request.args.get('app_type', '')
        status = request.args.get('status', '')
        
        sql = "SELECT * FROM asset_applications WHERE applicant_id = ?"
        params = [request.current_user_id]
        
        if app_type:
            sql += " AND app_type = ?"
            params.append(app_type)
        if status:
            sql += " AND status = ?"
            params.append(status)
        
        sql += " ORDER BY apply_time DESC"
        rows = conn.execute(sql, params).fetchall()
        return jsonify({'success': True, 'data': rows_to_list(rows)})
    finally:
        conn.close()


@app.route('/api/my/applications', methods=['POST'])
@login_required
def create_application():
    """提交申请（新申请/借用/归还/维修等）"""
    data = request.json or {}
    app_type = data.get('app_type')  # new_asset / borrow / return / repair / consumable
    
    valid_types = ['new_asset', 'borrow', 'return', 'repair', 'consumable', 'replace', 'special']
    if app_type not in valid_types:
        return jsonify({'success': False, 'error': '无效的申请类型'}), 400
    
    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (request.current_user_id,)).fetchone()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404

        app_id, app_no = insert_asset_application(
            conn, user, app_type,
            asset_id=data.get('asset_id'),
            asset_no=data.get('asset_no', ''),
            title=data.get('title') or '',
            reason=data.get('reason', ''),
            detail=data.get('detail', {}),
        )
        conn.commit()

        log_action('apply', 'asset_application', 'asset_applications', app_id,
                   app_no, data.get('title') or app_no)

        return jsonify({
            'success': True,
            'app_no': app_no,
            'id': app_id,
            'message': '申请已提交，请等待审批'
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/my/consumables', methods=['GET'])
@login_required
def my_consumables():
    """获取当前用户领用过的耗材（关联该用户的配件）"""
    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (request.current_user_id,)).fetchone()
        if not user:
            return jsonify({'success': False, 'error': '用户不存在'}), 404

        # 通过 consumable_logs 的领用(出库)记录，按姓名匹配关联人
        rows = conn.execute('''
            SELECT c.id, c.item_no, c.name, c.category, c.brand, c.model, c.unit, c.location,
                   SUM(cl.quantity) AS received_qty,
                   MAX(cl.log_time) AS last_receive_time
            FROM consumable_logs cl
            JOIN consumables c ON cl.item_id = c.id
            WHERE cl.target_user = ? AND cl.log_type = '出库'
            GROUP BY c.id
            ORDER BY last_receive_time DESC
        ''', (user['name'],)).fetchall()
        return jsonify({'success': True, 'data': rows_to_list(rows), 'total': len(rows)})
    finally:
        conn.close()


@app.route('/api/admin/applications', methods=['GET'])
@login_required
def list_all_applications():
    """获取所有申请列表（管理员）"""
    conn = get_conn()
    try:
        user = conn.execute("SELECT role FROM employees WHERE id = ?", (request.current_user_id,)).fetchone()
        if not user or user['role'] not in ('超管', '管理员'):
            return jsonify({'success': False, 'error': '无权限'}), 403
        
        status = request.args.get('status', '')
        app_type = request.args.get('app_type', '')
        
        sql = "SELECT * FROM asset_applications WHERE 1=1"
        params = []
        if status:
            sql += " AND status = ?"
            params.append(status)
        if app_type:
            sql += " AND app_type = ?"
            params.append(app_type)
        sql += " ORDER BY apply_time DESC"
        
        rows = conn.execute(sql, params).fetchall()
        return jsonify({'success': True, 'data': rows_to_list(rows)})
    finally:
        conn.close()


@app.route('/api/admin/applications/<int:app_id>/review', methods=['POST'])
@login_required
def review_application(app_id):
    """审批申请（管理员/超管）—— 原子抢占，与 approve_aggregated 一致"""
    conn = get_conn()
    try:
        user = conn.execute("SELECT * FROM employees WHERE id = ?", (request.current_user_id,)).fetchone()
        if not user or user['role'] not in ('超管', '管理员'):
            return jsonify({'success': False, 'error': '无权限'}), 403
        
        data = request.json or {}
        action = data.get('action')  # approve / reject
        comment = data.get('comment', '')
        
        if action not in ('approve', 'reject'):
            return jsonify({'success': False, 'error': '无效操作'}), 400

        cur = conn.execute(
            "UPDATE asset_applications SET status='processing' WHERE id=? AND status='pending'",
            (app_id,))
        if cur.rowcount == 0:
            return jsonify({'success': False, 'error': '申请不存在或已处理'}), 400
        appl = conn.execute("SELECT * FROM asset_applications WHERE id = ?", (app_id,)).fetchone()
        
        new_status = 'approved' if action == 'approve' else 'rejected'
        if action == 'approve':
            try:
                execute_asset_application(conn, dict(appl))
            except AssetAppExecError as e:
                conn.execute("UPDATE asset_applications SET status='pending' WHERE id=?", (app_id,))
                conn.commit()
                return jsonify({'success': False, 'error': str(e)}), 400
        conn.execute('''
            UPDATE asset_applications SET
                status = ?, reviewer_id = ?, reviewer_name = ?,
                review_comment = ?, review_time = datetime('now','localtime')
            WHERE id = ?
        ''', (new_status, user['id'], user['name'], comment, app_id))
        conn.commit()
        
        log_action('approve' if action == 'approve' else 'reject',
                   'asset_application', 'asset_applications', app_id,
                   appl['app_no'], appl['title'])
        
        return jsonify({'success': True, 'message': '审批完成'})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/approvals/fulfillable', methods=['GET'])
@login_required
def list_fulfillable_applications():
    """待发放：已通过且未发放的新资产/更换/特殊申请"""
    user = get_current_user()
    if not user or user.get('role') not in ADMIN_ROLES:
        return jsonify({'success': False, 'error': '需要管理权限', 'code': 403}), 403
    conn = get_conn()
    try:
        rows = conn.execute("""
            SELECT a.*, p.order_no AS po_order_no
            FROM asset_applications a
            LEFT JOIN purchase_orders p ON p.id = a.po_id
            WHERE a.status='approved' AND IFNULL(a.fulfilled,0)=0
              AND a.app_type IN ('new_asset','replace','special')
            ORDER BY a.review_time DESC, a.id DESC
        """).fetchall()
        return jsonify({'success': True, 'data': rows_to_list(rows), 'total': len(rows)})
    finally:
        conn.close()


@app.route('/api/approvals/<int:app_id>/create-order', methods=['POST'])
@login_required
def create_order_from_application(app_id):
    """待发放申请一键生成采购订单（金额先记 0，不猜单价）"""
    user = get_current_user()
    if not user or user.get('role') not in ADMIN_ROLES:
        return jsonify({'success': False, 'error': '需要管理权限', 'code': 403}), 403
    conn = get_conn()
    try:
        appr = conn.execute("SELECT * FROM asset_applications WHERE id=?", (app_id,)).fetchone()
        if not appr:
            return jsonify({'success': False, 'error': '申请不存在'}), 404
        if appr['status'] != 'approved' or (appr['fulfilled'] or 0) != 0:
            return jsonify({'success': False, 'error': '仅已通过且未发放的申请可生成订单'}), 400
        if appr['app_type'] not in ('new_asset', 'replace', 'special'):
            return jsonify({'success': False, 'error': '该申请类型不支持生成采购订单'}), 400
        if appr['po_id']:
            existing = conn.execute(
                "SELECT order_no FROM purchase_orders WHERE id=?", (appr['po_id'],)
            ).fetchone()
            return jsonify({
                'success': False,
                'error': '已关联采购订单',
                'order_id': appr['po_id'],
                'order_no': existing['order_no'] if existing else None,
            }), 400

        detail = {}
        try:
            detail = json.loads(appr['detail'] or '{}')
        except Exception:
            detail = {}
        category = (detail.get('category') or '').strip() or '固定资产'
        title = (appr['title'] or '').strip() or f"采购申请 {appr['app_no'] or app_id}"
        notes = f"来自申请 {appr['app_no'] or app_id} / {appr['applicant_name'] or ''}"
        if appr['reason']:
            notes = f"{notes}；{appr['reason']}"

        order_no = gen_order_no()
        cur = conn.execute('''
            INSERT INTO purchase_orders
            (order_no, title, supplier, order_date, total_amount, currency, status,
             purchase_by, notes, created_by)
            VALUES (?,?,?,?,?,?,?,?,?,?)
        ''', (
            order_no, title, '', datetime.now().strftime('%Y-%m-%d'), 0, 'CNY', '待到货',
            user.get('name') or '', notes, user.get('name') or '',
        ))
        order_id = cur.lastrowid
        conn.execute('''
            INSERT INTO order_items (order_id, item_type, item_name, category, quantity, unit_price, notes)
            VALUES (?,?,?,?,?,?,?)
        ''', (order_id, 'asset', title, category, 1, 0, appr['reason'] or ''))
        conn.execute("UPDATE asset_applications SET po_id=? WHERE id=?", (order_id, app_id))
        conn.commit()
        log_action('create', 'order', 'purchase_orders', order_id, order_no, user.get('name'),
                   {'from_application': appr['app_no'] or app_id})
        return jsonify({
            'success': True, 'message': '已生成采购订单',
            'order_id': order_id, 'order_no': order_no,
        })
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


@app.route('/api/approvals/<int:app_id>/fulfill', methods=['POST'])
@login_required
def fulfill_application(app_id):
    """确认发放：选定在库资产出库给申请人，并标记申请已发放"""
    user = get_current_user()
    if not user or user.get('role') not in ADMIN_ROLES:
        return jsonify({'success': False, 'error': '需要管理权限', 'code': 403}), 403
    data = request.json or {}
    asset_id = data.get('asset_id')
    if not asset_id:
        return jsonify({'success': False, 'error': '请选择要发放的在库资产'}), 400
    conn = get_conn()
    try:
        # 抢占未发放的已通过申请
        cur = conn.execute(
            """UPDATE asset_applications SET fulfilled=2 WHERE id=? AND status='approved'
               AND IFNULL(fulfilled,0)=0 AND app_type IN ('new_asset','replace','special')""",
            (app_id,))
        if cur.rowcount == 0:
            return jsonify({'success': False, 'error': '申请不存在、未通过或已发放'}), 400
        appr = conn.execute("SELECT * FROM asset_applications WHERE id=?", (app_id,)).fetchone()
        asset = conn.execute("SELECT * FROM assets WHERE id=?", (asset_id,)).fetchone()
        if not asset:
            conn.execute("UPDATE asset_applications SET fulfilled=0 WHERE id=?", (app_id,))
            conn.commit()
            return jsonify({'success': False, 'error': '资产不存在'}), 404
        if asset['status'] != '在库':
            conn.execute("UPDATE asset_applications SET fulfilled=0 WHERE id=?", (app_id,))
            conn.commit()
            return jsonify({'success': False, 'error': f'资产当前状态为【{asset["status"]}】，无法发放'}), 400
        applicant = appr['applicant_name']
        conn.execute(
            "UPDATE assets SET status='在用', assignee=?, updated_at=datetime('now','localtime') WHERE id=?",
            (applicant, asset_id))
        conn.execute('''INSERT INTO stock_logs(asset_id,asset_no,log_type,operator,target_user,from_location,to_location,remark)
                        VALUES(?,?,?,?,?,?,?,?)''',
                     (asset_id, asset['asset_no'], '出库', user.get('name') or '系统', applicant,
                      asset['location'], asset['location'],
                      f"发放申请 {appr['app_no'] or app_id}"))
        conn.execute(
            "UPDATE asset_applications SET fulfilled=1, asset_id=?, asset_no=?, close_time=datetime('now','localtime') WHERE id=?",
            (asset_id, asset['asset_no'], app_id))
        conn.commit()
        log_action('approve', 'approval', 'asset_application', app_id, appr['app_no'], applicant,
                   {'fulfill_asset': asset['asset_no']})
        return jsonify({'success': True, 'message': '已发放', 'asset_no': asset['asset_no']})
    except Exception as e:
        conn.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        conn.close()


if __name__ == '__main__':
    import sys
    import threading
    sys.path.insert(0, os.path.dirname(__file__))
    try:
        init_db()
    except Exception as e:
        print(f"[WARN] init_db: {e}", flush=True)

    # 启动每日续期提醒定时任务（每天 09:00 推送）
    def _schedule_expire_reminder():
        import time as _time
        while True:
            now = datetime.now()
            target = now.replace(hour=9, minute=0, second=0, microsecond=0)
            if now >= target:
                from datetime import timedelta
                target += timedelta(days=1)
            _time.sleep((target - now).total_seconds())
            try:
                push_expire_reminder()
            except Exception as e:
                print(f"[expire-reminder] 推送失败: {e}")

    threading.Thread(target=_schedule_expire_reminder, daemon=True).start()
    print("[OK] expire-reminder scheduler started (daily 09:00)", flush=True)

    _port = int(os.environ.get('PORT', 8088))
    app.run(host='0.0.0.0', port=_port, debug=False)
