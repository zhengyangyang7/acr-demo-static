# -*- coding: utf-8 -*-
"""手填官方目标：从 JSON / Excel 读入，按团队长姓名写入分层表。"""
import json
import os
from dataclasses import dataclass

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


@dataclass
class FillResult:
    written_teams: int
    company_written: bool
    column: int
    sheet: str


def _norm_header(value):
    if value is None:
        return ""
    return str(value).replace("\n", "").strip()


def _find_h2_task_col(ws):
    for cell in ws[2]:
        v = _norm_header(cell.value)
        if (
            "财富中心" in v
            and "下半年" in v
            and "任务" in v
            and "完成率" not in v
            and "差距" not in v
        ):
            return cell.column
    raise ValueError("未找到「财富中心下半年任务」列")


def _pick_week_sheet(wb):
    latest = None
    latest_key = -1
    for name in wb.sheetnames:
        if name.startswith("周度团队数据-折标+考核"):
            suffix = name.replace("周度团队数据-折标+考核", "")
            try:
                key = int(suffix)
            except ValueError:
                key = 0
            if key >= latest_key:
                latest_key = key
                latest = name
    if latest:
        return wb[latest]
    return wb.active


def _scan_team_anchors(ws):
    """返回 [(leader_name, anchor_row, sum_row), ...]，公司合计行不含在内。"""
    groups = []
    cur_leader = None
    cur_start = None
    cur_end = None
    for row_idx in range(4, ws.max_row + 1):
        a_val = ws.cell(row_idx, 1).value
        b_val = ws.cell(row_idx, 2).value
        is_sum = a_val and "合计" in str(a_val)
        has_name = b_val and isinstance(b_val, str) and b_val.strip() and not is_sum
        if is_sum:
            if cur_start is not None:
                leader = (cur_leader or "").strip()
                if not leader:
                    name = ws.cell(cur_start, 2).value
                    leader = str(name).strip() if name else ""
                groups.append((leader, cur_start, row_idx))
            cur_leader = None
            cur_start = None
            cur_end = None
        elif has_name:
            new_leader = str(a_val).strip() if a_val else None
            if cur_start is None:
                cur_leader = new_leader
                cur_start = row_idx
            cur_end = row_idx
    return groups


def _company_goal_row(ws, team_sum_rows):
    company_row = None
    for r in range(4, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a and "合计" in str(a) and r not in team_sum_rows:
            company_row = r
    if company_row is None:
        for r in range(4, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if a and "合计" in str(a):
                company_row = r
    return company_row + 1 if company_row else None


def _set_value(ws, row, col, value):
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= row <= mr.max_row and mr.min_col <= col <= mr.max_col:
                ws.cell(mr.min_row, mr.min_col).value = value
                return
        return
    cell.value = value


def _load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    teams = {}
    for item in data.get("teams") or []:
        leader = str(item.get("leader") or "").strip()
        if not leader:
            continue
        teams[leader] = item["target"]
    return teams, data.get("company_goal")


def _header_row_map(ws):
    mapping = {}
    for cell in ws[1]:
        v = _norm_header(cell.value)
        if v:
            mapping[v] = cell.column
    return mapping


def _load_excel(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = _header_row_map(ws)
    leader_col = None
    target_col = None
    for name, col in headers.items():
        if name == "团队长" or name == "理财师":
            leader_col = col
        if "下半年任务" in name.replace(" ", "") and "完成率" not in name and "差距" not in name:
            target_col = col
    if not leader_col or not target_col:
        wb.close()
        raise ValueError("目标 Excel 需包含「团队长」和「下半年任务」列")
    teams = {}
    company_goal = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        leader = row[leader_col - 1]
        target = row[target_col - 1]
        if leader is None or target is None:
            continue
        name = str(leader).strip()
        if name == "公司目标":
            company_goal = target
        else:
            teams[name] = target
    wb.close()
    return teams, company_goal


def load_official_targets(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".json":
        return _load_json(path)
    if ext in (".xlsx", ".xlsm"):
        return _load_excel(path)
    raise ValueError(f"不支持的目标文件类型: {ext}")


def fill_official_targets(workbook_path, targets_path, output_path=None):
    teams, company_goal = load_official_targets(targets_path)
    wb = load_workbook(workbook_path)
    try:
        ws = _pick_week_sheet(wb)
        col = _find_h2_task_col(ws)
        groups = _scan_team_anchors(ws)
        written = 0
        for leader, anchor, _sum_row in groups:
            if leader in teams:
                _set_value(ws, anchor, col, teams[leader])
                written += 1
        team_sum_rows = {g[2] for g in groups}
        company_written = False
        goal_row = _company_goal_row(ws, team_sum_rows)
        if company_goal is not None and goal_row:
            _set_value(ws, goal_row, col, company_goal)
            company_written = True
        dest = output_path or workbook_path
        sheet_title = ws.title
        wb.save(dest)
    finally:
        wb.close()
    return FillResult(
        written_teams=written,
        company_written=company_written,
        column=col,
        sheet=sheet_title,
    )
